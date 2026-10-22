#!/usr/bin/env python3
"""Cobertura documental por proveedor (sin modelo de IA).

Mercado Publico es desordenado: hay proveedores que suben solo anexos tecnicos, solo economicos o solo
administrativos, un PDF por producto, o escaneados sin texto. Este modulo clasifica los documentos de
cada proveedor y, para cada producto dudoso, distingue de quien es la falta:

  COMPLETO               el producto tiene sus datos y paso la validacion.
  DATO_NO_EXISTE         el proveedor no entrego el dato en ningun documento (nada que extraer).
  NO_ENCONTRADO          el dato probablemente existe en un documento legible pero el pipeline no lo extrajo.
  PENDIENTE_OCR          el dato podria estar en documentos sin texto utilizable (escaneados, ilegibles,
                         formatos no soportados) que aun no se leen.
  DATOS_CONTRADICTORIOS  hay datos pero los documentos se contradicen (validacion "inconsistente").
  REVISAR                tiene todos los datos pero la validacion dejo alertas.

Se usa desde 3_extraer_ia.py (columnas "Cobertura documental" y "Motivo" del Excel) y como informe:

  python cobertura_documental.py --dir lotes_z8/2025-11/ofertas --excel cobertura.xlsx
"""

from __future__ import annotations

import argparse
import glob
import importlib.util
import json
import re
import sys
import types
import unicodedata
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parent

COMPLETO = "COMPLETO"
DATO_NO_EXISTE = "DATO_NO_EXISTE"
NO_ENCONTRADO = "NO_ENCONTRADO"
PENDIENTE_OCR = "PENDIENTE_OCR"
CONTRADICTORIO = "DATOS_CONTRADICTORIOS"
REVISAR = "REVISAR"

ROLES_POR_PREFIJO = {"economico": "economico", "tecnico": "tecnico", "administrativo": "administrativo"}
EXTENSIONES_COMPRIMIDAS = (".zip", ".rar", ".7z")
EXTENSIONES_IMAGEN = (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp")

PATRON_MONTO = re.compile(r"(?<![\w.,])(\d{1,3}(?:[.,]\d{3})+)(?![\d])")
PATRON_CLAVE_PRECIO = re.compile(
    r"\$|\b(?:neto|iva|valor\s*(?:unit|total|neto)|precio|p\.?\s*unit|monto|subtotal)", re.I)
PALABRAS_GENERICAS = {
    "notebook", "notebooks", "computador", "computadores", "equipo", "equipos", "monitor", "monitores", "impresora",
    "impresoras", "intel", "core", "para", "con", "similar", "marca", "modelo", "pulgadas", "pantalla", "memoria",
    "disco", "windows", "procesador", "unidad", "unidades", "tipo", "linea", "serie", "nuevo", "nueva"}
PATRON_CLAVE_CANTIDAD = re.compile(r"\b(?:cantidad|cant\.?|unidades|unid\.?|qty|n[°º]\s*de\s*equipos)\b", re.I)
PATRON_ADMINISTRATIVO = re.compile(
    r"declaraci[oó]n\s+jurada|garant[ií]a\s+de\s+seriedad|poder\s+notarial|certificado\s+de|"
    r"inhabilidad|pacto\s+de\s+integridad|formulario\s+de\s+identificaci", re.I)
PATRON_TECNICO = re.compile(
    r"procesador|memoria\s+ram|\bram\b|\bssd\b|disco\s+duro|pantalla|sistema\s+operativo|resoluci[oó]n|"
    r"garant[ií]a\s+\d|especificaci[oó]n", re.I)

_E = None


def cargar_extractor():
    """3_extraer_ia.py como modulo (su nombre empieza con un digito y no se puede importar a secas)."""
    global _E
    if _E is None:
        spec = importlib.util.spec_from_file_location("extractor_ia_cobertura", ROOT / "3_extraer_ia.py")
        modulo = importlib.util.module_from_spec(spec)
        sys.modules.setdefault("extractor_ia_cobertura", modulo)
        spec.loader.exec_module(modulo)
        _E = modulo
    return _E


ARGS_LECTURA = types.SimpleNamespace(
    max_paginas=200, ocr=False, vision=False, tesseract_cmd=None, idioma_ocr="spa+eng",
    dpi_ocr=200, vision_min_filas=2, max_filas_excel=5000, max_columnas_excel=40,
)


def sin_tildes(texto):
    return unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii").lower()


# ---------------------------------------------------------------------------
# Analisis de un documento
# ---------------------------------------------------------------------------
def montos_plausibles(texto):
    """Montos escritos con separador de miles (1.234.567) entre 1.000 y 2.000.000.000."""
    montos = []
    for coincidencia in PATRON_MONTO.finditer(texto):
        digitos = re.sub(r"[.,]", "", coincidencia.group(1))
        if 1000 <= int(digitos) <= 2_000_000_000:
            montos.append(int(digitos))
    return montos


def rol_del_documento(nombre, texto=""):
    """Rol segun el prefijo que pone el scraper (economico__, tecnico__, administrativo__) y, si no lo
    trae, segun el nombre y el contenido."""
    prefijo = str(nombre).replace("\\", "/").rsplit("/", 1)[-1].split("__", 1)[0].lower()
    if prefijo in ROLES_POR_PREFIJO:
        return ROLES_POR_PREFIJO[prefijo]
    tipo = cargar_extractor().tipo_documental(nombre)
    if tipo in ("economico", "tecnico"):
        return tipo
    if texto:
        if PATRON_ADMINISTRATIVO.search(texto):
            return "administrativo"
        if montos_plausibles(texto) and PATRON_CLAVE_PRECIO.search(texto):
            return "economico"
        if len(PATRON_TECNICO.findall(texto)) >= 3:
            return "tecnico"
    return "otro"


def lineas_con_precio(lineas):
    """Indices de lineas con un indicio monetario ($, neto, iva, valor, precio...) y un monto plausible."""
    return [i for i, linea in enumerate(lineas)
            if PATRON_CLAVE_PRECIO.search(linea) and montos_plausibles(linea)]


def analizar_documento(ruta, relativo, conservar_lineas=False):
    """Resumen de un archivo: rol, estado de lectura y que datos contiene (con sus lineas si se piden)."""
    E = cargar_extractor()
    ruta = Path(ruta)
    extension = ruta.suffix.lower()
    resumen = {"archivo": relativo, "rol": rol_del_documento(relativo), "lectura": "legible", "montos": 0,
               "tiene_precios": False, "tiene_cantidad": False, "lineas_equipo": 0, "marcas": []}
    if extension in EXTENSIONES_COMPRIMIDAS:
        resumen["lectura"] = "comprimido"
        return resumen
    try:
        texto, estado = E.extraer_archivo(ruta, ARGS_LECTURA)
    except Exception as exc:  # documento corrupto: se informa, no se detiene el informe
        resumen["lectura"] = "error_lectura"
        resumen["detalle"] = str(exc)[:120]
        return resumen
    if not texto.strip():
        escaneable = extension == ".pdf" or extension in EXTENSIONES_IMAGEN
        resumen["lectura"] = "sin_texto" if escaneable else "no_soportado"
        resumen["detalle"] = str(estado or "")
        return resumen
    if E.texto_ilegible(texto):
        resumen["lectura"] = "ilegible"
        return resumen
    resumen["rol"] = "historial" if E.documento_de_historial(relativo, texto) else rol_del_documento(relativo, texto)
    lineas = [linea.strip() for linea in texto.splitlines() if linea.strip()]
    resumen["montos"] = len(montos_plausibles(texto))
    resumen["lineas_precio"] = len(lineas_con_precio(lineas))
    resumen["tiene_precios"] = resumen["lineas_precio"] > 0
    resumen["tiene_cantidad"] = bool(PATRON_CLAVE_CANTIDAD.search(texto))
    resumen["lineas_equipo"] = sum(1 for linea in lineas if len(linea) < 350 and E.PATRON_PRODUCTO.search(linea))
    plano = sin_tildes(texto)
    resumen["marcas"] = sorted(
        marca for clave, marca in E.MARCAS_CONOCIDAS.items()
        if len(clave) >= 3 and re.search(r"\b" + re.escape(sin_tildes(clave)) + r"\b", plano))
    if conservar_lineas:
        resumen["_lineas"] = lineas
    return resumen


# ---------------------------------------------------------------------------
# Clasificacion del proveedor
# ---------------------------------------------------------------------------
def marcar_uno_por_producto(documentos):
    """Un PDF por producto: 3 o mas archivos del mismo rol, con pocas lineas de equipo cada uno."""
    por_rol = {}
    for doc in documentos:
        if doc["lectura"] == "legible" and doc["rol"] != "historial":
            por_rol.setdefault(doc["rol"], []).append(doc)
    for rol, docs in por_rol.items():
        if rol in ("economico", "tecnico", "otro") and len(docs) >= 3 and all(d["lineas_equipo"] <= 3 for d in docs):
            for doc in docs:
                doc["uno_por_producto"] = True


def resumir_proveedor(documentos, necesita_ocr=(), no_soportados=()):
    marcar_uno_por_producto(documentos)
    legibles = [d for d in documentos if d["lectura"] == "legible" and d["rol"] != "historial"]
    con_historial = any(d["rol"] == "historial" for d in documentos)
    sin_texto = [d["archivo"] for d in documentos
                 if d["lectura"] in ("sin_texto", "ilegible", "no_soportado", "error_lectura")]
    sin_texto += [a for a in necesita_ocr if a not in sin_texto]
    sin_texto += [a for a in no_soportados
                  if a not in sin_texto and Path(a).suffix.lower() not in EXTENSIONES_COMPRIMIDAS]
    # Carnets, RUT, certificados y declaraciones escaneados del anexo administrativo no son la oferta: que no se
    # puedan leer no deja ningun dato del producto "pendiente de OCR".
    es_ruido = cargar_extractor().es_ruido_administrativo
    sin_texto = [a for a in sin_texto if not es_ruido(a)]
    roles = sorted({d["rol"] for d in legibles})
    if not documentos:
        clase = "sin_documentos"
    elif not legibles:
        clase = "ningun_documento_legible"
    elif roles == ["administrativo"]:
        clase = "solo_administrativos"
    elif "economico" in roles and "tecnico" in roles:
        clase = "economico_y_tecnico"
    elif "economico" in roles:
        clase = "solo_economico"
    elif "tecnico" in roles:
        clase = "solo_tecnico"
    else:
        clase = "otros"
    if any(d.get("uno_por_producto") for d in legibles):
        clase += "+uno_por_producto"
    if con_historial:
        clase += "+con_historial"
    if sin_texto:
        clase += "+con_documentos_sin_texto"
    return {
        "clase": clase,
        "n_documentos": len(documentos),
        "roles": roles,
        "con_precios": sorted(d["archivo"] for d in legibles if d["tiene_precios"]),
        "con_cantidad": sorted(d["archivo"] for d in legibles if d["tiene_cantidad"]),
        "marcas": sorted({m for d in legibles for m in d["marcas"]}),
        "sin_texto": sorted(set(sin_texto)),
        "documentos": documentos,
    }


def analizar_proveedor(carpeta, resultado=None, productos=None):
    """Lee los documentos de la carpeta de un proveedor y los clasifica (no usa OCR ni el modelo).
    Si se pasan `productos`, agrega resumen["cobertura"]: [(cobertura, motivo)] en el mismo orden."""
    E = cargar_extractor()
    carpeta = Path(carpeta)
    resultado = resultado or {}
    documentos = []
    for ruta in E.listar_archivos_oferta(carpeta):
        relativo = str(ruta.relative_to(carpeta))
        if ruta.suffix.lower() in EXTENSIONES_COMPRIMIDAS:
            extraido = carpeta / "_extraidos" / ruta.stem
            if extraido.is_dir() and any(extraido.rglob("*.*")):
                continue  # comprimido ya expandido: sus archivos se analizan aparte
        if ruta.suffix.lower() in (".json", ".html"):
            continue
        documentos.append(analizar_documento(ruta, relativo, conservar_lineas=productos is not None))
    resumen = resumir_proveedor(documentos, resultado.get("necesita_ocr") or [], resultado.get("no_soportados") or [])
    if productos is not None:
        resumen["cobertura"] = [cobertura_producto(p, resumen, resultado.get("total_oferta")) for p in productos]
    for doc in resumen["documentos"]:
        doc.pop("_lineas", None)  # no se arrastran las lineas a los procesos que reciben el resumen
        doc.pop("_planas", None)
    return resumen


# ---------------------------------------------------------------------------
# Cobertura de un producto
# ---------------------------------------------------------------------------
def vacio(valor):
    return valor in (None, "")


def campos_faltantes(producto):
    cantidad, unitario, total = (producto.get(c) for c in ("cantidad", "precio_unitario", "precio_total"))
    faltan = []
    if vacio(unitario) and vacio(total):
        faltan.append("precio")
    if vacio(cantidad) and (vacio(unitario) or vacio(total)):
        faltan.append("cantidad")
    if vacio(producto.get("marca")):
        faltan.append("marca")
    return faltan


def _plano(doc):
    if "_planas" not in doc:
        doc["_planas"] = [sin_tildes(linea) for linea in doc.get("_lineas", [])]
    return doc["_planas"]


def apariciones_del_producto(producto, documentos):
    """[(doc, indice_de_linea)] donde el documento habla de este producto: por codigo de modelo o, si no tiene,
    por varias palabras distintivas de su descripcion en la misma linea (con sus vecinas)."""
    E = cargar_extractor()
    codigos = E.codigos_de_modelo(producto)
    texto = sin_tildes(f"{producto.get('producto') or ''} {producto.get('modelo') or ''}")
    palabras = {t for t in re.findall(r"[a-z0-9]{3,}", texto)
                if t not in PALABRAS_GENERICAS and (len(t) >= 5 or re.search(r"\d", t))}
    if not codigos and len(palabras) < 2:
        return []
    exigidas = min(3, len(palabras))
    hallazgos = []
    for doc in documentos:
        if doc["rol"] == "historial":
            continue  # ordenes de compra pasadas: no son lo que se oferta
        planas = _plano(doc)
        for i in range(len(planas)):
            ventana = " ".join(planas[max(0, i - 1):i + 2])
            if codigos:
                coincide = any(codigo in ventana for codigo in codigos)
            else:
                coincide = sum(palabra in ventana for palabra in palabras) >= exigidas
            if coincide:
                hallazgos.append((doc, i))
    return hallazgos


PATRON_LINEA_GENERICA = re.compile(
    r"\b(?:computador(?:es)?|notebooks?|equipos?|monitor(?:es)?|impresoras?|all\s*in\s*one|aio|pc|laptops?)\b", re.I)
PATRON_VALOR_UNITARIO = re.compile(r"(?:valor|precio)\s*unitario|p\.\s*unit", re.I)
PATRON_ENTERO_CHICO =re.compile(r"(?<![\w.,/-])\d{1,4}(?![\w.,/-])")


def _ventana(doc, indice):
    lineas = doc.get("_lineas", [])
    return " ".join(lineas[max(0, indice - 1):indice + 2])


def _evidencia_campo(campo, hallazgos, documentos, producto):
    """(archivo, nota) si algun documento legible trae el dato de este producto; None si no."""
    E = cargar_extractor()
    for doc, indice in hallazgos:
        ventana = _ventana(doc, indice)
        if campo == "precio" and montos_plausibles(ventana) and (
                PATRON_CLAVE_PRECIO.search(ventana) or len(montos_plausibles(doc["_lineas"][indice])) >= 2):
            return doc["archivo"], "el producto aparece junto a montos"
        if campo == "marca":
            plano = sin_tildes(ventana)
            for clave, marca in E.MARCAS_CONOCIDAS.items():
                if len(clave) >= 3 and re.search(r"\b" + re.escape(sin_tildes(clave)) + r"\b", plano):
                    return doc["archivo"], f"la linea del producto menciona {marca}"
        if campo == "cantidad":
            sin_codigos = re.sub(r"[A-Za-z0-9][A-Za-z0-9\-/]{3,}", " ", ventana)
            if PATRON_CLAVE_CANTIDAD.search(ventana) or (doc["rol"] == "economico" and PATRON_ENTERO_CHICO.search(sin_codigos)):
                return doc["archivo"], "la linea del producto trae una cantidad"
    if campo == "precio":
        # Linea economica generica ("computador", "equipo"...) con precio: el dato esta pero no se vincula a este
        # producto (caso del cruce tecnico-economico).
        for doc in documentos:
            if doc["rol"] != "economico":
                continue
            lineas = doc.get("_lineas", [])
            for i in lineas_con_precio(lineas):
                linea = lineas[i]
                if PATRON_VALOR_UNITARIO.search(linea):
                    return doc["archivo"], "formulario con valor unitario (falta vincularlo al producto)"
                if PATRON_LINEA_GENERICA.search(linea) and not re.search(r"\b(?=\w*\d)(?=\w*[A-Za-z])\w{5,}\b", linea):
                    return doc["archivo"], "linea economica generica con precio (falta cruzarla con la ficha tecnica)"
    return None


def _evaluar_campo(campo, resumen, producto, hallazgos, total_oferta):
    """(categoria, motivo) para un campo faltante del producto."""
    encontrado = _evidencia_campo(campo, hallazgos, resumen["documentos"], producto)
    if encontrado:
        return NO_ENCONTRADO, f"{campo}: {encontrado[1]} en {encontrado[0]}"
    pendientes = resumen["sin_texto"]
    if pendientes:
        return PENDIENTE_OCR, f"{campo}: {len(pendientes)} documento(s) sin texto utilizable ({pendientes[0]})"
    if resumen["clase"] == "sin_documentos":
        return DATO_NO_EXISTE, f"{campo}: el proveedor no subio documentos descargables"
    donde = (f"el producto aparece en {hallazgos[0][0]['archivo']} pero sin {campo}" if hallazgos
             else "el producto no figura con ese dato en los documentos legibles")
    extra = ""
    if campo == "precio" and total_oferta not in (None, ""):
        extra = f" (solo el total de la oferta en el portal: {total_oferta})"
    return DATO_NO_EXISTE, f"{campo}: {donde} ({resumen['clase']}){extra}"


def cobertura_producto(producto, resumen, total_oferta=None):
    """(cobertura, motivo) de un producto frente a los documentos de su proveedor.
    `resumen["documentos"]` debe traer las lineas (`_lineas`) para ubicar el producto."""
    estado = str(producto.get("estado_validacion") or "")
    faltan = campos_faltantes(producto)
    if not faltan:
        if estado == "inconsistente":
            alertas = "; ".join(producto.get("alertas") or [])
            return CONTRADICTORIO, alertas or "los documentos del proveedor traen datos contradictorios"
        if estado in ("", "ok"):
            return COMPLETO, ""
        return REVISAR, "; ".join(producto.get("alertas") or []) or f"estado {estado}"
    hallazgos = apariciones_del_producto(producto, resumen["documentos"])
    evaluados = [_evaluar_campo(campo, resumen, producto, hallazgos, total_oferta) for campo in faltan]
    for categoria in (NO_ENCONTRADO, PENDIENTE_OCR):
        if any(c == categoria for c, _ in evaluados):
            return categoria, " | ".join(m for c, m in evaluados)
    return DATO_NO_EXISTE, " | ".join(m for _, m in evaluados)


# ---------------------------------------------------------------------------
# Proveedores de un lote
# ---------------------------------------------------------------------------
def carpeta_del_proveedor(carpeta_licitacion, rut):
    rut = str(rut or "").strip()
    if not rut:
        return None
    for carpeta in sorted(Path(carpeta_licitacion).iterdir()):
        if carpeta.is_dir() and not carpeta.name.startswith("_") and rut in carpeta.name:
            return carpeta
    return None


def _analizar_tarea(tarea):
    carpeta, resultado, productos = tarea
    try:
        return analizar_proveedor(carpeta, resultado, productos)
    except Exception as exc:  # no tumbar el lote por un proveedor
        return {"clase": "error_analisis", "n_documentos": 0, "roles": [], "con_precios": [], "con_cantidad": [],
                "marcas": [], "sin_texto": [], "documentos": [], "error": str(exc)[:200],
                "cobertura": [("", f"error al analizar los documentos: {str(exc)[:80]}")] * len(productos)}


def analizar_proveedores(tareas, workers=2):
    """tareas: lista de (carpeta_proveedor, resultado_checkpoint, productos) -> resumenes en el mismo orden,
    cada uno con "cobertura": [(cobertura, motivo)] alineada con `productos`."""
    if workers <= 1 or len(tareas) < 4:
        return [_analizar_tarea(t) for t in tareas]
    with ProcessPoolExecutor(max_workers=workers) as pool:
        return list(pool.map(_analizar_tarea, tareas, chunksize=4))


def anotar_productos(productos, raiz_ofertas, workers=2, resultados=None):
    """Agrega 'cobertura_documental' y 'cobertura_motivo' a cada producto.

    productos: dicts con 'codigo' (licitacion) y 'rut'. resultados: {(codigo, rut): resultado_checkpoint}
    opcional (necesita_ocr / no_soportados / total_oferta)."""
    resultados = resultados or {}
    raiz_ofertas = Path(raiz_ofertas)
    por_proveedor = {}
    for producto in productos:
        clave = (producto.get("codigo"), str(producto.get("rut") or "").strip())
        por_proveedor.setdefault(clave, []).append(producto)
    tareas, indices = [], []
    for clave, lista in sorted(por_proveedor.items(), key=lambda kv: (str(kv[0][0]), kv[0][1])):
        carpeta_licitacion = raiz_ofertas / str(clave[0])
        carpeta = carpeta_del_proveedor(carpeta_licitacion, clave[1]) if carpeta_licitacion.is_dir() else None
        if carpeta is None:
            for producto in lista:
                producto["cobertura_documental"], producto["cobertura_motivo"] = "", "carpeta del proveedor no encontrada"
        else:
            tareas.append((carpeta, resultados.get(clave, {}), lista))
            indices.append(clave)
    resumenes = dict(zip(indices, analizar_proveedores(tareas, workers)))
    for clave, resumen in resumenes.items():
        for producto, (cobertura, motivo) in zip(por_proveedor[clave], resumen["cobertura"]):
            producto["cobertura_documental"], producto["cobertura_motivo"] = cobertura, motivo
    return resumenes


# ---------------------------------------------------------------------------
# Informe independiente
# ---------------------------------------------------------------------------
def cargar_checkpoints(directorios):
    """[(codigo_licitacion, carpeta_licitacion, resultado)] de todos los extraccion_ia.json."""
    E = cargar_extractor()
    salida = []
    for patron in directorios:
        for directorio in sorted(glob.glob(patron)):
            for ruta in sorted(Path(directorio).glob("*/extraccion_ia.json")):
                try:
                    resultados = json.loads(ruta.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                for resultado in resultados:
                    salida.append((ruta.parent.name, ruta.parent, resultado))
    return salida, E


def escribir_excel(filas_productos, filas_proveedores, resumen, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()

    def hoja(titulo, columnas, filas, primera=False):
        ws = wb.active if primera else wb.create_sheet()
        ws.title = titulo
        ws.append(columnas)
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
        for fila in filas:
            ws.append([fila.get(c) for c in columnas])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    hoja("Resumen", ["Estado", "Cobertura", "Productos"], resumen, primera=True)
    hoja("Productos", ["licitacion", "proveedor", "producto", "estado", "cantidad", "precio_unitario",
                       "precio_total", "marca", "cobertura", "motivo"], filas_productos)
    hoja("Proveedores", ["licitacion", "proveedor", "clase", "documentos", "con_precios", "sin_texto", "productos"],
         filas_proveedores)
    wb.save(ruta)


def main():
    parser = argparse.ArgumentParser(description="Cobertura documental de los proveedores de un lote")
    parser.add_argument("--dir", nargs="+", required=True, help="carpeta(s) ofertas (acepta comodines)")
    parser.add_argument("--excel", default="cobertura_documental.xlsx")
    parser.add_argument("--workers", type=int, default=2)
    args = parser.parse_args()

    checkpoints, E = cargar_checkpoints(args.dir)
    tareas, validos = [], []
    for indice, (codigo, carpeta_licitacion, resultado) in enumerate(checkpoints):
        carpeta = carpeta_del_proveedor(carpeta_licitacion, resultado.get("rut"))
        if carpeta is None:
            continue
        productos = E.validar_productos(resultado.get("productos", []), resultado.get("productos_parciales", []),
                                        resultado.get("total_oferta"))
        tareas.append((carpeta, resultado, productos))
        validos.append(indice)
    print(f"Proveedores: {len(checkpoints)} | con carpeta: {len(validos)} | analizando con {args.workers} procesos...")
    resumenes = analizar_proveedores(tareas, args.workers)

    filas_productos, filas_proveedores, conteo = [], [], Counter()
    for indice, (carpeta, _, productos), resumen in zip(validos, tareas, resumenes):
        codigo, _, resultado = checkpoints[indice]
        for producto, (cobertura, motivo) in zip(productos, resumen["cobertura"]):
            conteo[(producto.get("estado_validacion"), cobertura)] += 1
            filas_productos.append({
                "licitacion": codigo, "proveedor": resultado.get("proveedor"), "producto": producto.get("producto"),
                "estado": producto.get("estado_validacion"), "cantidad": producto.get("cantidad"),
                "precio_unitario": producto.get("precio_unitario"), "precio_total": producto.get("precio_total"),
                "marca": producto.get("marca"), "cobertura": cobertura, "motivo": motivo})
        filas_proveedores.append({
            "licitacion": codigo, "proveedor": resultado.get("proveedor"), "clase": resumen["clase"],
            "documentos": resumen["n_documentos"], "con_precios": "; ".join(resumen["con_precios"]),
            "sin_texto": "; ".join(resumen["sin_texto"]), "productos": len(productos)})
    resumen_tabla = [{"Estado": e, "Cobertura": c, "Productos": n} for (e, c), n in sorted(conteo.items())]
    escribir_excel(filas_productos, filas_proveedores, resumen_tabla, args.excel)

    print("\nProductos por estado y cobertura:")
    for fila in resumen_tabla:
        print(f"  {fila['Estado']:<13} {fila['Cobertura']:<22} {fila['Productos']}")
    clases = Counter(f["clase"] for f in filas_proveedores)
    print("\nProveedores por clase documental:")
    for clase, n in clases.most_common():
        print(f"  {clase:<55} {n}")
    print(f"\nExcel: {args.excel}")


if __name__ == "__main__":
    main()
