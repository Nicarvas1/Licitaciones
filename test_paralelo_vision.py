"""
Pruebas del modo vision, el paralelismo, el reproceso selectivo, la conservacion
de checkpoints, la cancelacion con Ctrl+C y el fallback tecnico del scraper.

No necesitan LM Studio: levantan un servidor falso compatible con
/v1/chat/completions en un puerto libre.

Ejecutar desde la raiz del proyecto:
    python -m unittest tests.test_paralelo_vision -v
"""

import contextlib
import importlib.util
import io
import json
import os
import signal
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest import mock

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))


def cargar_modulo(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, RAIZ / archivo)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


ex = cargar_modulo("extractor_ia", "3_extraer_ia.py")
import pymupdf as fitz  # noqa: E402  (despues de cargar el extractor, que valida dependencias)


# ---------------------------------------------------------------------------
# Servidor falso de LM Studio
# ---------------------------------------------------------------------------
PRODUCTO = {
    "item": "1", "producto": "Notebook HP ProBook 440 G11", "marca": "HP", "modelo": "ProBook 440 G11",
    "cantidad": 10, "cantidad_fuente": "explicita", "precio_unitario": 650000, "precio_total": 6500000,
    "precio_total_tipo": "linea", "moneda": "CLP", "categoria": "equipo", "pagina": 1,
    "fila_fuente": "FILA 2", "evidencia": "Notebook HP ProBook 440 G11 10 650.000", "confianza": "alta",
}


class ServidorFalso:
    def __init__(self, retardo=0.0):
        self.retardo = retardo
        self.lock = threading.Lock()
        self.llamadas = 0
        self.con_imagen = 0
        self.activas = 0
        self.max_activas = 0
        servidor = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                cuerpo = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                contenido = cuerpo["messages"][1]["content"]
                tiene_imagen = isinstance(contenido, list) and any(
                    parte.get("type") == "image_url" for parte in contenido
                )
                with servidor.lock:
                    servidor.llamadas += 1
                    servidor.con_imagen += int(tiene_imagen)
                    servidor.activas += 1
                    servidor.max_activas = max(servidor.max_activas, servidor.activas)
                time.sleep(servidor.retardo)
                with servidor.lock:
                    servidor.activas -= 1
                respuesta = {
                    "model": cuerpo["model"],
                    "choices": [{"message": {"content": json.dumps({"productos": [PRODUCTO], "observaciones": None})}}],
                    "usage": {"prompt_tokens": 100, "completion_tokens": 10, "total_tokens": 110},
                }
                datos = json.dumps(respuesta).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(datos)))
                self.end_headers()
                self.wfile.write(datos)

        self.http = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.http.daemon_threads = True
        self.url = f"http://127.0.0.1:{self.http.server_address[1]}/v1/chat/completions"
        self.hilo = threading.Thread(target=self.http.serve_forever, daemon=True)
        self.hilo.start()

    def cerrar(self):
        self.http.shutdown()
        self.http.server_close()


# ---------------------------------------------------------------------------
# Documentos de prueba (solo PyMuPDF)
# ---------------------------------------------------------------------------
FILAS = [
    ["Item", "Descripcion", "Cant.", "Precio unitario", "Total"],
    ["1", "Notebook HP ProBook 440 G11", "10", "$ 650.000", "$ 6.500.000"],
    ["2", "Monitor HP P24 G5", "10", "$ 120.000", "$ 1.200.000"],
]
COLUMNAS_X = [50, 90, 300, 350, 450, 545]


def pdf_tabla(ruta, bordes):
    doc = fitz.open()
    pagina = doc.new_page(width=595, height=842)
    pagina.insert_text((50, 60), "Oferta economica - adquisicion de notebooks", fontsize=12)
    for numero, fila in enumerate(FILAS):
        y = 100 + numero * 25
        for columna, valor in enumerate(fila):
            pagina.insert_text((COLUMNAS_X[columna] + 3, y + 17), valor, fontsize=9)
    if bordes:
        for numero in range(len(FILAS) + 1):
            y = 100 + numero * 25
            pagina.draw_line((COLUMNAS_X[0], y), (COLUMNAS_X[-1], y))
        for x in COLUMNAS_X:
            pagina.draw_line((x, 100), (x, 100 + len(FILAS) * 25))
    doc.save(str(ruta))
    doc.close()


def pdf_escaneado(ruta):
    doc = fitz.open()
    pagina = doc.new_page(width=595, height=842)
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 200), 0)
    pixmap.clear_with(255)
    pagina.insert_image(pagina.rect, pixmap=pixmap)
    doc.save(str(ruta))
    doc.close()


def crear_proveedor(raiz, licitacion, numero, tipo_pdf="bordes", extra_docx=False):
    rut = f"7{numero:07d}-{numero % 10}"
    carpeta = raiz / licitacion / f"{rut}__PROVEEDOR_{numero}"
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / "oferta.json").write_text(json.dumps(
        {"proveedor": f"Proveedor {numero}", "rut": rut, "total": "$ 7.700.000", "estado": "Aceptada"}
    ), encoding="utf-8")
    ruta_pdf = carpeta / "economico__01__oferta.pdf"
    if tipo_pdf == "escaneado":
        pdf_escaneado(ruta_pdf)
    else:
        pdf_tabla(ruta_pdf, bordes=(tipo_pdf == "bordes"))
    if extra_docx:
        from docx import Document
        documento = Document()
        documento.add_paragraph("Ficha tecnica Notebook HP ProBook 440 G11")
        documento.save(str(carpeta / "tecnico__01__ficha.docx"))
    return rut, carpeta


class BaseConServidor(unittest.TestCase):
    retardo = 0.0

    def setUp(self):
        self.servidor = ServidorFalso(self.retardo)
        self.parche_url = mock.patch.object(ex, "LMSTUDIO_URL", self.servidor.url)
        self.parche_url.start()
        self.tmp = tempfile.TemporaryDirectory()
        self.raiz = Path(self.tmp.name)
        self.ofertas = self.raiz / "ofertas"
        ex.DETENER.clear()

    def tearDown(self):
        ex.DETENER.clear()
        self.parche_url.stop()
        self.servidor.cerrar()
        self.tmp.cleanup()

    def correr(self, *extra):
        argv = [
            "3_extraer_ia.py", "--dir", str(self.ofertas), "--backend", "lmstudio",
            "--modelo", "qwen/qwen3.6-35b-a3b", "--excel", str(self.raiz / "salida.xlsx"),
            "--metadata-csv", str(self.raiz / "no_existe.csv"),
            "--reintentos-modelo", "0", "--espera-reintento", "0", "--sin-calentamiento", "--rampa", "0", *extra,
        ]
        salida = io.StringIO()
        with mock.patch.object(sys, "argv", argv), contextlib.redirect_stdout(salida):
            ex.main()
        return salida.getvalue()

    def checkpoint(self, licitacion):
        return json.loads((self.ofertas / licitacion / "extraccion_ia.json").read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# 1. Payloads multimodales
# ---------------------------------------------------------------------------
class TestPayloads(unittest.TestCase):
    def _respuesta(self, backend):
        respuesta = mock.Mock(status_code=200)
        if backend == "lmstudio":
            respuesta.json.return_value = {"choices": [{"message": {"content": '{"productos": []}'}}], "usage": {}}
        else:
            respuesta.json.return_value = {"response": '{"productos": []}'}
        return respuesta

    def test_lmstudio_con_imagenes(self):
        imagen = {"mime": "image/png", "base64": "QUJD"}
        with mock.patch.object(ex.requests, "post", return_value=self._respuesta("lmstudio")) as post:
            ex.consultar_modelo("PROMPT", "m", 10, 8192, "lmstudio", imagenes=[imagen, imagen])
        contenido = post.call_args.kwargs["json"]["messages"][1]["content"]
        self.assertEqual([parte["type"] for parte in contenido], ["image_url", "image_url", "text"])
        self.assertEqual(contenido[0]["image_url"]["url"], "data:image/png;base64,QUJD")
        self.assertEqual(contenido[-1]["text"], "PROMPT")

    def test_lmstudio_sin_imagenes_no_cambia(self):
        with mock.patch.object(ex.requests, "post", return_value=self._respuesta("lmstudio")) as post:
            ex.consultar_modelo("PROMPT", "m", 10, 8192, "lmstudio")
        self.assertEqual(post.call_args.kwargs["json"]["messages"][1]["content"], "PROMPT")

    def test_ollama_con_imagenes(self):
        with mock.patch.object(ex.requests, "post", return_value=self._respuesta("ollama")) as post:
            ex.consultar_modelo("PROMPT", "m", 10, 8192, "ollama", imagenes=[{"mime": "image/png", "base64": "QUJD"}])
        self.assertEqual(post.call_args.kwargs["json"]["images"], ["QUJD"])

    def test_detencion_impide_llamadas(self):
        ex.DETENER.set()
        try:
            with mock.patch.object(ex.requests, "post") as post:
                with self.assertRaises(ex.CorridaDetenida):
                    ex.consultar_modelo("PROMPT", "m", 10, 8192, "lmstudio")
            post.assert_not_called()
        finally:
            ex.DETENER.clear()


# ---------------------------------------------------------------------------
# 2. Regla de vision
# ---------------------------------------------------------------------------
class TestErroresTransitorios(unittest.TestCase):
    def test_channel_error_se_reintenta(self):
        falla = mock.Mock(status_code=400, text='{"error": "Error: Channel Error"}')
        exito = mock.Mock(status_code=200)
        exito.json.return_value = {"choices": [{"message": {"content": '{"productos": []}'}}], "usage": {}}
        with mock.patch.object(ex.requests, "post", side_effect=[falla, exito]) as post:
            _, estado, _, _ = ex.consultar_modelo("P", "m", 10, 8192, "lmstudio", reintentos=2, espera_reintento=0)
        self.assertEqual(post.call_count, 2)
        self.assertIn(estado, ex.ESTADOS_IA_OK)

    def test_error_de_contenido_no_se_reintenta(self):
        falla = mock.Mock(status_code=400, text='{"error": "invalid request"}')
        with mock.patch.object(ex.requests, "post", return_value=falla) as post:
            _, estado, _, _ = ex.consultar_modelo("P", "m", 10, 8192, "lmstudio", reintentos=2, espera_reintento=0)
        self.assertEqual(post.call_count, 1)
        self.assertTrue(estado.startswith("http_400"))


class TestArranque(BaseConServidor):
    retardo = 0.2

    def test_calentamiento_y_rampa(self):
        for numero in range(1, 4):
            crear_proveedor(self.ofertas, "LIC-1", numero)
        argv = ["3_extraer_ia.py", "--dir", str(self.ofertas), "--backend", "lmstudio", "--modelo", "m",
                "--excel", str(self.raiz / "x.xlsx"), "--metadata-csv", str(self.raiz / "no.csv"),
                "--paralelo", "3", "--rampa", "1", "--por-proveedor"]
        salida = io.StringIO()
        inicio = time.time()
        with mock.patch.object(sys, "argv", argv), contextlib.redirect_stdout(salida):
            ex.main()
        self.assertIn("Modelo listo", salida.getvalue())
        self.assertEqual(self.servidor.llamadas, 4)  # 1 de calentamiento + 3 proveedores
        self.assertGreaterEqual(time.time() - inicio, 2)  # el tercero espero 2 s de rampa


class TestReglaVision(unittest.TestCase):
    FILAS_NOTEBOOK = (
        "Item Descripcion Cant. Precio unitario Total\n"
        "1 Notebook HP ProBook 440 G11 10 $ 650.000 $ 6.500.000\n"
        "2 Monitor HP P24 G5 10 $ 120.000 $ 1.200.000\n"
    )

    def test_tabla_sin_bordes_de_equipos(self):
        self.assertEqual(ex.pagina_requiere_vision(self.FILAS_NOTEBOOK, [], 0), "tabla_sin_bordes")

    def test_formulario_administrativo_con_montos_sueltos(self):
        texto = ("Formulario de garantia para adquisicion de equipos computacionales\n"
                 "Monto garantia: $ 1.000.000\nTotal oferta: $ 9.282.000\n")
        self.assertIsNone(ex.pagina_requiere_vision(texto, [], 0))

    def test_tabla_de_soportes_tv_fuera_de_alcance(self):
        texto = ("Soporte TV precio unitario total\n"
                 "1 Soporte TV 55 pulgadas 2 $ 35.000 $ 70.000\n"
                 "2 Soporte TV 65 pulgadas 1 $ 45.000 $ 45.000\n")
        self.assertIsNone(ex.pagina_requiere_vision(texto, [], 0))

    def test_min_filas_configurable(self):
        una_fila = "Precio unitario\n1 Notebook HP ProBook 440 G11 10 $ 650.000 $ 6.500.000\n"
        self.assertIsNone(ex.pagina_requiere_vision(una_fila, [], 0, min_filas=2))
        self.assertEqual(ex.pagina_requiere_vision(una_fila, [], 0, min_filas=1), "tabla_sin_bordes")

    def test_pagina_escaneada(self):
        self.assertEqual(ex.pagina_requiere_vision("", [], 0), "sin_texto")

    def test_tabla_colapsada_solo_con_equipos_y_precios(self):
        con_equipos = "Anexo economico: Notebook HP ProBook 440 G11, precio unitario y total por item"
        sin_equipos = "Declaracion jurada simple de socios y accionistas de la empresa oferente"
        self.assertEqual(ex.pagina_requiere_vision(con_equipos, [], 1), "tabla_colapsada")
        self.assertIsNone(ex.pagina_requiere_vision(sin_equipos, [], 1))

    def test_pdf_real_con_y_sin_bordes(self):
        with tempfile.TemporaryDirectory() as tmp:
            con_bordes, sin_bordes = Path(tmp) / "con.pdf", Path(tmp) / "sin.pdf"
            pdf_tabla(con_bordes, bordes=True)
            pdf_tabla(sin_bordes, bordes=False)
            detalle_con, detalle_sin = {}, {}
            ex.extraer_pdf(con_bordes, 20, detalle=detalle_con)
            ex.extraer_pdf(sin_bordes, 20, detalle=detalle_sin)
        self.assertEqual(detalle_con["paginas_vision"], {})
        self.assertEqual(detalle_sin["paginas_vision"], {0: "tabla_sin_bordes"})


# ---------------------------------------------------------------------------
# 3. Paralelismo, reproceso selectivo y checkpoints
# ---------------------------------------------------------------------------
class TestParaleloYReproceso(BaseConServidor):
    retardo = 0.2

    def test_paralelo_usa_varias_peticiones_a_la_vez(self):
        for numero in range(1, 5):
            crear_proveedor(self.ofertas, "LIC-1", numero)
        self.correr("--paralelo", "4")
        self.assertGreaterEqual(self.servidor.max_activas, 2)
        self.assertEqual(len(self.checkpoint("LIC-1")), 4)

    def test_vision_envia_imagenes_solo_donde_corresponde(self):
        crear_proveedor(self.ofertas, "LIC-1", 1, "bordes")
        crear_proveedor(self.ofertas, "LIC-1", 2, "sin_bordes")
        crear_proveedor(self.ofertas, "LIC-1", 3, "escaneado")
        self.correr("--paralelo", "3", "--vision")
        self.assertEqual(self.servidor.con_imagen, 2)
        registros = {r["proveedor"]: r["resultados_archivos"][0] for r in self.checkpoint("LIC-1")}
        self.assertNotIn("paginas_vision", registros["Proveedor 1"])
        self.assertEqual(registros["Proveedor 2"]["paginas_vision"], {"1": "tabla_sin_bordes"})
        self.assertEqual(registros["Proveedor 3"]["paginas_vision"], {"1": "sin_texto"})

    def test_reproceso_parcial_solo_archivo_fallido(self):
        crear_proveedor(self.ofertas, "LIC-1", 1, "bordes", extra_docx=True)
        self.correr()
        resultados = self.checkpoint("LIC-1")
        for registro in resultados[0]["resultados_archivos"]:
            if registro["archivo"].endswith(".docx"):
                registro["estado_ia"] = "error_lmstudio: simulado"
        (self.ofertas / "LIC-1" / "extraccion_ia.json").write_text(json.dumps(resultados), encoding="utf-8")

        antes = self.servidor.llamadas
        self.correr()
        self.assertEqual(self.servidor.llamadas - antes, 2)  # el docx + la consolidacion
        registros = {r["archivo"]: r for r in self.checkpoint("LIC-1")[0]["resultados_archivos"]}
        self.assertTrue(registros["economico__01__oferta.pdf"].get("reutilizado"))
        self.assertEqual(registros["tecnico__01__ficha.docx"]["estado_ia"], "ok")

    def test_archivo_nuevo_se_procesa_sin_repetir_el_resto(self):
        _, carpeta = crear_proveedor(self.ofertas, "LIC-1", 1, "bordes")
        self.correr()
        pdf_tabla(carpeta / "tecnico__02__nuevo.pdf", bordes=True)
        antes = self.servidor.llamadas
        self.correr()
        self.assertEqual(self.servidor.llamadas - antes, 2)
        archivos = [r["archivo"] for r in self.checkpoint("LIC-1")[0]["resultados_archivos"]]
        self.assertIn("tecnico__02__nuevo.pdf", archivos)

    def test_segunda_corrida_no_llama_al_modelo(self):
        crear_proveedor(self.ofertas, "LIC-1", 1)
        self.correr("--vision")
        antes = self.servidor.llamadas
        self.correr("--vision")
        self.assertEqual(self.servidor.llamadas, antes)

    def _preparar_dos_por_revisar(self):
        rut1, _ = crear_proveedor(self.ofertas, "LIC-1", 1)
        rut2, _ = crear_proveedor(self.ofertas, "LIC-1", 2)
        self.correr()
        resultados = self.checkpoint("LIC-1")
        for resultado in resultados:
            resultado["resultados_archivos"][0]["estado_ia"] = "error_lmstudio: simulado"
        (self.ofertas / "LIC-1" / "extraccion_ia.json").write_text(json.dumps(resultados), encoding="utf-8")
        return rut1, rut2

    def _correr_con_fallo(self, rut_que_falla, excepcion):
        original = ex.procesar_oferta

        def procesar_con_fallo(carpeta, *args, **kwargs):
            if carpeta.name.startswith(rut_que_falla):
                raise excepcion
            time.sleep(0.3)  # el otro proveedor termina despues y reescribe el checkpoint
            return original(carpeta, *args, **kwargs)

        with mock.patch.object(ex, "procesar_oferta", procesar_con_fallo):
            return self.correr("--paralelo", "2")

    def test_fallo_de_reproceso_no_borra_checkpoint_previo(self):
        rut1, rut2 = self._preparar_dos_por_revisar()
        self._correr_con_fallo(rut1, ValueError("fallo simulado"))
        por_rut = {r["rut"]: r for r in self.checkpoint("LIC-1")}
        self.assertEqual(set(por_rut), {rut1, rut2})
        self.assertEqual(por_rut[rut1]["resultados_archivos"][0]["estado_ia"], "error_lmstudio: simulado")
        self.assertEqual(por_rut[rut2]["resultados_archivos"][0]["estado_ia"], "ok")

    def test_backend_caido_no_borra_checkpoint_previo(self):
        rut1, rut2 = self._preparar_dos_por_revisar()
        self._correr_con_fallo(rut1, RuntimeError("lmstudio dejo de responder"))
        self.assertEqual({r["rut"] for r in self.checkpoint("LIC-1")}, {rut1, rut2})

    def test_resultado_previo_sigue_en_excel_si_falla(self):
        rut1, _ = self._preparar_dos_por_revisar()
        self._correr_con_fallo(rut1, ValueError("fallo simulado"))
        from openpyxl import load_workbook
        libro = load_workbook(self.raiz / "salida.xlsx", read_only=True)
        filas = list(libro["Resumen"].iter_rows(values_only=True))
        libro.close()
        encabezado = filas[0]
        ruts = {dict(zip(encabezado, fila)).get("rut") for fila in filas[1:]}
        self.assertIn(rut1, ruts)


# ---------------------------------------------------------------------------
# 3b. Modo por proveedor (--por-proveedor)
# ---------------------------------------------------------------------------
def pdf_administrativo(ruta):
    doc = fitz.open()
    pagina = doc.new_page()
    for i, linea in enumerate(["DECLARACION JURADA SIMPLE", "Adquisicion de computadores para la Municipalidad",
                               "El oferente declara no tener conflictos de interes."]):
        pagina.insert_text((40, 60 + i * 16), linea, fontsize=10)
    doc.save(str(ruta))
    doc.close()


class TestPorProveedor(BaseConServidor):
    def test_una_llamada_por_proveedor_y_descarta_administrativos(self):
        for numero in range(1, 4):
            _, carpeta = crear_proveedor(self.ofertas, "LIC-1", numero, extra_docx=True)
            pdf_administrativo(carpeta / "administrativo__01__declaracion.pdf")
        self.correr("--por-proveedor")
        self.assertEqual(self.servidor.llamadas, 3)  # antes: 3 archivos por proveedor + consolidacion
        for resultado in self.checkpoint("LIC-1"):
            self.assertEqual(resultado["modo"], "por_proveedor")
            self.assertEqual(resultado["llamadas_modelo"], 1)
            registros = {r["archivo"]: r for r in resultado["resultados_archivos"]}
            self.assertEqual(registros["administrativo__01__declaracion.pdf"]["paginas_enviadas"], [])
            self.assertEqual(registros["economico__01__oferta.pdf"]["paginas_enviadas"], [1])
            self.assertEqual(len(resultado["productos"]), 1)

    def test_si_el_filtro_descarta_todo_igual_se_envia(self):
        """Regresion (prueba de 5 licitaciones): un proveedor sin paginas 'relevantes'
        quedaba sin ninguna llamada y se perdian sus productos."""
        rut, carpeta = crear_proveedor(self.ofertas, "LIC-1", 1)
        (carpeta / "economico__01__oferta.pdf").unlink()
        doc = fitz.open()
        pagina = doc.new_page()
        for i, linea in enumerate(["Cotizacion", "Equipo: computador tipo 1, diez unidades",
                                   "Valor por unidad seiscientos cincuenta mil pesos"]):
            pagina.insert_text((40, 60 + i * 16), linea, fontsize=10)
        doc.save(str(carpeta / "economico__01__carta.pdf"))
        doc.close()
        self.correr("--por-proveedor")
        resultado = self.checkpoint("LIC-1")[0]
        self.assertEqual(resultado["llamadas_modelo"], 1)
        registro = resultado["resultados_archivos"][0]
        self.assertTrue(registro.get("respaldo_sin_filtro"))
        self.assertEqual(registro["paginas_enviadas"], [1])

    def test_si_no_cabe_divide_y_consolida(self):
        crear_proveedor(self.ofertas, "LIC-1", 1, extra_docx=True)
        self.correr("--por-proveedor", "--max-chars-proveedor", "300")
        resultado = self.checkpoint("LIC-1")[0]
        self.assertGreaterEqual(resultado["llamadas_modelo"], 3)  # 2 llamadas de extraccion + consolidacion
        self.assertEqual(resultado["estado_consolidacion"], "ok")

    def test_escaneado_va_como_imagen_con_vision(self):
        crear_proveedor(self.ofertas, "LIC-1", 1, "escaneado")
        self.correr("--por-proveedor", "--vision")
        self.assertEqual((self.servidor.llamadas, self.servidor.con_imagen), (1, 1))

    def test_imagenes_van_en_jpeg(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "a.pdf"
            pdf_escaneado(ruta)
            imagen = ex.renderizar_paginas_pdf(ruta, [0], 120)[0]
        self.assertEqual(imagen["mime"], "image/jpeg")

    def test_vision_solo_escaneadas(self):
        """Solo la pagina escaneada va como imagen (no la tabla sin bordes), y al activarla
        sobre un lote ya procesado solo se reprocesa el proveedor con OCR pendiente."""
        crear_proveedor(self.ofertas, "LIC-1", 1, "sin_bordes")
        rut_escaneado, _ = crear_proveedor(self.ofertas, "LIC-1", 2, "escaneado")
        self.correr("--por-proveedor")
        pendientes = {r["rut"]: r["necesita_ocr"] for r in self.checkpoint("LIC-1")}
        self.assertTrue(pendientes[rut_escaneado])
        antes = self.servidor.llamadas
        self.correr("--por-proveedor", "--vision", "--vision-solo-escaneadas")
        self.assertEqual(self.servidor.llamadas - antes, 1)  # solo el proveedor escaneado
        self.assertEqual(self.servidor.con_imagen, 1)
        pendientes = {r["rut"]: r["necesita_ocr"] for r in self.checkpoint("LIC-1")}
        self.assertEqual(pendientes[rut_escaneado], [])

    def test_cambiar_de_modo_rehace_el_proveedor(self):
        crear_proveedor(self.ofertas, "LIC-1", 1)
        self.correr()
        self.assertEqual(self.checkpoint("LIC-1")[0]["modo"], "por_archivo")
        self.correr("--por-proveedor")
        self.assertEqual(self.checkpoint("LIC-1")[0]["modo"], "por_proveedor")
        antes = self.servidor.llamadas
        self.correr("--por-proveedor")
        self.assertEqual(self.servidor.llamadas, antes)  # ya procesado en este modo: no repite

    def test_consumo_cuenta_una_llamada(self):
        crear_proveedor(self.ofertas, "LIC-1", 1, extra_docx=True)
        self.correr("--por-proveedor")
        from openpyxl import load_workbook
        libro = load_workbook(self.raiz / "salida.xlsx", read_only=True)
        filas = list(libro["Resumen"].iter_rows(values_only=True))
        libro.close()
        self.assertEqual(dict(zip(filas[0], filas[1]))["llamadas"], 1)


# ---------------------------------------------------------------------------
# 4. Ctrl+C (proceso real + SIGINT). En Windows se prueba a mano.
# ---------------------------------------------------------------------------
@unittest.skipUnless(os.name == "posix", "SIGINT a un subproceso solo se prueba en Linux/macOS")
class TestCtrlC(unittest.TestCase):
    def test_ctrl_c_detiene_rapido_y_deja_checkpoint_valido(self):
        servidor = ServidorFalso(retardo=1.0)
        try:
            with tempfile.TemporaryDirectory() as tmp:
                ofertas = Path(tmp) / "ofertas"
                for numero in range(1, 13):
                    crear_proveedor(ofertas, "LIC-1", numero)
                entorno = {**os.environ, "LMSTUDIO_URL": servidor.url, "PYTHONUNBUFFERED": "1"}
                proceso = subprocess.Popen(
                    [sys.executable, str(RAIZ / "3_extraer_ia.py"), "--dir", str(ofertas),
                     "--backend", "lmstudio", "--modelo", "m", "--paralelo", "2",
                     "--excel", str(Path(tmp) / "x.xlsx"), "--metadata-csv", str(Path(tmp) / "no.csv")],
                    cwd=str(RAIZ), env=entorno, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                )
                time.sleep(4)
                inicio = time.time()
                proceso.send_signal(signal.SIGINT)
                salida, _ = proceso.communicate(timeout=30)
                demora = time.time() - inicio

                # Sin la correccion, el executor seguia con los 12 proveedores (~12 s mas).
                self.assertEqual(proceso.returncode, 130, salida)
                self.assertLess(demora, 5, salida)
                self.assertLess(servidor.llamadas, 24)  # 12 proveedores x 2 llamadas
                checkpoint = ofertas / "LIC-1" / "extraccion_ia.json"
                if checkpoint.exists():
                    json.loads(checkpoint.read_text(encoding="utf-8"))  # debe ser JSON valido
        finally:
            servidor.cerrar()


# ---------------------------------------------------------------------------
# 5. Fallback tecnico del scraper
# ---------------------------------------------------------------------------
try:
    scraper = cargar_modulo("scraper_ofertas", "2_scraper_ofertas.py")
except SystemExit:
    scraper = None


@unittest.skipIf(scraper is None, "faltan playwright/bs4 para importar el scraper")
class TestFallbackTecnico(unittest.TestCase):
    def test_combina_resultados_y_no_omite_pendientes(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = Path(tmp)
            for rut in ("1-1", "2-2", "3-3"):
                (carpeta / f"{rut}__P").mkdir()
                (carpeta / f"{rut}__P" / "oferta.json").write_text(json.dumps({"rut": rut}), encoding="utf-8")
            con_precio = {"rut": "1-1", "productos": [{"precio_unitario": 1000}]}
            sin_precio = {"rut": "2-2", "productos": [{"precio_unitario": None}]}
            (carpeta / "extraccion_ia.json").write_text(json.dumps([sin_precio]), encoding="utf-8")
            (carpeta / "extraccion_openai.json").write_text(json.dumps([con_precio]), encoding="utf-8")

            excluir = scraper.ruts_con_precio([carpeta / "extraccion_ia.json", carpeta / "extraccion_openai.json"])
            self.assertEqual(excluir, {"1-1"})
            # 2-2 (sin precio) y 3-3 (aun no procesado) quedan para el fallback.
            self.assertEqual(scraper.ruts_en_disco(carpeta) - excluir, {"2-2", "3-3"})


if __name__ == "__main__":
    unittest.main()
