"""
Pruebas del extractor por reglas (3_extraer_reglas.py).

Ejecutar desde la raiz del proyecto:
    python -m unittest tests.test_extraer_reglas -v
"""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

spec = importlib.util.spec_from_file_location("extraer_reglas", RAIZ / "3_extraer_reglas.py")
reglas = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reglas)

import pymupdf as fitz  # noqa: E402
from openpyxl import Workbook  # noqa: E402

ARGS = SimpleNamespace(max_paginas=20)


class TestCatalogo(unittest.TestCase):
    def test_modelos_por_familia(self):
        casos = {
            "Notebook HP ProBook 440 G11 Core i5 16GB": ("HP", "ProBook 440 G11", "equipo"),
            "Notebook Lenovo ThinkPad E14 Gen 5 i5": ("Lenovo", "ThinkPad E14 Gen5", "equipo"),
            "Notebook Dell Latitude 3540 i5 16GB": ("Dell", "Latitude 3540", "equipo"),
            "Impresora HP LaserJet Pro 4003dn": ("HP", "LaserJet Pro 4003dn", "impresora"),
            "Monitor Dell P2422H 24\"": ("Dell", "P2422H", "monitor"),
            "Monitor HP P24 G5": ("HP", "P24 G5", "monitor"),
        }
        for texto, esperado in casos.items():
            with self.subTest(texto=texto):
                self.assertEqual(reglas.detectar_modelo(texto), esperado)

    def test_modelo_despues_de_marca_sin_familia(self):
        self.assertEqual(reglas.modelo_tras_marca("Monitor LG 24MK430H 24 pulgadas"), "24MK430H")
        self.assertIsNone(reglas.modelo_tras_marca("Monitor LG de 24 pulgadas"))


class TestMontos(unittest.TestCase):
    def test_formatos_de_monto(self):
        casos = {"$ 650.000 c/u": 650000, "650.000 + IVA": 650000, "650.000.-": 650000,
                 "$650.000 neto": 650000, "650.000,00": 650000, "10 unidades": 10,
                 "Core i5 16GB": None, "2 x 650.000": None}
        for texto, esperado in casos.items():
            with self.subTest(texto=texto):
                self.assertEqual(reglas.numero(texto), esperado)


class TestTipoEquipo(unittest.TestCase):
    def test_subcategorias(self):
        casos = {"NOTEBOOK": "notebook", "Computador HP AIO 27- cr0011la": "all-in-one",
                 "HP Pavilion 32-B1002LA": "all-in-one", "WORKSTATION": "workstation",
                 "PC Armado Ci9 14": "desktop", "Monitor HP P24": "monitor",
                 "Multifuncional Brother DCP-1617NW": "impresora", "MacBook Pro de 14 pulgadas": "notebook"}
        for texto, esperado in casos.items():
            with self.subTest(texto=texto):
                self.assertEqual(reglas.ia.subcategoria_producto({"producto": texto}), esperado)


class TestEncabezados(unittest.TestCase):
    def test_prefiere_neto_y_descripcion_ofertada(self):
        filas = [
            ["ANEXO 5"], [],
            ["N°", "Descripción solicitada", "Descripción ofertada", "Cantidad",
             "Precio Unitario Neto", "Precio Unitario con IVA", "Total Neto"],
        ]
        inicio, mapeo = reglas.buscar_encabezado(filas)
        self.assertEqual(inicio, 3)
        self.assertEqual(mapeo["descripcion"], 2)
        self.assertEqual(mapeo["unitario"], 4)
        self.assertEqual(mapeo["total"], 6)

    def test_encabezado_en_dos_filas(self):
        filas = [["Item", "Producto", "Cant.", "Precio", "Valor"], ["", "", "", "Unitario", "Total"],
                 [1, "Notebook", 10, 700000, 7000000]]
        inicio, mapeo = reglas.buscar_encabezado(filas)
        self.assertEqual(inicio, 2)
        self.assertEqual((mapeo["unitario"], mapeo["total"], mapeo["cantidad"]), (3, 4, 2))

    def test_tabla_tecnica_sin_precios_no_tiene_encabezado_valido(self):
        self.assertIsNone(reglas.buscar_encabezado([["Característica", "Requerido", "Ofertado"]]))


class TestFilasCuadradas(unittest.TestCase):
    def test_detecta_cantidad_unitario_total(self):
        resultado = reglas.fila_cuadrada("1  Notebook Acer TravelMate P2 14 i5   10   $ 600.000   $ 6.000.000")
        self.assertEqual(resultado, ("Notebook Acer TravelMate P2 14 i5", 10, 600000, 6000000))

    def test_no_confunde_resolucion_ni_modelo(self):
        resultado = reglas.fila_cuadrada("Monitor LG 24MK430H 1920 x 1080   2   $ 120.000   $ 240.000")
        self.assertEqual(resultado[1:], (2, 120000, 240000))

    def test_no_toma_neto_iva_total(self):
        self.assertIsNone(reglas.fila_cuadrada("Neto $ 6.000.000  IVA $ 1.140.000  Total $ 7.140.000"))

    def test_no_inventa_sin_cuadratura(self):
        self.assertIsNone(reglas.fila_cuadrada("10 notebooks por un total de $ 6.500.000"))


class TestProveedor(unittest.TestCase):
    def _proveedor(self, raiz):
        carpeta = raiz / "LIC" / "76000001-1__PROV"
        carpeta.mkdir(parents=True)
        (carpeta / "oferta.json").write_text(json.dumps({"proveedor": "Prov", "rut": "76000001-1", "total": "$ 9.000.000"}))
        return carpeta

    def test_excel_con_accesorios_y_garantia(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            libro = Workbook()
            hoja = libro.active
            hoja.append(["Descripción", "Cantidad", "Precio unitario", "Total"])
            hoja.append(["Notebook Asus ExpertBook B1 B1402 i5", 10, 620000, 6200000])
            hoja.append(["Mouse inalámbrico", 10, 8000, 80000])
            hoja.append(["Garantía extendida 3 años", 10, 50000, 500000])
            hoja.append(["Total", None, None, 6780000])
            libro.save(carpeta / "economico__01__oferta.xlsx")
            resultado = reglas.procesar_proveedor(carpeta, ARGS, None)
        self.assertEqual(resultado["estado_reglas"], "resuelto")
        self.assertEqual([(p["marca"], p["precio_unitario"]) for p in resultado["productos"]], [("Asus", 620000)])

    def test_marca_desde_ficha_tecnica(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            libro = Workbook()
            hoja = libro.active
            hoja.append(["Descripción", "Cantidad", "Valor unitario neto", "Valor total neto"])
            hoja.append(["Computador de escritorio tipo 1", 10, 800000, 8000000])
            libro.save(carpeta / "economico__01__anexo.xlsx")
            doc = fitz.open()
            pagina = doc.new_page()
            pagina.insert_text((40, 60), "Equipo ofertado: HP ProDesk 400 G9 SFF", fontsize=10)
            doc.save(str(carpeta / "tecnico__01__ficha.pdf"))
            doc.close()
            resultado = reglas.procesar_proveedor(carpeta, ARGS, None)
        producto = resultado["productos"][0]
        self.assertEqual((producto["marca"], producto["modelo"]), ("HP", "ProDesk 400 G9 SFF"))
        self.assertEqual(producto["fuente_producto"], "tecnico__01__ficha.pdf")

    def test_no_presta_modelo_a_filas_que_no_son_equipos(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            libro = Workbook()
            hoja = libro.active
            hoja.append(["Descripción", "Cantidad", "Precio unitario", "Total"])
            hoja.append(["Computador de escritorio tipo 1", 10, 800000, 8000000])
            hoja.append(["Disco Duro SSD 480GB", 10, 43548, 435480])
            hoja.append(["UPS 1000VA", 2, 142950, 285900])
            hoja.append(["Despacho", 1, 290723, 290723])
            libro.save(carpeta / "economico__01__anexo.xlsx")
            doc = fitz.open()
            doc.new_page().insert_text((40, 60), "Equipo ofertado: HP ProDesk 400 G9 SFF", fontsize=10)
            doc.save(str(carpeta / "tecnico__01__ficha.pdf"))
            doc.close()
            resultado = reglas.procesar_proveedor(carpeta, ARGS, None)
        self.assertEqual([p["producto"] for p in resultado["productos"]], ["Computador de escritorio tipo 1"])
        self.assertEqual(resultado["productos"][0]["modelo"], "ProDesk 400 G9 SFF")

    def test_total_con_iva_se_convierte_a_neto(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            libro = Workbook()
            hoja = libro.active
            hoja.append(["Descripción", "Cantidad", "Precio unitario neto", "Total"])
            hoja.append(["Notebook HP ProBook 440 G11", 10, 650000, 7735000])
            libro.save(carpeta / "economico__01__oferta.xlsx")
            producto = reglas.procesar_proveedor(carpeta, ARGS, None)["productos"][0]
        self.assertEqual(producto["precio_total"], 6500000)
        self.assertEqual(producto["nota_iva"], "total_con_iva_convertido_a_neto")
        self.assertEqual(producto["estado_validacion"], "ok")

    def test_descripcion_en_varias_lineas(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            doc = fitz.open()
            pagina = doc.new_page()
            for i, linea in enumerate(["Item Descripcion Cant. Unitario Total",
                                       "Notebook HP 250 G10 Intel Core i5 16GB",
                                       "SSD 512GB Windows 11 Pro",
                                       "Licencia - OEM - DVD-ROM - PC  30  $ 854.597  $ 25.637.910"]):
                pagina.insert_text((40, 60 + i * 14), linea, fontsize=9)
            doc.save(str(carpeta / "economico__01__cotizacion.pdf"))
            doc.close()
            producto = reglas.procesar_proveedor(carpeta, ARGS, None)["productos"][0]
        self.assertTrue(producto["producto"].startswith("Notebook HP 250 G10"))
        self.assertEqual((producto["marca"], producto["modelo"]), ("HP", "250 G10"))

    def test_solo_total_queda_parcial_sin_inventar_unitario(self):
        with tempfile.TemporaryDirectory() as tmp:
            carpeta = self._proveedor(Path(tmp))
            libro = Workbook()
            hoja = libro.active
            hoja.append(["Descripción", "Valor total neto"])
            hoja.append(["Notebooks HP 240 G9 (lote)", 6000000])
            libro.save(carpeta / "economico__01__oferta.xlsx")
            resultado = reglas.procesar_proveedor(carpeta, ARGS, None)
        self.assertEqual(resultado["estado_reglas"], "parcial")
        self.assertIsNone(resultado["productos"][0]["precio_unitario"])


if __name__ == "__main__":
    unittest.main()
