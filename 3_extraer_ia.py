#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extractor IA local para ofertas de Mercado Publico.

Flujo:
  1. Extrae el texto de CADA archivo por separado.
    2. Preserva tablas y divide documentos largos sin perder filas.
    3. Consulta Ollama o LM Studio; con --paralelo N procesa N proveedores a la vez.
       Con --vision, las paginas PDF con tablas dificiles o escaneadas se envian
       como imagen (+ su texto) a un modelo con vision.
    4. Consolida hallazgos por proveedor sin truncar registros.
    5. Valida alcance, evidencia y coherencia matematica.
    6. Guarda checkpoints, alertas y un log JSONL para reanudar corridas.

Formatos: PDF, DOCX, XLSX, XLSM, TXT, CSV, TSV, ZIP, RAR y 7Z.
Los PDF sin texto quedan registrados en necesita_ocr o se procesan con
Tesseract cuando se usa --ocr.

Instalacion:
    pip install -r requirements.txt

Ejemplos:
    python 3_extraer_ia.py --dir ofertas/3572-22-LE25 --backend lmstudio --modelo openai/gpt-oss-20b --rehacer --debug
    python 3_extraer_ia.py --dir ofertas/3572-22-LE25 --solo-texto --rehacer
"""

import argparse
import base64
import logging
import csv
import hashlib
import importlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import unicodedata
import zipfile
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

# pdfminer (usado por pdfplumber) avisa "Could not get FontBBox..." cuando una fuente del PDF
# no declara sus dimensiones; usa valores por defecto y sigue leyendo. Solo ensucia la consola.
logging.getLogger("pdfminer").setLevel(logging.ERROR)
from datetime import datetime
from pathlib import Path

try:
    import pymupdf as fitz
    import pdfplumber
    import requests
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError as exc:
    sys.exit(
        f"Falta dependencia: {exc}. Instala: "
        "pip install pymupdf pdfplumber requests python-docx openpyxl"
    )

try:
    pytesseract = importlib.import_module("pytesseract")
    Image = importlib.import_module("PIL.Image")
except ModuleNotFoundError:
    pytesseract = None
    Image = None

# Se pueden cambiar con variables de entorno (ej. LM Studio en otro puerto o en otro equipo).
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434/api/generate")
LMSTUDIO_URL = os.environ.get("LMSTUDIO_URL", "http://127.0.0.1:1234/v1/chat/completions")
DEBUG = False
ESTADOS_IA_OK = {"ok", "ok_reparado", "ok_extraido"}
ESTADOS_IA_NEUTROS = ("no_ejecutada", "solo_texto", "omitido_sin_equipos", "omitido_duplicado")
# PyMuPDF no es seguro entre hilos: todo uso de fitz pasa por este lock.
# Solo protege abrir/renderizar PDFs (rapido); las llamadas al modelo quedan fuera.
FITZ_LOCK = threading.Lock()
LOG_LOCK = threading.Lock()  # el log JSONL se escribe desde varios hilos
# Senal de detencion para los hilos de trabajo (Ctrl+C o backend caido). Los hilos
# la revisan antes de cada llamada al modelo y antes de cada archivo.
DETENER = threading.Event()
RESULTADOS_PROVEEDOR = {}  # {(codigo licitacion, rut): resultado} para la cobertura documental


class CorridaDetenida(Exception):
    """Un hilo abandono su proveedor porque la corrida se esta deteniendo."""


def verificar_detencion():
    if DETENER.is_set():
        raise CorridaDetenida("corrida detenida")
EXTENSIONES_IMAGEN = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
PATRON_PALABRA_PRECIO = re.compile(r"precio|valor|unitario|total|neto|subtotal|monto", re.I)
PATRON_NUMERO = re.compile(r"\d+(?:[.,]\d+)*")
PATRON_EQUIPO_GENERICO = re.compile(r"\bequipos?\b", re.I)
PATRON_MONTO = re.compile(r"\$\s*[\d.]|\b\d{1,3}(?:\.\d{3})+(?:,\d+)?\b")

MARCAS = [
    "HP", "Lenovo", "Dell", "Asus", "Acer", "Apple", "Samsung", "Lexmark",
    "Epson", "Canon", "Brother", "MSI", "Huawei", "Kingston", "Logitech",
    "LG", "ThinkCentre", "ThinkPad", "ProOne", "ProDesk", "OptiPlex",
    "Latitude", "Pavilion", "Ozxen", "OMA", "SNAKE"
]

import cobertura_documental
from prompts_extraccion import (
        NOTA_IMAGENES_PROVEEDOR,
        PROMPT_ARCHIVO,
        PROMPT_ARCHIVO_COMPACTO,
        PROMPT_CONSOLIDAR,
        PROMPT_PROVEEDOR,
        REGLAS_VISION,
)

PATRON_PRODUCTO = re.compile(
    r"notebook|computador|laptop|all.?in.?one|\baio\b|desktop|monitor|impresora|"
    r"workstation|servidor|pc\b|equipo|modelo|procesador|ryzen|core\s*i[3579]|"
    r"\bram\b|ssd|hdd|pulgadas|\bgb\b", re.I
)
PATRON_PRECIO = re.compile(
    r"precio|valor|unitario|cantidad|subtotal|total|neto|iva|\$\s*[\d.]|"
    r"\b\d{1,3}(?:\.\d{3})+(?:,\d+)?\b", re.I
)


def limpiar_control(texto):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(texto))


def ruta_sistema(path):
    """Ruta para E/S de Windows sin MAX_PATH; nunca se guarda en fuentes o JSON."""
    ruta = os.fspath(path)
    if os.name != "nt":
        return ruta
    absoluta = os.path.abspath(ruta)
    if absoluta.startswith("\\\\?\\"):
        return absoluta
    if absoluta.startswith("\\\\"):
        return "\\\\?\\UNC\\" + absoluta.lstrip("\\")
    return "\\\\?\\" + absoluta


def cargar_json(path):
    try:
        with open(ruta_sistema(path), encoding="utf-8") as archivo:
            return json.load(archivo)
    except Exception:
        return {}


def escribir_json_atomico(path, datos):
    temporal = path.with_suffix(path.suffix + ".tmp")
    with open(ruta_sistema(temporal), "w", encoding="utf-8") as archivo:
        json.dump(datos, archivo, ensure_ascii=False, indent=2)
    os.replace(ruta_sistema(temporal), ruta_sistema(path))


def escribir_texto(path, contenido):
    with open(ruta_sistema(path), "w", encoding="utf-8") as archivo:
        archivo.write(contenido)


def cargar_metadata_licitaciones(ruta):
    ruta = Path(ruta)
    if not os.path.isfile(ruta_sistema(ruta)):
        return {}
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(ruta_sistema(ruta), encoding=encoding, newline="") as archivo:
                filas = list(csv.DictReader(archivo))
            break
        except UnicodeDecodeError:
            continue
    else:
        return {}
    return {
        (fila.get("codigo") or "").strip(): {
            "nombre_licitacion": (fila.get("nombre") or "").strip(),
            "fecha_publicacion": (fila.get("fecha_publicacion") or "").strip(),
            "estado_licitacion": (fila.get("estado") or "").strip(),
            "organismo": (fila.get("organismo") or "").strip(),
        }
        for fila in filas
        if (fila.get("codigo") or "").strip()
    }


def cargar_metadata_para_corrida(licitaciones, ruta_explicitada=None):
    """Combina catalogo global y CSV mensuales; un CSV indicado por el usuario prevalece."""
    catalogo = Path(__file__).resolve().parent / "lotes" / "catalogo_licitaciones.csv"
    fuentes = [catalogo]
    meses = sorted({licitacion.parent.parent for licitacion in licitaciones
                    if licitacion.parent.name.lower() == "ofertas"}, key=str)
    fuentes.extend(mes / "para_scrapear.csv" for mes in meses)
    if ruta_explicitada:
        fuentes.append(Path(ruta_explicitada))
    metadata = {}
    for fuente in dict.fromkeys(fuentes):
        metadata.update(cargar_metadata_licitaciones(fuente))
    return metadata


RUTAS_TESSERACT_WINDOWS = (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                           r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe")
TESSDATA_LOCAL = Path(__file__).resolve().parent / "tessdata"


def configurar_tesseract(tesseract_cmd=None):
    """Ubica Tesseract aunque no este en el PATH y usa la carpeta tessdata del proyecto si tiene idiomas (asi se
    agrega el espanol sin permisos de administrador: python verificar_entorno.py --preparar-tessdata)."""
    if pytesseract is None:
        return
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    elif pytesseract.pytesseract.tesseract_cmd == "tesseract" and not shutil.which("tesseract"):
        for ruta in RUTAS_TESSERACT_WINDOWS:
            if os.path.isfile(ruta):
                pytesseract.pytesseract.tesseract_cmd = ruta
                break
    if "TESSDATA_PREFIX" not in os.environ and any(TESSDATA_LOCAL.glob("*.traineddata")):
        os.environ["TESSDATA_PREFIX"] = str(TESSDATA_LOCAL)


def verificar_ocr(tesseract_cmd=None, idioma="spa+eng"):
    """(ok, mensaje). Con --ocr, un Tesseract ausente no da error: los PDF escaneados quedan "sin leer" y se
    pierden productos sin aviso. Esta comprobacion se hace al arrancar para que falle de forma visible."""
    if pytesseract is None or Image is None:
        return False, "falta la libreria pytesseract o Pillow (pip install -r requirements.txt)"
    configurar_tesseract(tesseract_cmd)
    try:
        version = str(pytesseract.get_tesseract_version())
        idiomas = set(pytesseract.get_languages(config=""))
    except Exception as exc:
        return False, (f"no se pudo ejecutar Tesseract ({str(exc)[:80]}). Instalalo (con el idioma espanol) o "
                       f"indica su ruta con --tesseract-cmd")
    faltan = [i for i in idioma.split("+") if i and i not in idiomas]
    if faltan:
        # Sin el paquete de idioma espanol el OCR sigue sirviendo (cifras y marcas), pero lee peor los acentos.
        return True, (f"Tesseract {version}; AVISO: falta el idioma {', '.join(faltan)} y se usara "
                      f"{idioma_ocr_efectivo(idioma)} (disponibles: {', '.join(sorted(idiomas))})")
    return True, f"Tesseract {version} ({', '.join(sorted(idiomas))})"


_IDIOMAS_OCR = {}


def idioma_ocr_efectivo(idioma):
    """Idioma pedido, reducido a los paquetes que Tesseract realmente tiene (se consulta una sola vez)."""
    if idioma not in _IDIOMAS_OCR:
        try:
            disponibles = set(pytesseract.get_languages(config=""))
            validos = [i for i in idioma.split("+") if i in disponibles]
            _IDIOMAS_OCR[idioma] = "+".join(validos) if validos else idioma
        except Exception:
            _IDIOMAS_OCR[idioma] = idioma
    return _IDIOMAS_OCR[idioma]


def extraer_pdf_ocr(
    path,
    max_paginas,
    tesseract_cmd=None,
    idioma="spa+eng",
    dpi=180,
    paginas=None,
    texto_paginas=None,
):
    if pytesseract is None or Image is None:
        return "", "ocr_dependencia_no_disponible"
    configurar_tesseract(tesseract_cmd)
    idioma = idioma_ocr_efectivo(idioma)
    try:
        # Renderizar bajo lock (PyMuPDF no es thread-safe); el OCR corre fuera del lock.
        renders = []
        escala = dpi / 72
        with FITZ_LOCK:
            doc = fitz.open(ruta_sistema(path))
            try:
                for index, page in enumerate(doc):
                    if index >= max_paginas:
                        break
                    if paginas is not None and index not in paginas:
                        continue
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(escala, escala), alpha=False)
                    renders.append((index, pixmap.tobytes("png")))
            finally:
                doc.close()
        partes = []
        for index, png in renders:
            with Image.open(io.BytesIO(png)) as imagen:
                try:
                    texto = pytesseract.image_to_string(imagen, lang=idioma)
                except pytesseract.TesseractError:
                    texto = pytesseract.image_to_string(imagen, lang="eng")
            if texto.strip():
                pagina_ocr = limpiar_control(f"[PAGINA {index + 1} - OCR]\n{texto.strip()}")
                partes.append(pagina_ocr)
                if texto_paginas is not None:
                    texto_paginas[index] = pagina_ocr
        texto = limpiar_control("\n\n".join(partes)).strip()
        return texto, "texto_ocr" if len(texto) > 40 else "necesita_ocr"
    except Exception as exc:
        return "", f"error_ocr: {exc}"


def paginas_pdf_blancas(path, indices):
    """Detecta solo paginas realmente blancas con una miniatura gris barata."""
    if not indices:
        return set()
    blancas = set()
    with FITZ_LOCK:
        doc = fitz.open(ruta_sistema(path))
        try:
            for indice in indices:
                if not 0 <= indice < len(doc):
                    continue
                pixmap = doc[indice].get_pixmap(matrix=fitz.Matrix(0.25, 0.25),
                                                colorspace=fitz.csGRAY, alpha=False)
                if all(pixel >= 250 for pixel in pixmap.samples):
                    blancas.add(indice)
        finally:
            doc.close()
    return blancas


def tabla_estructurada(filas):
    filas_con_datos = [fila for fila in filas if any(str(celda or "").strip() for celda in fila)]
    if not filas_con_datos:
        return False
    filas_multicolumna = sum(
        sum(bool(str(celda or "").strip()) for celda in fila) >= 2
        for fila in filas_con_datos
    )
    return filas_multicolumna / len(filas_con_datos) >= 0.5


def objeto_fuera_de_tablas(objeto, cajas):
    objeto_izquierda = objeto.get("x0", 0)
    objeto_derecha = objeto.get("x1", 0)
    objeto_arriba = objeto.get("top", 0)
    objeto_abajo = objeto.get("bottom", 0)
    return not any(
        objeto_izquierda < derecha
        and objeto_derecha > izquierda
        and objeto_arriba < abajo
        and objeto_abajo > arriba
        for izquierda, arriba, derecha, abajo in cajas
    )


def linea_tipo_fila(linea):
    """Linea que parece fila de tabla de precios: un monto y al menos otro numero
    (cantidad, item o segundo monto). Descarta lineas sueltas como 'Total $ 9.282.000'
    o 'Monto garantia $ 1.000.000'."""
    return bool(PATRON_MONTO.search(linea)) and len(PATRON_NUMERO.findall(linea)) >= 2


def pagina_en_alcance(texto_pagina):
    """La pagina menciona algun equipo del alcance comercial."""
    if PATRON_EQUIPO_RELEVANTE.search(texto_pagina) or PATRON_MONITOR.search(texto_pagina) \
            or PATRON_IMPRESORA.search(texto_pagina):
        return True
    # "equipo" generico cuenta solo si la pagina no habla de TV.
    return bool(PATRON_EQUIPO_GENERICO.search(texto_pagina)) and not PATRON_TV.search(texto_pagina)


def pagina_requiere_vision(texto_pagina, tablas_validas, tablas_colapsadas, min_filas=2):
    """Decide si una pagina conviene leerla como imagen en vez de solo texto.

    - sin_texto: pagina escaneada (sin imagen no hay nada que leer).
    - tabla_colapsada: pdfplumber detecto una tabla sin columnas recuperables, y la
      pagina menciona un equipo del alcance y precios.
    - tabla_sin_bordes: sin tablas detectadas, pero con al menos `min_filas` lineas
      que parecen filas de precios, una palabra de precio y un equipo del alcance.

    Las paginas que no califican NO se pierden: siguen por la ruta de texto.
    """
    if len(texto_pagina) <= 40 and not tablas_validas:
        return "sin_texto"
    if not pagina_en_alcance(texto_pagina):
        return None
    habla_de_precios = bool(PATRON_PALABRA_PRECIO.search(texto_pagina))
    if tablas_colapsadas:
        if habla_de_precios or PATRON_MONTO.search(texto_pagina):
            return "tabla_colapsada"
        return None
    if not tablas_validas and habla_de_precios:
        filas = sum(1 for linea in texto_pagina.splitlines() if linea_tipo_fila(linea))
        if filas >= min_filas:
            return "tabla_sin_bordes"
    return None


PALABRAS_CORTAS_VALIDAS = {"de", "el", "la", "en", "un", "y", "a", "o", "e", "al", "se", "su", "es", "no", "lo",
                           "por", "con", "los", "las", "del", "una", "que", "si", "ni", "ml", "kg", "gb", "tb", "hp",
                           "pc", "tv", "lg", "ia", "uf", "iva"}


def texto_ilegible(texto):
    """True si el texto extraido de una pagina es basura (letras sueltas de columnas rotadas o con
    fuentes mal codificadas): mas del 60% de sus 'palabras' tiene 1 o 2 letras. Esas paginas se
    tratan como escaneadas (OCR o vision) en vez de enviar ruido al modelo."""
    palabras = re.findall(r"[A-Za-zÁÉÍÓÚáéíóúñÑ]+", texto)
    if len(palabras) < 40:
        return False
    cortas = [w for w in palabras if len(w) <= 2 and w.lower() not in PALABRAS_CORTAS_VALIDAS]
    return len(cortas) / len(palabras) > 0.6


def extraer_pdf(path, max_paginas, usar_ocr=False, tesseract_cmd=None, idioma_ocr="spa+eng", dpi_ocr=180,
                detalle=None, min_filas_vision=2):
    """Extrae texto y tablas de un PDF.

    Si se entrega `detalle` (dict), se llena con informacion por pagina para el
    modo vision: texto_paginas {indice: texto} y paginas_vision {indice: motivo}.
    """
    try:
        partes = []
        paginas_sin_texto = []
        texto_paginas = {}
        paginas_vision = {}
        with pdfplumber.open(ruta_sistema(path)) as pdf:
            for index, page in enumerate(pdf.pages):
                if index >= max_paginas:
                    break
                try:
                    tablas = page.find_tables()
                except Exception:
                    tablas = []
                tablas_validas = []
                tablas_colapsadas = 0
                for tabla in tablas:
                    filas = tabla.extract() or []
                    if tabla_estructurada(filas):
                        tablas_validas.append((tabla, filas))
                    else:
                        tablas_colapsadas += 1

                if tablas_validas:
                    cajas = [tabla.bbox for tabla, _ in tablas_validas]
                    pagina_sin_tablas = page.filter(lambda objeto: objeto_fuera_de_tablas(objeto, cajas))
                    texto_pagina = (pagina_sin_tablas.extract_text() or "").strip()
                else:
                    texto_pagina = (page.extract_text(layout=True) or page.extract_text() or "").strip()

                if texto_pagina and texto_ilegible(texto_pagina):
                    texto_pagina = ""
                partes_pagina = [f"[PAGINA {index + 1}]"]
                if texto_pagina:
                    partes_pagina.append(texto_pagina)
                if tablas_colapsadas:
                    partes_pagina.append(f"[AVISO: {tablas_colapsadas} tabla(s) sin columnas recuperables]")
                for num_tabla, (_, filas) in enumerate(tablas_validas, 1):
                    partes_pagina.append(f"[TABLA {num_tabla} - pagina {index + 1}]")
                    for num_fila, fila in enumerate(filas, 1):
                        valores = [(celda or "").strip().replace("\n", " ") for celda in fila]
                        if any(valores):
                            partes_pagina.append(f"FILA {num_fila}: " + " || ".join(valores))
                partes.extend(partes_pagina)
                texto_paginas[index] = limpiar_control("\n".join(partes_pagina)).strip()
                motivo_vision = pagina_requiere_vision(
                    texto_pagina, tablas_validas, tablas_colapsadas, min_filas_vision
                )
                if motivo_vision:
                    paginas_vision[index] = motivo_vision
                if len(texto_pagina) <= 40 and not tablas_validas:
                    paginas_sin_texto.append(index)
        texto = limpiar_control("\n".join(partes)).strip()
        try:
            paginas_vacias = paginas_pdf_blancas(path, paginas_sin_texto)
        except Exception:
            # La deteccion de blancos es una optimizacion, no debe borrar texto util.
            paginas_vacias = set()
        paginas_sin_texto = [indice for indice in paginas_sin_texto if indice not in paginas_vacias]
        for indice in paginas_vacias:
            paginas_vision.pop(indice, None)
    except Exception:
        texto = ""
        paginas_sin_texto = []
        texto_paginas = {}
        paginas_vision = {}
        paginas_vacias = set()

    if detalle is not None:
        detalle["texto_paginas"] = texto_paginas
        detalle["paginas_vision"] = paginas_vision
        detalle["paginas_vacias"] = sorted(paginas_vacias)
        if paginas_vision:
            # En modo vision no hace falta OCR ni respaldo: las paginas dificiles
            # se leen como imagen y el resto ya tiene texto estructurado.
            return texto, "texto_vision"

    if usar_ocr and paginas_sin_texto:
        texto_ocr, estado_ocr = extraer_pdf_ocr(
            path,
            max_paginas,
            tesseract_cmd,
            idioma_ocr,
            dpi_ocr,
            set(paginas_sin_texto),
        )
        if texto_ocr:
            texto = limpiar_control(f"{texto}\n\n{texto_ocr}").strip()
            return texto, "texto_ocr"
        if len(texto) > 40:
            return texto, f"texto_con_paginas_ocr_pendiente:{estado_ocr}"
        return "", estado_ocr

    if len(texto) > 40:
        estado = "texto_con_paginas_ocr_pendiente" if paginas_sin_texto else "texto"
        return texto, estado

    # Respaldo: pdfplumber no extrajo nada usable (ej. PDF escaneado sin capa de texto normal).
    try:
        partes = []
        with FITZ_LOCK:
            doc = fitz.open(ruta_sistema(path))
            try:
                for index, page in enumerate(doc):
                    if index >= max_paginas:
                        break
                    partes.append(page.get_text("text"))
            finally:
                doc.close()
        texto = limpiar_control("\n".join(partes)).strip()
        if len(texto) > 40:
            return texto, "texto"
    except Exception as exc:
        if not usar_ocr:
            return "", f"error_pdf: {exc}"
    if usar_ocr:
        return extraer_pdf_ocr(path, max_paginas, tesseract_cmd, idioma_ocr, dpi_ocr)
    return "", "necesita_ocr"


def encontrar_antiword():
    """antiword (viene con Git for Windows) lee los .doc antiguos de Word."""
    ruta = shutil.which("antiword")
    if ruta:
        return ruta
    for candidato in (r"C:\Program Files\Git\mingw64\bin\antiword.exe", r"C:\Program Files\Git\usr\bin\antiword.exe"):
        if os.path.isfile(candidato):
            return candidato
    return None


def extraer_doc(path):
    """Documento Word antiguo (.doc): texto con antiword. Sin antiword queda como no soportado."""
    antiword = encontrar_antiword()
    if not antiword:
        return "", "no_soportado"
    temporal = None
    ruta = str(path)
    try:
        if len(ruta) > 240:  # antiword no entiende el prefijo de rutas largas: se trabaja sobre una copia
            import tempfile
            temporal = Path(tempfile.gettempdir()) / f"antiword_{os.getpid()}_{abs(hash(ruta))}.doc"
            shutil.copyfile(ruta_sistema(path), temporal)
            ruta = str(temporal)
        resultado = subprocess.run([antiword, ruta], capture_output=True, timeout=60)
    except (OSError, subprocess.SubprocessError):
        return "", "no_soportado"
    finally:
        if temporal is not None:
            temporal.unlink(missing_ok=True)
    if resultado.returncode != 0:
        return "", "no_soportado"
    texto = limpiar_control(resultado.stdout.decode("cp1252", errors="replace")).strip()
    return (texto, "texto") if len(texto) > 20 else ("", "vacio")


def extraer_docx(path):
    try:
        doc = Document(ruta_sistema(path))
        partes = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for numero_tabla, tabla in enumerate(doc.tables, 1):
            partes.append(f"[TABLA {numero_tabla}]")
            for numero_fila, fila in enumerate(tabla.rows, 1):
                valores = [limpiar_control(celda.text).strip().replace("\n", " | ") for celda in fila.cells]
                if any(valores):
                    partes.append(f"FILA {numero_fila}: " + " || ".join(valores))
        texto = "\n".join(partes).strip()
        return texto, "texto" if len(texto) > 20 else "vacio"
    except Exception as exc:
        return "", f"error_docx: {exc}"


def extraer_excel(path, max_filas, max_columnas):
    try:
        wb = load_workbook(
            ruta_sistema(path),
            data_only=True,
            read_only=True,
            keep_vba=path.suffix.lower() == ".xlsm"
        )
        partes = []
        for hoja in wb.worksheets:
            partes.append(f"[HOJA {hoja.title}]")
            for numero_fila, fila in enumerate(hoja.iter_rows(values_only=True), 1):
                if numero_fila > max_filas:
                    break
                valores = []
                for valor in fila[:max_columnas]:
                    valores.append("" if valor is None else limpiar_control(valor).strip())
                if any(valores):
                    partes.append(f"FILA {numero_fila}: " + " || ".join(valores))
        wb.close()
        texto = "\n".join(partes).strip()
        return texto, "texto" if len(texto) > 20 else "vacio"
    except Exception as exc:
        return "", f"error_excel: {exc}"


def extraer_texto_plano(path):
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(ruta_sistema(path), encoding=encoding) as archivo:
                texto = limpiar_control(archivo.read()).strip()
            return texto, "texto" if len(texto) > 20 else "vacio"
        except Exception:
            pass
    return "", "error_texto"


def encontrar_7zip(ruta_explicitada=None):
    candidatos = [ruta_explicitada, shutil.which("7z"), shutil.which("7z.exe")]
    if os.name == "nt":
        candidatos.extend([
            r"C:\Program Files\7-Zip\7z.exe",
            r"C:\Program Files (x86)\7-Zip\7z.exe"
        ])
    return next((str(x) for x in candidatos if x and Path(x).exists()), None)


def descomprimir(carpeta, seven_zip):
    registros = []
    raiz_extraidos = carpeta / "_extraidos"
    with os.scandir(ruta_sistema(carpeta)) as entradas:
        archivos = [carpeta / entrada.name for entrada in entradas
                    if entrada.is_file() and Path(entrada.name).suffix.lower() in {".zip", ".rar", ".7z"}]
    archivos.sort(key=lambda archivo: archivo.name.lower())
    for archivo in archivos:
        destino = raiz_extraidos / re.sub(r"[^\w.-]", "_", archivo.stem)
        marcador = destino / ".extraido_ok"
        if os.path.isfile(ruta_sistema(marcador)):
            registros.append({"archivo": archivo.name, "estado": "ya_extraido"})
            continue
        Path(ruta_sistema(destino)).mkdir(parents=True, exist_ok=True)
        try:
            if archivo.suffix.lower() == ".zip":
                with zipfile.ZipFile(ruta_sistema(archivo)) as compacto:
                    compacto.extractall(ruta_sistema(destino))
            elif seven_zip:
                proceso = subprocess.run(
                    [seven_zip, "x", "-y", f"-o{ruta_sistema(destino)}", ruta_sistema(archivo)],
                    capture_output=True,
                    text=True,
                    timeout=300
                )
                if proceso.returncode != 0:
                    raise RuntimeError((proceso.stderr or proceso.stdout)[-500:])
            else:
                raise RuntimeError("7-Zip no encontrado")
            with open(ruta_sistema(marcador), "w", encoding="utf-8") as registro:
                registro.write("ok")
            registros.append({"archivo": archivo.name, "estado": "ok"})
        except Exception as exc:
            registros.append({"archivo": archivo.name, "estado": f"error: {exc}"})
    return registros


def extraer_archivo(path, args, detalle=None):
    extension = path.suffix.lower()
    if extension == ".pdf":
        return extraer_pdf(
            path,
            args.max_paginas,
            usar_ocr=args.ocr,
            tesseract_cmd=args.tesseract_cmd,
            idioma_ocr=args.idioma_ocr,
            dpi_ocr=args.dpi_ocr,
            detalle=detalle,
            min_filas_vision=getattr(args, "vision_min_filas", 2),
        )
    if extension in EXTENSIONES_IMAGEN:
        return "", "imagen" if getattr(args, "vision", False) else "no_soportado"
    if extension == ".docx":
        return extraer_docx(path)
    if extension == ".doc":
        return extraer_doc(path)
    if extension in (".xlsx", ".xlsm"):
        return extraer_excel(path, args.max_filas_excel, args.max_columnas_excel)
    if extension in (".txt", ".csv", ".tsv"):
        return extraer_texto_plano(path)
    if extension == ".xls":
        return "", "xls_legacy_no_soportado"
    return "", "no_soportado"


def tipo_documental(nombre):
    texto = unicodedata.normalize("NFKD", str(nombre)).encode("ascii", "ignore").decode("ascii").lower()
    if any(x in texto for x in ("econom", "cotiza", "presupuesto", "precio")):
        return "economico"
    if any(x in texto for x in ("tecnic", "ficha", "especific", "producto")):
        return "tecnico"
    return "otro"


def prioridad(path):
    tipo = tipo_documental(path.name)
    orden = {"economico": 0, "tecnico": 1, "otro": 2}[tipo]
    return orden, path.name.lower()


def pistas(texto):
    lineas = [linea.strip() for linea in texto.splitlines() if linea.strip()]
    productos = [linea for linea in lineas if PATRON_PRODUCTO.search(linea) and len(linea) < 350]
    precios = [linea for linea in lineas if PATRON_PRECIO.search(linea) and len(linea) < 350]
    return productos[:40], precios[:40]


def dividir_texto(texto, limite, lineas_solapadas=3):
    if len(texto) <= limite:
        return [texto]
    lineas = [linea for linea in texto.splitlines() if linea.strip()]
    fragmentos = []
    actual = []
    caracteres = 0
    for linea in lineas:
        trozos = [linea[indice:indice + limite] for indice in range(0, len(linea), limite)] or [""]
        for trozo in trozos:
            largo = len(trozo) + 1
            if actual and caracteres + largo > limite:
                fragmentos.append("\n".join(actual))
                actual = actual[-lineas_solapadas:]
                caracteres = sum(len(valor) + 1 for valor in actual)
            actual.append(trozo)
            caracteres += largo
    if actual:
        fragmentos.append("\n".join(actual))
    return fragmentos


def combinar_consumos(consumos, modelo):
    return {
        "modelo": next((fila.get("modelo") for fila in consumos if fila.get("modelo")), modelo),
        "prompt_tokens": sum(fila.get("prompt_tokens") or 0 for fila in consumos),
        "completion_tokens": sum(fila.get("completion_tokens") or 0 for fila in consumos),
        "total_tokens": sum(fila.get("total_tokens") or 0 for fila in consumos),
        "llamadas": len(consumos),
    }


def parsear_json(respuesta):
    texto = respuesta.strip()
    texto = re.sub(r"^```(?:json)?\s*|\s*```$", "", texto, flags=re.I | re.S)
    try:
        return json.loads(texto), "ok"
    except Exception:
        pass
    decoder = json.JSONDecoder()
    errores = []
    candidatos = []
    for coincidencia in re.finditer(r"\{", texto):
        fragmento = texto[coincidencia.start():]
        try:
            candidato, _ = decoder.raw_decode(fragmento)
            if isinstance(candidato, dict):
                candidatos.append(candidato)
        except json.JSONDecodeError as exc:
            errores.append(str(exc))
    for candidato in candidatos:
        if isinstance(candidato.get("productos"), list):
            return candidato, "ok_extraido"
    if candidatos:
        return candidatos[0], "ok_extraido"

    coincidencia = re.search(r"\{.*\}", texto, re.S)
    if coincidencia:
        reparado = re.sub(r",\s*([}\]])", r"\1", coincidencia.group(0))
        try:
            return json.loads(reparado), "ok_reparado"
        except Exception as exc:
            errores.append(str(exc))
    return None, f"json_invalido: {errores[-1]}" if errores else "sin_json"


PATRON_ERROR_TRANSITORIO = re.compile(r"channel\s*error|model\s+(?:is\s+)?(?:loading|unloaded)|"
                                      r"no\s+models?\s+loaded|connection\s+(?:reset|closed)", re.I)


def es_error_transitorio(texto):
    """Errores de LM Studio que suelen resolverse reintentando (canal interno o LM Link
    recien conectado, modelo cargandose)."""
    return bool(PATRON_ERROR_TRANSITORIO.search(str(texto or "")))


ESPERA_BACKEND_MINUTOS = 30
OCR_DISPONIBLE = False  # se confirma al arrancar con verificar_ocr
_LOCK_ESPERA_BACKEND = threading.Lock()


def url_modelos(backend):
    if backend == "lmstudio":
        return LMSTUDIO_URL.split("/v1/")[0] + "/v1/models"
    return OLLAMA_URL.split("/api/")[0] + "/api/tags"


def backend_responde(backend, timeout=10):
    """True si el servidor del modelo contesta (aunque una peticion concreta haya fallado)."""
    try:
        return requests.get(url_modelos(backend), timeout=timeout).status_code < 500
    except requests.exceptions.RequestException:
        return False


def esperar_backend(backend, minutos):
    """Espera hasta `minutos` a que el servidor vuelva (sondea cada 30 s). Un solo hilo espera; los demas
    esperan a que termine. Devuelve True si volvio."""
    with _LOCK_ESPERA_BACKEND:
        if backend_responde(backend):
            return True
        limite = time.monotonic() + minutos * 60
        print(f"\n  {backend} no responde; se espera hasta {minutos} min a que vuelva...")
        while time.monotonic() < limite:
            if DETENER.wait(30):
                return False
            if backend_responde(backend):
                print(f"  {backend} volvio a responder; se continua.")
                return True
        return False


def calentar_modelo(args, intentos=4, espera=10):
    """Peticion minima antes de la carga real: activa el enlace (LM Link) y confirma que
    el modelo responde. Devuelve los segundos que tardo o None si no respondio."""
    for intento in range(1, intentos + 1):
        inicio = time.perf_counter()
        try:
            _, estado, _, _ = consultar_modelo(
                'Responde solo {"ok": true}', args.modelo, args.timeout, args.num_ctx, args.backend,
                32, 0, 0)
        except RuntimeError as exc:
            estado = str(exc)
        if estado in ESTADOS_IA_OK:
            return round(time.perf_counter() - inicio, 1)
        print(f"  El modelo aun no responde ({estado[:120]}); reintento {intento}/{intentos} en {espera} s...")
        time.sleep(espera)
    return None


def consultar_modelo(
    prompt,
    modelo,
    timeout,
    num_ctx,
    backend="ollama",
    max_tokens=4096,
    reintentos=2,
    espera_reintento=5,
    imagenes=None,
):
    """imagenes: lista opcional de dicts {"mime": ..., "base64": ...} para modelos con vision."""
    verificar_detencion()
    if backend == "lmstudio":
        if imagenes:
            contenido_usuario = [
                {"type": "image_url", "image_url": {"url": f"data:{img['mime']};base64,{img['base64']}"}}
                for img in imagenes
            ] + [{"type": "text", "text": prompt}]
        else:
            contenido_usuario = prompt
        payload = {
            "model": modelo,
            "messages": [
                {"role": "system", "content": "Devuelve exclusivamente un objeto JSON valido."},
                {"role": "user", "content": contenido_usuario}
            ],
            "temperature": 0.0,
            "max_tokens": max_tokens,
            "reasoning_effort": "none",
            "stream": False,
            "response_format": {"type": "text"},
            "chat_template_kwargs": {"enable_thinking": False},
        }
        url = LMSTUDIO_URL
    else:
        payload = {
            "model": modelo,
            "prompt": prompt,
            "stream": False,
            "format": "json",
            "options": {
                "temperature": 0.1,
                "num_ctx": num_ctx,
                "num_predict": max_tokens
            }
        }
        if imagenes:
            payload["images"] = [img["base64"] for img in imagenes]
        url = OLLAMA_URL
    ultimo_error = None
    for intento in range(reintentos + 1):
        try:
            respuesta = requests.post(url, json=payload, timeout=timeout)
            if respuesta.status_code >= 400:
                estado = f"http_{respuesta.status_code}: {respuesta.text[:500]}"
                transitorio = respuesta.status_code in (429, 500, 502, 503, 504) or es_error_transitorio(respuesta.text)
                if not transitorio or intento >= reintentos:
                    return None, estado, "", {}
                ultimo_error = estado
            else:
                datos_respuesta = respuesta.json()
                if backend == "lmstudio":
                    cruda = ((datos_respuesta.get("choices") or [{}])[0]
                             .get("message", {}).get("content") or "")
                else:
                    cruda = datos_respuesta.get("response", "")
                datos, estado = parsear_json(cruda)
                if backend == "lmstudio":
                    uso = datos_respuesta.get("usage") or {}
                    consumo = {
                        "modelo": datos_respuesta.get("model", modelo),
                        "prompt_tokens": uso.get("prompt_tokens"),
                        "completion_tokens": uso.get("completion_tokens"),
                        "total_tokens": uso.get("total_tokens"),
                    }
                else:
                    consumo = {
                        "modelo": modelo,
                        "prompt_tokens": datos_respuesta.get("prompt_eval_count"),
                        "completion_tokens": datos_respuesta.get("eval_count"),
                        "total_tokens": (
                            (datos_respuesta.get("prompt_eval_count") or 0)
                            + (datos_respuesta.get("eval_count") or 0)
                        ),
                    }
                return datos, estado, cruda, consumo
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
            ultimo_error = str(exc)
            if intento >= reintentos:
                if backend_responde(backend):
                    # El servidor esta vivo: es esta peticion la que no termina (documento enorme o modelo
                    # atascado en ella). Se registra como fallo del archivo y la corrida sigue.
                    return None, f"timeout_{backend}: {str(exc)[:200]}", "", {}
                # El servidor no contesta (LM Studio reiniciandose, LM Link caido...): se espera a que vuelva
                # antes de dar la corrida por perdida. Corridas desatendidas sobreviven a cortes cortos.
                if esperar_backend(backend, ESPERA_BACKEND_MINUTOS):
                    return consultar_modelo(prompt, modelo, timeout, num_ctx, backend, max_tokens,
                                            reintentos, espera_reintento, imagenes)
                raise RuntimeError(f"{backend} dejo de responder tras {reintentos + 1} intentos y "
                                   f"{ESPERA_BACKEND_MINUTOS} min de espera: {exc}")
        except Exception as exc:
            return None, f"error_{backend}: {exc}", "", {}
        time.sleep(espera_reintento)
    return None, f"error_{backend}: {ultimo_error}", "", {}


def normalizar_numero(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, (int, float)):
        return valor
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    texto = re.sub(r"[.\-\u2013\u2014]+$", "", texto)  # "214.555.-" -> "214.555"
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", texto):
        texto = texto.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:,\d{3})+", texto):
        texto = texto.replace(",", "")
    elif "," in texto and "." not in texto:
        texto = texto.replace(",", ".")
    texto = re.sub(r"[^\d.]", "", texto)
    try:
        numero = float(texto)
        return int(numero) if numero.is_integer() else numero
    except Exception:
        return None


PATRON_CELDA_NUMERICA = re.compile(r"\s*(?:CLP)?\s*\$?\s*[\d.,]+(?:\.-|-)?\s*(?:CLP)?\s*", re.I)
PATRON_FILA_DE_TOTALES = re.compile(r"\b(?:total|neto|iva|bruto|impuestos?|subtotal)\b", re.I)


def terna_en_celdas(celdas):
    """(posicion_total, cantidad, unitario, total) con cantidad x unitario = total entre las celdas
    numericas de una fila, sin depender de su posicion ni de columnas vacias intermedias."""
    numeros = []
    for indice, celda in enumerate(celdas):
        if celda and PATRON_CELDA_NUMERICA.fullmatch(celda):
            valor = normalizar_numero(celda)
            if isinstance(valor, (int, float)) and valor > 0:
                numeros.append((indice, valor))
    visibles = {valor for _, valor in numeros}
    mejor = None
    for indice_u, unitario in numeros:
        if unitario < 1000:
            continue
        for indice_t, total in numeros:
            if indice_t <= indice_u or total < unitario:
                continue
            cantidad = round(total / unitario)
            if cantidad < 1 or abs(cantidad * unitario - total) > max(2, total * 0.001):
                continue
            if cantidad == 1 and unitario != total:
                continue
            if cantidad != 1 and cantidad not in visibles:
                continue
            if mejor is None or indice_t > mejor[0]:
                mejor = (indice_t, cantidad, unitario, total)
    return mejor


def filas_precio_economico(bloques):
    """Precios de filas con cantidad, unitario y total comprobables en el anexo."""
    filas = []
    for bloque in bloques:
        archivo = bloque.get("archivo") or ""
        if tipo_documental(archivo) != "economico":
            continue
        for linea in bloque.get("texto", "").splitlines():
            if not re.match(r"^FILA\s+\d+:", linea) or "||" not in linea:
                continue
            celdas = [celda.strip() for celda in linea.split("||")]
            celdas[0] = re.sub(r"^FILA\s+\d+:\s*", "", celdas[0])  # "FILA 4: 3" -> "3" (la cantidad)
            if len(celdas) < 4:
                continue
            unitario = normalizar_numero(celdas[-2])
            total = normalizar_numero(celdas[-1])
            terna = None
            if all(isinstance(valor, (int, float)) and valor > 0 for valor in (unitario, total)):
                cantidad = round(total / unitario)
                cantidades_visibles = [normalizar_numero(c) for c in celdas[:-2]]
                if (cantidad > 0 and cantidad in cantidades_visibles
                        and abs(cantidad * unitario - total) <= max(2, total * 0.001)):
                    terna = (len(celdas) - 1, cantidad, unitario, total)
            if terna is None:
                terna = terna_en_celdas(celdas)  # columnas vacias entre unitario y total, otro orden...
            if terna is None:
                continue
            posicion_total, cantidad, unitario, total = terna
            descripcion = " ".join(
                c for indice, c in enumerate(celdas)
                if re.search(r"[A-Za-z\u00c1\u00c9\u00cd\u00d3\u00da\u00e1\u00e9\u00ed\u00f3\u00fa]", c)
                and not PATRON_CELDA_NUMERICA.fullmatch(c))
            if not descripcion or PATRON_FILA_DE_TOTALES.search(descripcion):
                continue

            filas.append({"archivo": archivo, "pagina": bloque.get("pagina"),
                          "descripcion": descripcion, "cantidad": cantidad,
                          "precio_unitario": unitario, "precio_total": total})
    return filas


PATRON_TOKEN_NUMERICO = re.compile(r"^[\d.,$]+$")


def particion_cantidad_unitario_total(texto_digitos, maximo_cantidad=20000):
    """Busca (cantidad, unitario, total) con cantidad x unitario = total en una cadena de digitos
    (los separadores de miles y los espacios ya estan quitados). El orden puede ser
    unitario-cantidad-total o cantidad-unitario-total. Devuelve None si no hay solucion unica."""
    soluciones = set()
    n = len(texto_digitos)
    for i in range(1, n - 1):
        for j in range(i + 1, n):
            partes = texto_digitos[:i], texto_digitos[i:j], texto_digitos[j:]
            if any(p.startswith("0") for p in partes):
                continue
            a, b, total = (int(p) for p in partes)
            for cantidad, unitario in ((b, a), (a, b)):
                # el unitario puede venir redondeado (735.069,17 -> 735.069): se tolera menos de 1 peso por unidad
                if (1 <= cantidad <= maximo_cantidad and unitario >= 1000 and total >= unitario
                        and abs(cantidad * unitario - total) < cantidad):
                    soluciones.add((cantidad, unitario, total))
    return next(iter(soluciones)) if len(soluciones) == 1 else None


def filas_precio_texto_libre(bloques):
    """Filas economicas en texto plano con cifras partidas por espacios del PDF
    ("6 75.620 1 0 6 .756.200" = 10 x 675.620 = 6.756.200). Solo acepta una fila si existe
    una unica reconstruccion con cantidad x unitario = total."""
    filas = []
    for bloque in bloques:
        archivo = bloque.get("archivo") or ""
        if tipo_documental(archivo) != "economico":
            continue
        for linea in bloque.get("texto", "").splitlines():
            if "||" in linea or len(linea) > 400:
                continue
            fichas = linea.split()
            inicio = len(fichas)
            while inicio > 0 and PATRON_TOKEN_NUMERICO.match(fichas[inicio - 1]):
                inicio -= 1
            descripcion = " ".join(fichas[:inicio])
            if len(re.findall(r"[A-Za-zÁÉÍÓÚáéíóú]", descripcion)) < 6:
                continue
            for k in range(inicio, len(fichas) - 2):  # prueba colas cada vez mas cortas
                digitos = re.sub(r"\D", "", "".join(fichas[k:]))
                if not 8 <= len(digitos) <= 24:
                    continue
                solucion = particion_cantidad_unitario_total(digitos)
                if solucion:
                    cantidad, unitario, total = solucion
                    extra = " ".join(fichas[inicio:k])
                    filas.append({"archivo": archivo, "pagina": bloque.get("pagina"),
                                  "descripcion": f"{descripcion} {extra}".strip(), "cantidad": cantidad,
                                  "precio_unitario": unitario, "precio_total": total})
                    break
    return filas


def todas_las_filas_economicas(bloques):
    filas = filas_precio_economico(bloques)
    vistos = {(f["archivo"], f["precio_unitario"], f["cantidad"]) for f in filas}
    for fila in filas_precio_texto_libre(bloques):
        clave = (fila["archivo"], fila["precio_unitario"], fila["cantidad"])
        if clave not in vistos:
            vistos.add(clave)
            filas.append(fila)
    return filas


PATRON_NUMERO_EN_TEXTO = re.compile(r"\d[\d.,]*")


def numeros_del_texto(texto):
    """Todos los numeros enteros que aparecen escritos en un texto (con o sin separador de miles)."""
    encontrados = set()
    for coincidencia in PATRON_NUMERO_EN_TEXTO.finditer(texto):
        valor = normalizar_numero(coincidencia.group(0).rstrip(".,"))
        if isinstance(valor, (int, float)):
            encontrados.add(round(valor))
    return encontrados


def marcar_precios_sin_respaldo(productos, numeros_documentos, bloques_economicos):
    """Marca los productos cuyo precio unitario y total no aparecen en NINGUN documento del proveedor
    ni en una fila verificable: el modelo pudo inventarlo o copiarlo de otro producto."""
    pares = {(f["precio_unitario"], f["precio_total"]) for f in todas_las_filas_economicas(bloques_economicos)}
    for producto in productos:
        unitario, total = producto.get("precio_unitario"), producto.get("precio_total")
        valores = [round(v) for v in (unitario, total) if isinstance(v, (int, float)) and v > 0]
        if not valores or (unitario, total) in pares:
            continue
        if producto.get("precio_desde_linea_agregada"):
            continue  # unitario deducido de "N equipos || total": el documento trae el total, no el unitario
        if not any(v in numeros_documentos for v in valores):
            producto["precio_sin_respaldo"] = True
    return productos


def reconciliar_precios_con_anexos(productos, bloques, total_oferta=None, anular=True):
    """Corrige solo filas economicas inequivocas; anula precios sin respaldo en ellas."""
    filas = todas_las_filas_economicas(bloques)
    por_archivo = {}
    for fila in filas:
        por_archivo.setdefault(fila["archivo"], []).append(fila)
    resultado = [dict(p) for p in productos]
    asignaciones = {}
    pares_verificados = {archivo: {(f["precio_unitario"], f["precio_total"]) for f in lista}
                         for archivo, lista in por_archivo.items()}
    numeros_por_archivo = {}
    for bloque in bloques:
        numeros_por_archivo.setdefault(bloque.get("archivo"), set()).update(numeros_del_texto(bloque.get("texto", "")))
    pares_en_productos = {(p.get("fuente_precio") or p.get("archivo_fuente"), p.get("precio_unitario"),
                           p.get("precio_total")) for p in resultado}
    for fila in filas:
        if (fila["archivo"], fila["precio_unitario"], fila["precio_total"]) in pares_en_productos:
            continue  # la fila ya tiene su producto en ese archivo: no puede ser robada por uno mal leido
        texto_fila = texto_normalizado(fila["descripcion"])
        puntuaciones = []
        for indice, producto in enumerate(resultado):
            if producto.get("fuente_precio") != fila["archivo"]:
                continue
            # Un producto cuyo par (unitario, total) NO esta entre las filas verificables de su anexo no
            # tiene respaldo (p.ej. cifras partidas leidas mal): puede corregirse con la fila que
            # describe el mismo equipo aunque la cantidad leida sea distinta.
            respaldado = (producto.get("precio_unitario"), producto.get("precio_total")) in pares_verificados.get(
                fila["archivo"], set())
            cantidad = producto.get("cantidad")
            if cantidad not in (None, "") and cantidad != fila["cantidad"] and respaldado:
                continue
            tokens = {t for t in texto_normalizado(producto.get("producto")).split()
                      if len(t) >= (4 if respaldado else 3)}
            comunes = {t for t in tokens if t in texto_fila.split()}
            modelo_base = re.sub(r"\s*\([^)]*\)", "", str(producto.get("modelo") or ""))
            coincide_modelo = len(modelo_base) >= 5 and texto_normalizado(modelo_base) in texto_fila
            codigo_comun = (not respaldado) and any(
                re.search(r"[a-z]", t) and re.search(r"\d", t) and len(t) >= 5 for t in comunes)
            if len(comunes) < 4 and not coincide_modelo and not codigo_comun:
                continue
            puntuacion = len(comunes) + (3 if coincide_modelo or codigo_comun else 0)
            puntuaciones.append((puntuacion, len(comunes), indice))
        puntuaciones.sort(reverse=True)
        if not puntuaciones or puntuaciones[0][0] < 3 or puntuaciones[0][1] < 2:
            continue
        if len(puntuaciones) > 1 and puntuaciones[0][0] <= puntuaciones[1][0] * 1.2:
            continue
        indice = puntuaciones[0][2]
        asignaciones.setdefault(indice, []).append(fila)

    # Un kit de una sola linea puede describirse genericamente en el economico:
    # el modelo del PC solo aparece en el tecnico. Asociar solo si la identidad
    # principal, la cantidad y el total general son univocos.
    if len(filas) == 1 and not asignaciones:
        fila = filas[0]
        descripcion = texto_normalizado(fila["descripcion"])
        referencia = normalizar_numero(total_oferta)
        candidatos = [
            indice for indice, producto in enumerate(resultado)
            if producto.get("categoria") == "equipo"
            and producto.get("fuente_precio") in (None, "", fila["archivo"])
            and re.search(r"\b(?:desktop|computador|notebook|pc)\b", texto_normalizado(producto.get("producto")))
            and (producto.get("cantidad") in (None, "", fila["cantidad"]))
        ]
        if (len(candidatos) == 1 and referencia is not None
                and abs(referencia - fila["precio_total"]) <= 2
                and re.search(r"\b(?:desktop|computador|notebook|pc)\b", descripcion)):
            asignaciones[candidatos[0]] = [fila]

    for indice, producto in enumerate(resultado):
        archivo = producto.get("fuente_precio")
        asociadas = asignaciones.get(indice) or []
        if len(asociadas) == 1:
            fila = asociadas[0]
            producto["fuente_precio"] = fila["archivo"]
            producto["cantidad"] = fila["cantidad"]
            if not producto.get("cantidad_fuente"):
                producto["cantidad_fuente"] = "anexo_economico"
            producto["precio_unitario"] = fila["precio_unitario"]
            producto["precio_total"] = fila["precio_total"]
            producto["precio_total_tipo"] = "linea"
            producto["pagina_precio"] = fila["pagina"]
            producto["precio_reconciliado_fuente"] = True
            producto.pop("precio_fuente_no_coincide", None)
        elif archivo in por_archivo and (producto.get("precio_unitario") not in (None, "")
                                       or producto.get("precio_total") not in (None, "")):
            pares = {(f["precio_unitario"], f["precio_total"]) for f in por_archivo[archivo]}
            par = (producto.get("precio_unitario"), producto.get("precio_total"))
            asignada_a_otro = any(
                f["precio_unitario"] == par[0] and f["precio_total"] == par[1]
                for otro_indice, usadas in asignaciones.items() if otro_indice != indice for f in usadas
            )
            en_el_texto = any(isinstance(v, (int, float)) and round(v) in numeros_por_archivo.get(archivo, set())
                              for v in par)
            if anular and (par not in pares or asignada_a_otro) and not (en_el_texto and not asignada_a_otro):
                producto["precio_unitario"] = None
                producto["precio_total"] = None
                producto["precio_total_tipo"] = None
                producto["precio_fuente_no_coincide"] = True
    return descartar_copias_mal_leidas(resultado, filas)


def descartar_copias_mal_leidas(productos, filas):
    """Si un producto lleva un par (unitario, total) que ninguna fila verificable respalda y existe
    otro producto con la MISMA descripcion y cantidad cuyo par si esta respaldado, el primero es una
    copia mal leida (cifras partidas): se descarta y su archivo queda como respaldo del buen producto."""
    verificados = {(f["precio_unitario"], f["precio_total"]) for f in filas}
    if not verificados:
        return productos
    respaldados = [p for p in productos if (p.get("precio_unitario"), p.get("precio_total")) in verificados]
    descartar = set()
    for indice, producto in enumerate(productos):
        if (producto.get("precio_unitario"), producto.get("precio_total")) in verificados:
            continue
        for bueno in respaldados:
            if (bueno is not producto
                    and (bueno.get("cantidad") == producto.get("cantidad") or producto.get("precio_unitario") in (None, ""))
                    and texto_normalizado(bueno.get("producto")) == texto_normalizado(producto.get("producto"))):
                fuentes = set(bueno.get("fuentes_respaldo") or [])
                fuentes.update(str(producto.get(c)) for c in ("archivo_fuente", "fuente_precio", "fuente_producto")
                               if producto.get(c))
                bueno["fuentes_respaldo"] = sorted(fuentes)
                descartar.add(indice)
                break
    return [p for i, p in enumerate(productos) if i not in descartar]


def fila_representada(fila, productos):
    """True si algun producto ya lleva el unitario de la fila (y su cantidad, si la tiene)."""
    return any(
        isinstance(p.get("precio_unitario"), (int, float))
        and abs(p["precio_unitario"] - fila["precio_unitario"]) <= 1
        and p.get("cantidad") in (None, "", fila["cantidad"])
        for p in productos
    )


def filas_economicas_sin_producto(productos, bloques):
    """Filas del anexo economico (cantidad x unitario = total, comprobables) que son equipos
    del alcance y ningun producto recoge. Devuelve [(fila, candidato)]."""
    faltantes = []
    for fila in todas_las_filas_economicas(bloques):
        if fila_representada(fila, productos):
            continue
        descripcion = re.sub(r"^FILA\s+\d+:\s*", "", fila["descripcion"]).strip()
        descripcion = re.sub(r"^\d+\s+(?=\S)", "", descripcion)
        descripcion = re.sub(r"\s+\d+\s*(?:meses|mes|a[nñ]os?|d[ií]as?)\b.*$", "", descripcion, flags=re.I).strip()
        candidato = {"producto": descripcion}
        categoria = clasificar_producto(candidato)
        if categoria:
            candidato["categoria"] = categoria
            faltantes.append((fila, candidato))
    return faltantes


def rescatar_filas_economicas(productos, parciales, bloques, proveedor=None, rut=None):
    """Incorpora filas economicas verificables que el modelo no asocio a ningun producto.

    - Si hay UN solo producto sin precio del mismo tipo (p.ej. el notebook de una ficha
      tecnica) y una sola fila de ese tipo, se completan sus precios.
    - Si no, se agrega la fila como producto propio; la marca solo se toma de las fichas
      del proveedor cuando todas las del mismo tipo comparten una unica marca.
    Todo lo rescatado queda marcado para revision; no se inventa ningun valor.
    """
    resultado = [dict(p) for p in productos]
    faltantes = filas_economicas_sin_producto(resultado, bloques)
    por_tipo = {}
    for fila, candidato in faltantes:
        por_tipo.setdefault(subcategoria_producto(candidato), []).append((fila, candidato))
    for tipo, filas in por_tipo.items():
        sin_precio = [
            p for p in resultado
            if p.get("precio_unitario") in (None, "") and p.get("precio_total") in (None, "")
            and subcategoria_producto(p) == tipo
        ]
        for fila, candidato in filas:
            datos_fila = {
                "cantidad": fila["cantidad"], "precio_unitario": fila["precio_unitario"],
                "precio_total": fila["precio_total"], "precio_total_tipo": "linea",
                "fuente_precio": fila["archivo"], "pagina_precio": fila["pagina"],
                "rescatado_fila_economica": True,
            }
            if len(filas) == 1 and len(sin_precio) == 1 and sin_precio[0].get("cantidad") in (None, "", fila["cantidad"]):
                sin_precio[0].update(datos_fila)
                sin_precio[0].pop("precio_fuente_no_coincide", None)
                sin_precio[0]["cantidad_fuente"] = sin_precio[0].get("cantidad_fuente") or "anexo_economico"
                continue
            marcas = {texto_normalizado(p.get("marca")): p.get("marca") for p in (*resultado, *parciales)
                      if p.get("marca") and subcategoria_producto(p) == tipo}
            nuevo = {
                "proveedor": proveedor, "rut": rut,
                "item": None, "producto": candidato["producto"], "modelo": None,
                "categoria": candidato["categoria"], "tipo_equipo": tipo,
                "marca": next(iter(marcas.values())) if len(marcas) == 1 else None,
                "metodo_marca": "tipo_con_marca_unica" if len(marcas) == 1 else None,
                "cantidad_fuente": "anexo_economico", "confianza": "media",
                "fuente_producto": fila["archivo"], "pagina_producto": fila["pagina"],
                "archivo_fuente": fila["archivo"],
                **datos_fila,
            }
            resultado.append(nuevo)
    return resultado


def normalizar_producto(producto, proveedor, rut, archivo, tipo_doc):
    descripcion = str(producto.get("producto") or "").strip()
    if not descripcion:
        return None
    cantidad = normalizar_numero(producto.get("cantidad"))
    precio_unitario = normalizar_numero(producto.get("precio_unitario"))
    precio_total = normalizar_numero(producto.get("precio_total"))
    precio_total_tipo = str(producto.get("precio_total_tipo") or "").strip().lower() or None
    if precio_total_tipo not in (None, "linea", "oferta"):
        precio_total_tipo = None
    cantidad_fuente = producto.get("cantidad_fuente")
    if (
        precio_unitario is None
        and precio_total is not None
        and cantidad not in (None, 0)
        and precio_total_tipo == "linea"
    ):
        precio_unitario = round(float(precio_total) / float(cantidad), 2)
    if (
        cantidad is None
        and precio_unitario not in (None, 0)
        and precio_total not in (None, 0)
        and precio_total_tipo == "linea"
    ):
        proporcion = float(precio_total) / float(precio_unitario)
        cantidad_entera = round(proporcion)
        if cantidad_entera > 0 and abs(proporcion - cantidad_entera) < 0.01:
            cantidad = cantidad_entera
            cantidad_fuente = "inferida_total_dividido_unitario"
    return {
        "proveedor": proveedor,
        "rut": rut,
        "item": producto.get("item"),
        "producto": descripcion,
        "marca": producto.get("marca"),
        "modelo": producto.get("modelo"),
        "categoria": producto.get("categoria"),
        "tipo_equipo": producto.get("tipo") or producto.get("tipo_equipo"),
        "cantidad": cantidad,
        "cantidad_fuente": cantidad_fuente,
        "precio_unitario": precio_unitario,
        "precio_total": precio_total,
        "precio_total_tipo": precio_total_tipo,
        "moneda": producto.get("moneda"),
        "confianza": producto.get("confianza"),
        "pagina": producto.get("pagina"),
        "pagina_producto": normalizar_numero(producto.get("pagina_producto")),
        "pagina_precio": normalizar_numero(producto.get("pagina_precio")),
        "fila_fuente": producto.get("fila_fuente"),
        "evidencia": producto.get("evidencia"),
        "fuente_producto": producto.get("fuente_producto"),
        "fuente_precio": producto.get("fuente_precio"),
        "fuente_marca": producto.get("fuente_marca"),
        "pagina_marca": normalizar_numero(producto.get("pagina_marca")),
        "metodo_marca": producto.get("metodo_marca"),
        "archivo_fuente": archivo,
        "tipo_documental": tipo_doc
    }


PATRON_EQUIPO_RELEVANTE = re.compile(
    r"notebook|laptop|port[aá]til|ultrabook|chromebook|all.?in.?one|\baio\b|"
    r"desktop|escritorio|\bpc\b|computador(?:a)?|workstation|thinkcentre|"
    r"thinkstation|thinkpad|thinkbook|prodesk|elitedesk|optiplex|latitude|pavilion|macbook|mac\s*mini|imac|"
    r"ideapad|ideacentre|vivobook|zenbook|expertbook|expertcenter|zbook|elitebook|probook|"
    r"proone|eliteone|vostro|inspiron|veriton|aspire|travelmate|surface\s+(?:laptop|pro)|"
    r"\bloq\b|\bvictus\b|\bomen\b", re.I
)
PATRON_TV = re.compile(r"\b(?:smart\s*)?t\.?v\.?\b|televisor", re.I)
PATRON_EXCLUIDO = re.compile(
    r"proyector|tablet|\btab\s+one\b|android|ipad|celular|smartphone|servidor|\bservers?\b|thinksystem|poweredge|proliant|uniserver|"
    r"storage|switch|router|"
    r"access point|\bred\b|rack|cable|"
    r"teclado|mouse|rat[oó]n|docking|dock|base de expansi[oó]n|"
    r"tinta|t[oó]ner|cartucho|tambor|repuesto|licencia|instalaci[oó]n|servicio",
    re.I
)
# Accesorios y servicios que nombran un equipo pero no lo son:
# "Estuche de cuero para MacBook Pro", "FortiSIEM All-In-One Subscription License".
PATRON_ACCESORIO_INICIAL = re.compile(
    r"^\W*(?:n.?\s*\d+\s+)?(?:\d+\W+)?(?:estuche|funda|malet[ií]n|mochila|bolso|cargador|adaptador|protector|"
    r"l[aá]mina|soporte|base\s+refrigerante|candado|hub|kit|lic(?:encia|ense)|renovaci[oó]n|"
    r"extensi[oó]n\s+de\s+garant[ií]a|garant[ií]a)\b", re.I)
# Perifericos y accesorios cuyo nombre contiene "notebook"/"pc" ("Carro de carga para 45
# notebooks", "Escaner Epson", "UPS", "Kensington Maletin"). Se evaluan al inicio de la
# descripcion o, para accesorios con marca delante, en sus primeras palabras.
PATRON_PERIFERICO_INICIAL = re.compile(
    r"^\W*(?:carro|carrito|esc[aá]ner|scanner|ups\b|webcam|c[aá]mara|concentrador|"
    r"parlante|aud[ií]fono|pizarra|lector\s+de\s+(?:c[oó]digo|barra|tarjeta)|gabinete|"
    r"estaci[oó]n\s+de\s+carga|bater[ií]a|fuente\s+de\s+poder)", re.I)
PATRON_ACCESORIO_CABEZA = re.compile(
    r"malet[ií]n|mochila|bolso|funda|estuche|cargador|adaptador|\bups\b|"
    r"carro\s+(?:de\s+)?(?:carga|transporte|tecnol)", re.I)
PATRON_SERVICIO_SOFTWARE = re.compile(r"subscription|suscripci[oó]n|\bsaas\b|per\s+device|por\s+dispositivo", re.I)
PATRON_MONITOR = re.compile(r"monitor|display", re.I)
PATRON_IMPRESORA = re.compile(r"impresora|multifuncional|plotter", re.I)
PATRON_ESPECIFICACION = re.compile(
    r"procesador|cpu|\bram\b|memoria(?:\s+(?:ram|ddr))?|\bssd\b|\bhdd\b|"
    r"disco(?:\s+duro)?|puertos?|conectividad|sistema operativo|windows|"
    r"fuente de poder|garant[ií]a|servicio|instalaci[oó]n|configuraci[oó]n|"
    r"servidor|storage|switch|router|access point|\bred\b|cable|rack", re.I
)


MARCAS_ALIAS = {
    "hewlett packard": "HP", "hewlett-packard": "HP", "hp inc": "HP", "hp inc.": "HP",
    "hewlett packard enterprise": "HP", "ozxen by exact": "OZXEN", "ozxen by exact s.a.": "OZXEN",
}
# Fabricantes de componentes o valores vacios que el modelo devuelve como marca del equipo.
MARCAS_NO_VALIDAS = {
    "intel", "amd", "nvidia", "qualcomm", "microsoft", "windows", "n/a", "na", "null", "none",
    "no aplica", "sin marca", "no especificada", "no especificado", "generica", "generico",
    "no indica", "s/m",
}


MARCAS_CONOCIDAS = {m.lower(): m for m in (
    "HP", "Lenovo", "ASUS", "Dell", "Apple", "Acer", "Samsung", "LG", "Epson", "Brother", "Canon",
    "MSI", "ViewSonic", "Xiaomi", "Sony", "Kyocera", "Ricoh", "Xerox", "Lexmark", "AOC", "Philips",
    "Gigabyte", "Huawei", "Logitech", "Zebra", "OZXEN", "Microsoft Surface")}


def normalizar_marca(marca):
    """Unifica alias de marca y descarta fabricantes de componentes; None si no es una marca."""
    texto = " ".join(str(marca or "").split())
    if not texto:
        return None
    clave = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii").lower()
    if clave in MARCAS_NO_VALIDAS or re.fullmatch(r"(?:marca\s+)?(?:generic[oa]?s?|sin marca|no (?:indica|especifica)\w*)", clave):
        return None
    partes = [parte for parte in re.split(r"\s*[/+,]\s*|\s+y\s+", clave) if parte]
    if len(partes) > 1:
        # "Lenovo/HP", "Asus/MSI/Logitech": varias marcas en un campo (componentes de un equipo armado o
        # lectura ambigua). Ninguna es "la" marca: se deja vacia y la inferencia por descripcion decide.
        return None
    return MARCAS_ALIAS.get(clave) or MARCAS_CONOCIDAS.get(clave, texto)


def clasificar_producto(producto):
    texto = " ".join(str(producto.get(c) or "") for c in ("producto", "modelo", "categoria"))
    if PATRON_TV.search(texto):
        return None
    descripcion = str(producto.get("producto") or "")
    if PATRON_ACCESORIO_INICIAL.search(descripcion) or PATRON_SERVICIO_SOFTWARE.search(descripcion):
        return None
    if PATRON_PERIFERICO_INICIAL.search(descripcion):
        return None
    if (re.search(r"esc[a\u00e1]ner|scanner", descripcion[:40], re.I)
            and not re.search(r"impresora|multifuncional", descripcion, re.I)):
        return None  # "Brother Escaner ADS-3100"; un multifuncional "con escaner" si es impresora
    if re.match(r"\W*(?:armado|ensamblaje|instalaci[o\u00f3]n|configuraci[o\u00f3]n)\b", descripcion, re.I):
        return None  # mano de obra: "ARMADO COMPUTADOR SIN SISTEMA OPERATIVO"
    if (PATRON_ACCESORIO_CABEZA.search(descripcion[:45])
            and not PATRON_EQUIPO_RELEVANTE.match(descripcion.lstrip())
            and not re.match(r"\W*(?:monitor|impresora|multifuncional|all.?in.?one)", descripcion, re.I)):
        return None
    # Una impresora "a tinta" o un monitor "con cable HDMI" siguen siendo el equipo: manda el sustantivo inicial.
    inicio = descripcion.lstrip(" 	-*0123456789.)(")
    if re.match(r"(?:impresoras?|multifuncional(?:es)?)\b", inicio, re.I):
        return "impresora"
    if re.match(r"monitor(?:es)?\b", inicio, re.I):
        return "monitor"
    if PATRON_EQUIPO_RELEVANTE.search(texto):
        return "equipo"
    if PATRON_EXCLUIDO.search(texto):
        return None
    if PATRON_ESPECIFICACION.search(texto) and not any(
        patron.search(texto) for patron in (PATRON_EQUIPO_RELEVANTE, PATRON_MONITOR, PATRON_IMPRESORA)
    ):
        return None
    if PATRON_MONITOR.search(texto):
        return "monitor"
    if PATRON_IMPRESORA.search(texto):
        return "impresora"
    categoria_declarada = str(producto.get("categoria") or "").strip().lower()
    if categoria_declarada in {"equipo", "monitor", "impresora"}:
        return categoria_declarada
    return None


PATRON_TIPO_AIO = re.compile(r"all.?in.?one|\baio\b|todo.{0,2}en.{0,2}uno|pro\s?one|elite\s?one|\bimac\b|pavilion\s+\d{2}-|\b\d{2}-cr\d", re.I)
PATRON_TIPO_WORKSTATION = re.compile(r"workstation|estaci[oó]n\s+de\s+trabajo|thinkstation|\bz[248]\s+g\d", re.I)
PATRON_TIPO_NOTEBOOK = re.compile(
    r"notebook|laptop|port[aá]til|ultrabook|chromebook|macbook|probook|elitebook|zbook|thinkpad|thinkbook|"
    r"ideapad|latitude|vostro|vivobook|zenbook|expertbook|travelmate|\b1[45]-[a-z]{2}\d", re.I)
PATRON_TIPO_DESKTOP = re.compile(
    r"desktop|escritorio|\btorre\b|\bpc\b|\bsff\b|mini\s*pc|prodesk|elitedesk|optiplex|thinkcentre|"
    r"veriton|expertcenter|mac\s*mini|computador", re.I)


def id_revision(licitacion, rut, producto):
    """Identificador estable de un producto para la revision: depende de su contenido,
    no de su posicion, asi las decisiones guardadas no se cruzan si cambia el orden."""
    import hashlib
    clave = "|".join(str(producto.get(c) or "") for c in ("producto", "precio_unitario", "cantidad", "archivo_fuente"))
    return f"{licitacion}__{rut}__{hashlib.sha1(clave.encode('utf-8')).hexdigest()[:8]}"


TIPOS_EQUIPO = {"notebook", "all-in-one", "desktop", "workstation", "monitor", "impresora"}
ALIAS_TIPO = {
    "laptop": "notebook", "portatil": "notebook", "ultrabook": "notebook", "chromebook": "notebook",
    "aio": "all-in-one", "all in one": "all-in-one", "allinone": "all-in-one", "todo en uno": "all-in-one",
    "pc": "desktop", "torre": "desktop", "computador": "desktop", "mini pc": "desktop", "sff": "desktop",
    "escritorio": "desktop", "estacion de trabajo": "workstation", "multifuncional": "impresora",
    "printer": "impresora", "pantalla": "monitor",
}


def normalizar_tipo(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode("ascii")
    texto = " ".join(texto.lower().replace("_", " ").replace("-", " ").split())
    if texto in TIPOS_EQUIPO:
        return texto
    if texto.replace(" ", "-") in TIPOS_EQUIPO:
        return texto.replace(" ", "-")
    return ALIAS_TIPO.get(texto)


PATRON_IMPRESORA_FUERTE = re.compile(
    r"impresora|multifuncional|printer|eco\s?tank|smart\s?tank|laserjet|deskjet|officejet|workforce|inkjet|"
    r"plotter|t[oó]ner|tinta", re.I)
MARCAS_SOLO_IMPRESION = {"epson", "brother", "canon", "ricoh", "kyocera", "xerox", "lexmark", "zebra", "pantum",
                         "konica minolta", "sharp", "oki"}


def subcategoria_producto(producto):
    """Tipo de equipo para analisis (ver _subcategoria_producto). Una impresora "AIO" (all-in-one de impresion:
    imprime, copia y escanea) no es un computador All-in-One: si el texto es de impresora o la marca solo
    fabrica impresoras, manda impresora."""
    tipo = _subcategoria_producto(producto)
    if tipo in ("all-in-one", "desktop", "equipo (sin especificar)", None):
        texto = " ".join(str(producto.get(c) or "") for c in ("producto", "modelo"))
        marca = str(producto.get("marca") or "").strip().lower()
        if PATRON_IMPRESORA_FUERTE.search(texto) or marca in MARCAS_SOLO_IMPRESION:
            return "impresora"
    return tipo


def _subcategoria_producto(producto):
    """Tipo de equipo para analisis: notebook, all-in-one, desktop, workstation,
    monitor o impresora. 'equipo (sin especificar)' si no se puede distinguir.
    Prefiere el tipo declarado por el modelo al leer el documento."""
    declarado = normalizar_tipo(producto.get("tipo_equipo") or producto.get("tipo"))
    if declarado:
        return declarado  # el modelo leyo el documento: un AIO descrito con "monitor" sigue siendo all-in-one
    categoria = producto.get("categoria") or clasificar_producto(producto)
    if categoria == "monitor":
        return "monitor"
    if categoria == "impresora":
        return "impresora"
    texto = " ".join(str(producto.get(c) or "") for c in ("producto", "modelo"))
    for tipo, patron in (("all-in-one", PATRON_TIPO_AIO), ("workstation", PATRON_TIPO_WORKSTATION),
                         ("notebook", PATRON_TIPO_NOTEBOOK), ("desktop", PATRON_TIPO_DESKTOP)):
        if patron.search(texto):
            return tipo
    return "equipo (sin especificar)" if categoria == "equipo" else None


def filtrar_productos_relevantes(productos):
    clasificados = []
    for producto in productos:
        categoria = clasificar_producto(producto)
        if categoria:
            producto["categoria"] = categoria
            clasificados.append(producto)
    return clasificados


def texto_normalizado(valor):
    texto = re.sub(r"[^a-z0-9]+", " ", str(valor or "").lower())
    return " ".join(texto.split())


def numeros_iguales(valor_a, valor_b, tolerancia=0.01):
    if valor_a in (None, "") or valor_b in (None, ""):
        return True
    valor_a = float(valor_a)
    valor_b = float(valor_b)
    return abs(valor_a - valor_b) <= max(1.0, abs(valor_b) * tolerancia)


def productos_equivalentes(producto_a, producto_b):
    if producto_a.get("categoria") != producto_b.get("categoria"):
        return False
    marca_a = texto_normalizado(producto_a.get("marca"))
    marca_b = texto_normalizado(producto_b.get("marca"))
    if marca_a and marca_b and marca_a != marca_b:
        return False
    item_a = texto_normalizado(producto_a.get("item"))
    item_b = texto_normalizado(producto_b.get("item"))
    if item_a and item_b and item_a != item_b:
        # Items distintos: lineas distintas, salvo que sea el mismo producto (descripcion especifica) listado
        # en dos archivos que numeran sus items por separado.
        fuente_a = producto_a.get("fuente_precio") or producto_a.get("archivo_fuente")
        fuente_b = producto_b.get("fuente_precio") or producto_b.get("archivo_fuente")
        mismo_producto_en_otro_archivo = (
            fuente_a and fuente_b and fuente_a != fuente_b
            and descripcion_especifica(producto_a.get("producto"))
            and texto_normalizado(producto_a.get("producto")) == texto_normalizado(producto_b.get("producto")))
        if not mismo_producto_en_otro_archivo:
            return False
    mismo_modelo = (
        len(texto_normalizado(producto_a.get("modelo"))) >= 4
        and texto_normalizado(producto_a.get("modelo"))
        == texto_normalizado(producto_b.get("modelo"))
    )
    misma_descripcion = (
        texto_normalizado(producto_a.get("producto"))
        == texto_normalizado(producto_b.get("producto"))
    )
    mismos_numeros = all(
        numeros_iguales(producto_a.get(campo), producto_b.get(campo))
        for campo in ("cantidad", "precio_unitario", "precio_total")
    )
    return mismos_numeros and (mismo_modelo or misma_descripcion)


def deduplicar_productos(productos):
    unicos = []
    for producto in productos:
        existente = next(
            (candidato for candidato in unicos if productos_equivalentes(candidato, producto)),
            None,
        )
        if existente is None:
            unicos.append(producto)
            continue
        fuentes = set(existente.get("fuentes_respaldo") or [])
        fuentes.update(producto.get("fuentes_respaldo") or [])
        for campo in ("archivo_fuente", "fuente_producto", "fuente_precio"):
            if existente.get(campo):
                fuentes.add(str(existente[campo]))
            if producto.get(campo):
                fuentes.add(str(producto[campo]))
        existente["fuentes_respaldo"] = sorted(fuentes)
        for campo in (
            "marca", "modelo", "item", "cantidad", "cantidad_fuente",
            "precio_unitario", "precio_total", "precio_total_tipo", "moneda", "evidencia",
            "pagina", "pagina_producto", "pagina_precio", "fila_fuente",
            "fuente_producto", "fuente_precio", "tipo_equipo",
            "fuente_marca", "pagina_marca", "metodo_marca",
        ):
            if not existente.get(campo) and producto.get(campo):
                existente[campo] = producto[campo]
    return unicos


def completar_marcas_desde_parciales(productos, parciales):
    """Recupera una marca omitida solo con identidad concreta y fuentes univocas.

    No usa precios, la marca mas frecuente del proveedor ni modelos de componentes.
    Los hallazgos deben proceder del mismo proveedor (como ocurre en cada corrida).
    """
    resultado = []
    for original in productos:
        producto = dict(original)
        resultado.append(producto)
        if producto.get("marca"):
            continue
        candidatos = []
        modelo = texto_normalizado(producto.get("modelo"))
        descripcion = texto_normalizado(producto.get("producto"))
        for parcial in parciales:
            if not parcial.get("marca") or producto.get("categoria") != parcial.get("categoria"):
                continue
            if any(producto.get(campo) and parcial.get(campo)
                   and texto_normalizado(producto[campo]) != texto_normalizado(parcial[campo])
                   for campo in ("item", "proveedor", "rut")):
                continue
            modelo_parcial = texto_normalizado(parcial.get("modelo"))
            if modelo and modelo_parcial and modelo != modelo_parcial:
                continue
            mismo_modelo = len(modelo) >= 4 and bool(re.search(r"\d", modelo)) and modelo == modelo_parcial
            misma_descripcion = (len(descripcion) >= 10 and bool(re.search(r"\d", descripcion))
                                 and descripcion == texto_normalizado(parcial.get("producto")))
            if mismo_modelo or misma_descripcion:
                candidatos.append(parcial)
        marcas = {texto_normalizado(p["marca"]) for p in candidatos}
        if len(marcas) != 1:
            continue
        fuente = candidatos[0]
        producto["marca"] = fuente["marca"]
        producto["fuente_marca"] = fuente.get("fuente_marca") or fuente.get("fuente_producto") or fuente.get("archivo_fuente")
        producto["pagina_marca"] = fuente.get("pagina_marca") or fuente.get("pagina_producto") or fuente.get("pagina")
        producto["metodo_marca"] = "cruce_identidad_exacta"
    return resultado


def coincide_con_parcial(producto, parcial):
    if producto.get("categoria") != parcial.get("categoria"):
        return False
    modelo = texto_normalizado(producto.get("modelo"))
    modelo_parcial = texto_normalizado(parcial.get("modelo"))
    if len(modelo) >= 4 and modelo == modelo_parcial:
        return True
    tokens = set(texto_normalizado(producto.get("producto")).split())
    tokens_parcial = set(texto_normalizado(parcial.get("producto")).split())
    if not tokens or not tokens_parcial:
        return False
    return len(tokens & tokens_parcial) / min(len(tokens), len(tokens_parcial)) >= 0.5


def descripcion_especifica(valor):
    """Una descripcion identifica un producto solo si trae un modelo/cifra o es larga."""
    texto = texto_normalizado(valor)
    return bool(re.search(r"\d", texto)) or len(texto.split()) >= 6


def numero_de_item(valor):
    """Numero de linea ("Item 4", "FILA 16", "4"); None si falta o es un codigo UNSPSC (>= 7 digitos)."""
    digitos = re.findall(r"\d+", str(valor or ""))
    if len(digitos) != 1 or len(digitos[0]) >= 7:
        return None
    return int(digitos[0])


def misma_identidad_producto(producto_a, producto_b):
    if producto_a.get("categoria") != producto_b.get("categoria"):
        return False
    item_a = numero_de_item(producto_a.get("item"))
    item_b = numero_de_item(producto_b.get("item"))
    if item_a is not None and item_b is not None and item_a != item_b:
        return False  # lineas distintas de la misma oferta, aunque repitan modelo
    modelo_a = texto_normalizado(producto_a.get("modelo"))
    modelo_b = texto_normalizado(producto_b.get("modelo"))
    if len(modelo_a) >= 4 and modelo_a == modelo_b:
        return True
    if item_a is not None and item_a == item_b:
        # Dos anexos numeran sus items por separado: el numero solo no basta.
        return coincide_con_parcial(producto_a, producto_b)
    descripcion_a = texto_normalizado(producto_a.get("producto"))
    # "Notebook" o "All in One" sueltos pueden ser lotes distintos, no el mismo item.
    return descripcion_especifica(descripcion_a) and descripcion_a == texto_normalizado(producto_b.get("producto"))


TOTAL_OFERTA_MINIMO_CLP = 50_000


def total_oferta_creible(total_oferta):
    """El total del cuadro de ofertas a veces viene corrupto ("$ 4"): solo se usa como referencia
    si es un monto plausible para una oferta de equipos."""
    valor = normalizar_numero(total_oferta)
    return valor if isinstance(valor, (int, float)) and valor >= TOTAL_OFERTA_MINIMO_CLP else None


PRECIO_MAXIMO_POR_TIPO = {  # un unitario mayor suele ser el total de la linea leido como unitario
    "all-in-one": 6_000_000, "notebook": 8_000_000, "desktop": 9_000_000, "monitor": 3_500_000,
    "impresora": 9_000_000,
}
PRECIO_UNITARIO_MINIMO_CLP = 20_000
PRECIO_UNITARIO_MAXIMO_CLP = 20_000_000


def completar_fuentes_desde_parciales(productos, parciales):
    """Fija de que archivo salio cada dato usando los hallazgos parciales (que si conocen su archivo).

    El modelo, al consolidar, a veces abrevia o inventa nombres de archivo y paginas. Aqui se
    reemplaza toda fuente que no sea un archivo real de los hallazgos:
      - fuente del precio: el archivo cuyo parcial tiene el mismo precio unitario (economico primero);
      - fuente del producto/marca: el parcial de la misma identidad (con marca, tecnico primero).
    Si no hay una unica respuesta razonable, se deja vacio (no se adivina).
    """
    archivos = {str(a) for p in parciales
                for a in (p.get("archivo_fuente"), p.get("fuente_producto"), p.get("fuente_precio"))
                if a and a != "consolidacion"}
    if not archivos:
        return productos

    def elegir(candidatos):
        por_archivo = {}
        for c in candidatos:
            c = {**c, "archivo_fuente": c.get("archivo_fuente") or c.get("fuente_producto") or c.get("fuente_precio")}
            if c["archivo_fuente"]:
                por_archivo.setdefault(str(c["archivo_fuente"]), c)
        if len(por_archivo) == 1:
            return next(iter(por_archivo.values()))
        economicos = [c for a, c in por_archivo.items() if tipo_documental(Path(a).name) == "economico"]
        return economicos[0] if len(economicos) == 1 else None

    for producto in productos:
        if producto.get("fuente_precio") not in archivos:
            unitario = producto.get("precio_unitario")
            fuente = None
            if isinstance(unitario, (int, float)):
                fuente = elegir([
                    p for p in parciales
                    if p.get("archivo_fuente") in archivos and isinstance(p.get("precio_unitario"), (int, float))
                    and abs(p["precio_unitario"] - unitario) <= 1
                    and p.get("cantidad") in (None, "", producto.get("cantidad"))])
            producto["fuente_precio"] = fuente["archivo_fuente"] if fuente else None
            producto["pagina_precio"] = (fuente.get("pagina") or fuente.get("pagina_precio")) if fuente else None
        if producto.get("fuente_producto") not in archivos:
            candidatos = [p for p in parciales
                          if p.get("archivo_fuente") in archivos and coincide_con_parcial(producto, p)
                          and (not producto.get("marca") or not p.get("marca")
                               or texto_normalizado(p["marca"]) == texto_normalizado(producto["marca"]))]
            con_marca = [p for p in candidatos if p.get("marca")]
            tecnicos = [p for p in (con_marca or candidatos) if tipo_documental(Path(p["archivo_fuente"]).name) == "tecnico"]
            fuente = elegir(tecnicos or con_marca or candidatos)
            producto["fuente_producto"] = fuente["archivo_fuente"] if fuente else None
            producto["pagina_producto"] = (fuente.get("pagina") or fuente.get("pagina_producto")) if fuente else None
    return productos


FAMILIA_MARCA = {
    "probook": "HP", "elitebook": "HP", "zbook": "HP", "proone": "HP", "prodesk": "HP", "elitedesk": "HP",
    "eliteone": "HP", "pavilion": "HP", "laserjet": "HP", "officejet": "HP", "deskjet": "HP",
    "smart tank": "HP", "victus": "HP", "omen": "HP",
    "thinkpad": "Lenovo", "thinkbook": "Lenovo", "ideapad": "Lenovo", "thinkcentre": "Lenovo",
    "legion": "Lenovo", "loq": "Lenovo", "yoga": "Lenovo",
    "latitude": "Dell", "optiplex": "Dell", "vostro": "Dell", "inspiron": "Dell",
    "vivobook": "ASUS", "expertbook": "ASUS", "zenbook": "ASUS", "expertcenter": "ASUS",
    "macbook": "Apple", "imac": "Apple",
    "aspire": "Acer", "swift": "Acer", "nitro": "Acer", "travelmate": "Acer", "veriton": "Acer",
    "ecotank": "Epson", "workforce": "Epson", "pixma": "Canon",
}
MARCAS_EN_TEXTO = ("HP", "Lenovo", "Dell", "ASUS", "Acer", "Apple", "Samsung", "Epson", "Canon", "Brother",
                   "MSI", "LG", "Xerox", "Kyocera", "Ricoh", "Lexmark", "ViewSonic", "AOC", "Philips",
                   "Gigabyte", "Huawei", "Xiaomi", "Zebra", "Logitech")
PATRON_REQUISITO = re.compile(r"equivalent|similar|\btipo\b|o superior|compatible|segun bases", re.I)


def marca_desde_descripcion(texto):
    """Marca nombrada en la propia descripcion ("Impresora Epson L6270", "LAPTOP HP PROBOOK"). Devuelve
    None si hay mas de una marca o si el texto es un requisito ("equivalente a", "tipo 1")."""
    texto = str(texto or "")
    if not texto or PATRON_REQUISITO.search(texto):
        return None
    encontradas = set()
    for marca in MARCAS_EN_TEXTO:
        if re.search(r"(?<![A-Za-z0-9])" + re.escape(marca) + r"(?![A-Za-z0-9])", texto, re.I):
            encontradas.add(normalizar_marca(marca))
    for familia, marca in FAMILIA_MARCA.items():
        if re.search(r"(?<![A-Za-z0-9])" + re.escape(familia) + r"(?![A-Za-z0-9])", texto, re.I):
            encontradas.add(marca)
    return next(iter(encontradas)) if len(encontradas) == 1 else None


def completar_marcas_por_descripcion_y_tipo(productos, parciales):
    """1) marca nombrada en la descripcion; 2) si todas las fichas tecnicas del proveedor del mismo
    tipo de equipo son de una sola marca, esa marca. Siempre queda registrada la forma de obtenerla."""
    for producto in productos:
        if producto.get("marca"):
            continue
        marca = marca_desde_descripcion(f"{producto.get('producto') or ''} {producto.get('modelo') or ''}")
        if marca:
            producto["marca"] = marca
            producto["metodo_marca"] = "descripcion"
    sin_marca = [p for p in productos if not p.get("marca")]
    if not sin_marca:
        return productos
    marcas_por_tipo = {}
    for parcial in parciales:
        if not parcial.get("marca"):
            continue
        if tipo_documental(Path(str(parcial.get("archivo_fuente") or "")).name) == "economico":
            continue  # solo fichas tecnicas: el economico generico no prueba la marca
        tipo = subcategoria_producto(parcial)
        marcas_por_tipo.setdefault(tipo, set()).add(normalizar_marca(parcial["marca"]))
    for producto in sin_marca:
        tipo = subcategoria_producto(producto)
        posibles = set(marcas_por_tipo.get(tipo, set()))
        if tipo == "notebook":
            posibles |= marcas_por_tipo.get("workstation", set())  # "workstation movil" = notebook
        posibles.discard(None)
        if len(posibles) == 1 and tipo and not tipo.startswith("equipo"):
            producto["marca"] = next(iter(posibles))
            producto["metodo_marca"] = "tipo_con_marca_unica"
    return productos


def completar_marcas_por_nombre_de_archivo(productos, nombres_archivos):
    """Marca deducida de los NOMBRES de las fichas tecnicas del proveedor ("HP_ProOne_245_G10_AIO.pdf"):
    si todas las fichas de un tipo de equipo declaran una sola marca, se asigna a los productos sin marca
    de ese tipo (p.ej. "All in One Tipo 1" en un anexo economico generico)."""
    marcas_por_tipo = {}
    for nombre in nombres_archivos:
        if tipo_documental(Path(str(nombre)).name) != "tecnico":
            continue
        base = Path(str(nombre)).stem.replace("_", " ").replace("-", " ")
        marca = marca_desde_descripcion(base)
        tipo = subcategoria_producto({"producto": base, "categoria": "equipo"})
        if marca and tipo and not tipo.startswith("equipo"):
            marcas_por_tipo.setdefault(tipo, set()).add(marca)
    for producto in productos:
        if producto.get("marca"):
            continue
        tipo = subcategoria_producto(producto)
        posibles = set(marcas_por_tipo.get(tipo, set()))
        if len(posibles) == 1 and tipo and not tipo.startswith("equipo"):
            producto["marca"] = next(iter(posibles))
            producto["metodo_marca"] = "nombre_de_ficha_tecnica"
    return productos


PATRON_NOMBRE_GENERICO = (r"notebooks?|laptops?|port[aá]tiles?|computador(?:es|as?)?|equipos?(?:\s+computacionales)?|"
                          r"monitor(?:es)?|pantallas?|impresoras?|multifuncional(?:es)?|all\s*in\s*one|aio|"
                          r"desktops?|pcs?|estaciones?\s+de\s+trabajo|workstations?")
PATRON_FILA_AGREGADA = re.compile(r"^\s*(\d{1,5})\s+(?:unidades?\s+(?:de\s+)?|u\.?\s+)?(" + PATRON_NOMBRE_GENERICO + r")\b(.*)$",
                                  re.I)
PATRON_BLOQUE_CON_IVA = re.compile(r"con\s+iva|iva\s+incluido|impuestos?\s+incluidos?|valor\s+bruto", re.I)


def filas_agregadas(bloques):
    """Lineas de anexo economico con cantidad y un solo monto para un rubro generico, del tipo
    "52 NOTEBOOKS || $ 45.172.400" o "52 || Notebook || $ 45.172.400". No traen precio unitario: el
    unitario se deduce (total / cantidad), pasando a neto si el anexo declara valores con IVA."""
    filas = []
    for bloque in bloques:
        archivo = bloque.get("archivo") or ""
        if tipo_documental(Path(str(archivo)).name) != "economico":
            continue
        texto = str(bloque.get("texto") or "")
        con_iva = bool(PATRON_BLOQUE_CON_IVA.search(texto))
        for linea in texto.splitlines():
            limpia = re.sub(r"^\s*FILA\s+\d+:\s*", "", linea, flags=re.I)
            celdas = [c.strip() for c in limpia.split("||") if c.strip()]
            if len(celdas) < 2:
                continue
            montos = [normalizar_numero(c) for c in celdas[1:] if PATRON_CELDA_NUMERICA.fullmatch(c)]
            montos = [m for m in montos if isinstance(m, (int, float)) and m >= 50_000]
            if len(montos) != 1:
                continue  # con unitario y total la fila ya la cubre rescatar_filas_economicas
            cantidad, descripcion = None, None
            m = PATRON_FILA_AGREGADA.match(celdas[0])
            if m:
                cantidad, descripcion = int(m.group(1)), f"{m.group(2)}{m.group(3)}".strip()
            elif (len(celdas) >= 3 and re.fullmatch(r"\d{1,5}", celdas[0])
                  and re.match(r"\W*(?:" + PATRON_NOMBRE_GENERICO + r")\b", celdas[1], re.I)):
                cantidad, descripcion = int(celdas[0]), celdas[1]
            if not cantidad or cantidad <= 0:
                continue
            tipo = subcategoria_producto({"producto": descripcion, "categoria": "equipo"})
            if not tipo or tipo.startswith("equipo"):
                continue
            filas.append({"archivo": archivo, "pagina": bloque.get("pagina"), "cantidad": cantidad,
                          "tipo": tipo, "total": montos[0], "con_iva": con_iva, "texto": limpia[:200]})
    return filas


def cruzar_filas_agregadas(productos, bloques):
    """Cruce tecnico-economico para anexos que solo traen "N <rubro> || total": el equipo (marca y modelo) lo
    trae la ficha tecnica y el precio esta en la linea agregada. Se asigna solo si hay exactamente un producto
    sin precio de ese tipo y una sola linea agregada de ese tipo, con cantidad compatible."""
    filas = filas_agregadas(bloques)
    if not filas:
        return productos
    for tipo in {f["tipo"] for f in filas}:
        de_tipo = [f for f in filas if f["tipo"] == tipo]
        unicas = {(f["cantidad"], f["total"]) for f in de_tipo}
        if len(unicas) != 1:
            continue
        fila = de_tipo[0]
        candidatos = [p for p in productos
                      if subcategoria_producto(p) == tipo
                      and p.get("precio_unitario") in (None, "") and p.get("precio_total") in (None, "")
                      and p.get("cantidad") in (None, "", fila["cantidad"])]
        con_precio_mismo_tipo = [p for p in productos if subcategoria_producto(p) == tipo
                                 and isinstance(p.get("precio_unitario"), (int, float))]
        if len(candidatos) != 1 or con_precio_mismo_tipo:
            continue
        total = fila["total"] / TASA_IVA if fila["con_iva"] else fila["total"]
        unitario = round(total / fila["cantidad"])
        maximo = PRECIO_MAXIMO_POR_TIPO.get(tipo, PRECIO_UNITARIO_MAXIMO_CLP)
        if not PRECIO_UNITARIO_MINIMO_CLP <= unitario <= maximo:
            continue
        producto = candidatos[0]
        producto["cantidad"] = fila["cantidad"]
        producto["precio_unitario"] = unitario
        producto["precio_total"] = round(unitario * fila["cantidad"])
        producto["fuente_precio"] = fila["archivo"]
        producto["pagina_precio"] = fila["pagina"]
        producto["precio_desde_linea_agregada"] = True
        if fila["con_iva"]:
            producto["total_con_iva_convertido"] = True
        producto.setdefault("evidencia", fila["texto"])
    return productos


PATRON_CODIGO_GENERICO = re.compile(
    r"^(?:\d+(?:gb|tb|mb|hz|mhz|ghz|w|wh|mah|mm|cm|pulg)|ddr\d\w*|lpddr\d\w*|i[3579]-?\d+\w*|r[3579]-?\d+\w*|"
    r"w1[01]\w*|win1[01]\w*|usb\d?\w*|hdmi\d?\w*|wifi\d?\w*|wi-fi\d?\w*|pcie\d?\w*|gen\d+|\d+th|\d+(?:a|v)|"
    r"fhd\w*|uhd\w*|wuxga\w*|4k\w*|rtx\w*|gtx\w*|nvme\w*|m\.?2\w*)$", re.I)


def codigos_distintivos(producto):
    """Codigos de modelo que identifican al equipo y no a un componente generico (16GB, i5-1335U, DDR5...)."""
    return {c for c in codigos_de_modelo(producto) if len(c) >= 5 and not PATRON_CODIGO_GENERICO.match(c)}


def completar_marcas_por_linea(productos, bloques):
    """Marca escrita junto al codigo del modelo en los documentos del proveedor. El modelo a veces deja la marca
    vacia aunque la linea dice "Impresora Epson EcoTank L4260". Solo se asigna si en todas las lineas donde
    aparece el codigo hay una unica marca (y no es un requisito del tipo "similar a")."""
    sin_marca = [p for p in productos if not p.get("marca")]
    if not sin_marca or not bloques:
        return productos
    lineas_por_bloque = [(b.get("archivo"), b.get("pagina"), [l for l in str(b.get("texto") or "").splitlines() if l.strip()])
                         for b in bloques]
    for producto in sin_marca:
        codigos = codigos_distintivos(producto)
        if not codigos:
            continue
        marcas, fuente = set(), None
        for archivo, pagina, lineas in lineas_por_bloque:
            minusculas = [l.lower() for l in lineas]
            for i, linea in enumerate(minusculas):
                if not any(codigo in linea for codigo in codigos):
                    continue
                marca = marca_desde_descripcion(" ".join(lineas[max(0, i - 1):i + 2]))
                if marca:
                    marcas.add(marca)
                    fuente = fuente or (archivo, pagina)
        if len(marcas) == 1:
            producto["marca"] = next(iter(marcas))
            producto["metodo_marca"] = "linea_del_producto"
            if fuente and fuente[0]:
                producto["fuente_marca"], producto["pagina_marca"] = fuente
    return productos


def corregir_marca_por_familia(productos):
    """Una familia de producto pertenece a una sola marca (ThinkPad/LOQ=Lenovo, ProBook=HP...). Si el
    modelo declaro otra marca (p.ej. leyo el nombre del renglon vecino), manda la familia."""
    for producto in productos:
        texto = f"{producto.get('producto') or ''} {producto.get('modelo') or ''}"
        if PATRON_REQUISITO.search(texto):
            continue
        familias = {marca for familia, marca in FAMILIA_MARCA.items()
                    if re.search(r"(?<![A-Za-z0-9])" + re.escape(familia) + r"(?![A-Za-z0-9])", texto, re.I)}
        declarada = normalizar_marca(producto.get("marca"))
        if len(familias) == 1 and declarada and declarada not in familias:
            producto["marca_declarada_por_modelo"] = declarada
            producto["marca"] = next(iter(familias))
            producto["metodo_marca"] = "familia_del_modelo"
    return productos


TASA_IVA = 1.19


def convertir_totales_con_iva(productos):
    """Si total = cantidad x unitario x 1,19 el total viene con IVA: se deja el neto. Debe correr ANTES
    de deduplicar, para que la misma linea (una vez con total neto y otra con IVA) se fusione."""
    for producto in productos:
        cantidad, unitario, total = producto.get("cantidad"), producto.get("precio_unitario"), producto.get("precio_total")
        if all(isinstance(v, (int, float)) and v > 0 for v in (cantidad, unitario, total)):
            neto = cantidad * unitario
            if abs(total - neto) > max(2.0, neto * 0.01) and abs(total - neto * TASA_IVA) <= max(3.0, neto * 0.004):
                producto["precio_total"] = round(neto)
                producto["total_con_iva_convertido"] = True
    return productos


def aplicar_logica_iva(productos):
    """Totales con IVA y duplicados con IVA:
    - si total = cantidad x unitario x 1,19 el total viene con IVA: se deja el neto;
    - si dos lineas del mismo producto difieren exactamente por 1,19, la mas cara es la version con
      IVA (p.ej. un formulario que repite el precio): se descarta y su fuente queda como respaldo."""
    productos = convertir_totales_con_iva(productos)
    descartar = set()
    for i, a in enumerate(productos):
        for j, b in enumerate(productos):
            if i == j or i in descartar or j in descartar:
                continue
            ua, ub = a.get("precio_unitario"), b.get("precio_unitario")
            if not (isinstance(ua, (int, float)) and isinstance(ub, (int, float)) and ua > 0 and ub > ua):
                continue
            if abs(ub - ua * TASA_IVA) > max(2.0, ua * 0.004):
                continue
            if (a.get("categoria") == b.get("categoria")
                    and a.get("cantidad") in (None, "", b.get("cantidad"))):
                fuentes = set(a.get("fuentes_respaldo") or [])
                fuentes.update(str(b.get(c)) for c in ("archivo_fuente", "fuente_precio", "fuente_producto") if b.get(c))
                a["fuentes_respaldo"] = sorted(fuentes)
                a["precio_con_iva_confirmado_en_otra_fuente"] = True
                if b.get("marca") and not a.get("marca"):
                    a["marca"] = b["marca"]
                descartar.add(j)
    return [p for k, p in enumerate(productos) if k not in descartar]


def linea_coherente(producto):
    cantidad, unitario, total = producto.get("cantidad"), producto.get("precio_unitario"), producto.get("precio_total")
    return (all(isinstance(v, (int, float)) and v > 0 for v in (cantidad, unitario, total))
            and abs(cantidad * unitario - total) <= max(2.0, total * 0.001))


def lineas_distintas_coherentes(a, b):
    """Dos renglones del mismo anexo, cada uno con cantidad x unitario = total, son lineas distintas
    (el mismo equipo pedido en dos items), no fuentes en conflicto."""
    fuente_a = a.get("fuente_precio") or a.get("archivo_fuente")
    fuente_b = b.get("fuente_precio") or b.get("archivo_fuente")
    return bool(fuente_a) and fuente_a == fuente_b and linea_coherente(a) and linea_coherente(b)


def corregir_total_leido_como_unitario(productos, total_oferta):
    """Si cantidad x "total" coincide con el total de la oferta (neto o con IVA), el valor que el modelo
    tomo como total de la linea es en realidad el precio unitario (formularios con una sola columna de
    precio): se corrige unitario y total."""
    referencia = total_oferta_creible(total_oferta)
    if not referencia:
        return productos
    for producto in productos:
        cantidad, total = producto.get("cantidad"), producto.get("precio_total")
        if not (isinstance(cantidad, (int, float)) and cantidad > 1 and isinstance(total, (int, float)) and total > 0):
            continue
        for objetivo in (referencia, referencia / TASA_IVA):
            if abs(cantidad * total - objetivo) <= objetivo * 0.005:
                producto["precio_unitario"] = total
                producto["precio_total"] = round(cantidad * total)
                producto["total_leido_como_unitario"] = True
                break
    return productos


PATRON_DESCRIPCION_GENERICA = re.compile(
    r"\btipo\s*\d|\bl[ií]nea\s*\d|^(?:notebooks?|computador(?:es)?|impresoras?|monitor(?:es)?|all in one|equipos?)$",
    re.I)


def fusionar_lineas_genericas(productos):
    """La misma linea de precios puede aparecer en el cuadro de ofertas con una descripcion generica
    ("Notebook Tipo 1") y en la cotizacion con la descripcion real: mismos cantidad, unitario y total
    en archivos distintos -> una sola linea, con la descripcion mas especifica."""
    descartar = set()
    for i, a in enumerate(productos):
        if i in descartar or not isinstance(a.get("precio_unitario"), (int, float)):
            continue
        for j, b in enumerate(productos):
            if j <= i or j in descartar:
                continue
            if (a.get("cantidad") != b.get("cantidad") or a.get("precio_unitario") != b.get("precio_unitario")
                    or a.get("precio_total") != b.get("precio_total") or a.get("categoria") != b.get("categoria")):
                continue
            fuente_a = a.get("fuente_precio") or a.get("archivo_fuente")
            fuente_b = b.get("fuente_precio") or b.get("archivo_fuente")
            if not fuente_a or not fuente_b or fuente_a == fuente_b:
                continue  # dos renglones del mismo archivo son lineas distintas
            generica_a = bool(PATRON_DESCRIPCION_GENERICA.search(texto_normalizado(a.get("producto"))))
            generica_b = bool(PATRON_DESCRIPCION_GENERICA.search(texto_normalizado(b.get("producto"))))
            if generica_a == generica_b:
                # Ambas especificas: la misma linea escrita como requisito ("COMPUTADOR: PROCESADOR INTEL i5...")
                # en un formulario y como oferta real en la cotizacion. Mismos numeros en archivos distintos y
                # mismo tipo de equipo => una sola linea; se conserva la que nombra marca o familia.
                if generica_a or grupo_de_tipo(a) != grupo_de_tipo(b):
                    continue
                if not (a.get("cantidad") and a.get("cantidad") > 1 or a.get("precio_unitario", 0) >= 100_000):
                    continue
                puntaje_a = bool(a.get("marca")) + bool(marca_desde_descripcion(a.get("producto")))
                puntaje_b = bool(b.get("marca")) + bool(marca_desde_descripcion(b.get("producto")))
                generica_a, generica_b = (puntaje_a < puntaje_b), (puntaje_b < puntaje_a)
                if generica_a == generica_b:
                    generica_a = len(str(a.get("producto") or "")) > len(str(b.get("producto") or ""))
            bueno, malo, indice_malo = (b, a, i) if generica_a else (a, b, j)
            if not bueno.get("marca") and malo.get("marca"):
                bueno["marca"] = malo["marca"]
            fuentes = set(bueno.get("fuentes_respaldo") or [])
            fuentes.update(str(malo.get(c)) for c in ("archivo_fuente", "fuente_precio", "fuente_producto") if malo.get(c))
            bueno["fuentes_respaldo"] = sorted(fuentes)
            descartar.add(indice_malo)
            if indice_malo == i:
                break
    return [p for k, p in enumerate(productos) if k not in descartar]


def descartar_lineas_no_ofertadas(productos):
    """En un formulario economico con lineas "Tipo 1..4" o "Linea 1..4", las que el proveedor deja sin
    precio no se ofertaron. Si un archivo economico tiene al menos un producto con precio, sus lineas
    genericas sin precio ni total se descartan."""
    archivos_con_precio = {
        str(p.get("fuente_precio") or p.get("archivo_fuente")) for p in productos
        if isinstance(p.get("precio_unitario"), (int, float))}
    conservados = []
    for producto in productos:
        archivo = str(producto.get("fuente_precio") or producto.get("archivo_fuente") or "")
        sin_precio = producto.get("precio_unitario") in (None, "") and producto.get("precio_total") in (None, "")
        if (sin_precio and archivo in archivos_con_precio and tipo_documental(Path(archivo).name) == "economico"
                and PATRON_DESCRIPCION_GENERICA.search(texto_normalizado(producto.get("producto")))):
            continue
        conservados.append(producto)
    return conservados


def grupo_de_tipo(producto):
    """Desktop y workstation de sobremesa se describen indistintamente en los anexos: mismo grupo."""
    tipo = subcategoria_producto(producto)
    return "sobremesa" if tipo in ("desktop", "workstation") else tipo


def codigos_de_modelo(producto):
    """Codigos alfanumericos (letras y digitos, >= 4) que identifican un modelo: 15-FC0022LA, B7KF9AT..."""
    texto = f"{producto.get('producto') or ''} {producto.get('modelo') or ''}"
    return {t.lower() for t in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-/]{3,}", texto)
            if re.search(r"[A-Za-z]", t) and re.search(r"\d", t)}


def fusionar_fichas_tecnicas(productos):
    """Sin consolidacion por modelo, cada ficha tecnica queda como un producto sin precio ni cantidad.
    Aqui se fusionan por codigo: la ficha aporta marca/modelo/fuente a la fila con precio que comparte
    codigo de modelo (o descripcion muy parecida) y luego se descartan las fichas sin cantidad ni precio
    si el proveedor ya tiene productos con precio (siguen existiendo en `productos_parciales`)."""
    con_precio = [p for p in productos if isinstance(p.get("precio_unitario"), (int, float))]
    if not con_precio:
        return productos
    descartar = set()
    for indice, ficha in enumerate(productos):
        if ficha in con_precio or ficha.get("precio_total") not in (None, ""):
            continue
        if tipo_documental(Path(str(ficha.get("archivo_fuente") or ficha.get("fuente_producto") or "")).name) != "tecnico":
            continue
        codigos = codigos_de_modelo(ficha)
        tipo = subcategoria_producto(ficha)
        if ficha.get("cantidad") not in (None, ""):
            # Un anexo tecnico que declara cantidad ("solicitud de 10 AIO") solo se funde con la fila de precio
            # del mismo tipo y la misma cantidad, si es la unica; si no, se conserva.
            gemelas = [f for f in con_precio if f.get("cantidad") == ficha.get("cantidad")
                       and grupo_de_tipo(f) == grupo_de_tipo(ficha) and f.get("categoria") == ficha.get("categoria")
                       and not (texto_normalizado(f.get("marca")) and texto_normalizado(ficha.get("marca"))
                                and texto_normalizado(f.get("marca")) != texto_normalizado(ficha.get("marca")))]
            if len(gemelas) != 1:
                continue
            if not gemelas[0].get("marca") and ficha.get("marca") and "gener" not in str(ficha["marca"]).lower():
                gemelas[0]["marca"] = ficha["marca"]
                gemelas[0]["metodo_marca"] = "ficha_tecnica_por_tipo_y_cantidad"
            fuentes = set(gemelas[0].get("fuentes_respaldo") or [])
            fuentes.update(str(ficha.get(c)) for c in ("archivo_fuente", "fuente_producto") if ficha.get(c))
            gemelas[0]["fuentes_respaldo"] = sorted(fuentes)
            descartar.add(indice)
            continue
        mejores = []
        for fila in con_precio:
            if fila.get("categoria") != ficha.get("categoria"):
                continue
            marca_f, marca_p = texto_normalizado(ficha.get("marca")), texto_normalizado(fila.get("marca"))
            if marca_f and marca_p and marca_f != marca_p:
                continue
            comunes = codigos & codigos_de_modelo(fila)
            parecida = coincide_con_parcial(fila, ficha) and subcategoria_producto(fila) == tipo
            if comunes or (parecida and not codigos):
                mejores.append((len(comunes), fila))
        if mejores:
            mejores.sort(key=lambda par: par[0], reverse=True)
            fila = mejores[0][1]
            if not fila.get("marca") and ficha.get("marca") and (len(mejores) == 1 or mejores[0][0] > mejores[1][0]):
                fila["marca"] = ficha["marca"]
                fila["metodo_marca"] = "ficha_tecnica_por_codigo"
                fila["fuente_marca"] = ficha.get("archivo_fuente") or ficha.get("fuente_producto")
            if not fila.get("modelo") and ficha.get("modelo"):
                fila["modelo"] = ficha["modelo"]
            fuentes = set(fila.get("fuentes_respaldo") or [])
            fuentes.update(str(ficha.get(c)) for c in ("archivo_fuente", "fuente_producto") if ficha.get(c))
            fila["fuentes_respaldo"] = sorted(fuentes)
        descartar.add(indice)
    return [p for i, p in enumerate(productos) if i not in descartar]


def validar_productos(productos, parciales, total_oferta):
    for producto in (*productos, *parciales):
        producto["marca"] = normalizar_marca(producto.get("marca"))
    productos = convertir_totales_con_iva(filtrar_productos_relevantes(productos))
    productos = deduplicar_productos(productos)
    parciales = filtrar_productos_relevantes(parciales)
    productos = completar_fuentes_desde_parciales(productos, parciales)
    productos = fusionar_fichas_tecnicas(productos)
    productos = fusionar_lineas_genericas(productos)
    productos = descartar_lineas_no_ofertadas(productos)
    productos = aplicar_logica_iva(productos)
    productos = corregir_total_leido_como_unitario(productos, total_oferta)
    productos = completar_marcas_por_descripcion_y_tipo(productos, parciales)
    productos = corregir_marca_por_familia(productos)
    referencia_oferta = total_oferta_creible(total_oferta)
    for producto in productos:
        fuentes_declaradas = {
            str(producto.get(campo))
            for campo in ("archivo_fuente", "fuente_producto", "fuente_precio")
            if producto.get(campo) and producto.get(campo) != "consolidacion"
        }
        fuentes = {
            str(parcial.get("archivo_fuente"))
            for parcial in parciales
            if parcial.get("archivo_fuente")
            and (
                str(parcial.get("archivo_fuente")) in fuentes_declaradas
                or coincide_con_parcial(producto, parcial)
            )
        }
        producto["fuentes_respaldo"] = sorted(fuentes)
        alertas = []
        cantidad = producto.get("cantidad")
        unitario = producto.get("precio_unitario")
        total = producto.get("precio_total")
        if not producto.get("marca"):
            alertas.append("marca_no_identificada")
        if producto.get("precio_fuente_no_coincide"):
            alertas.append("precio_fuente_no_coincide")
        if cantidad in (None, ""):
            alertas.append("cantidad_no_identificada")
        elif float(cantidad) <= 0 or abs(float(cantidad) - round(float(cantidad))) > 0.001:
            alertas.append("cantidad_invalida")
        if unitario in (None, ""):
            alertas.append("precio_unitario_no_identificado")
        elif float(unitario) <= 0:
            alertas.append("precio_unitario_invalido")
        if total not in (None, "") and producto.get("precio_total_tipo") is None:
            alertas.append("origen_precio_total_no_identificado")
        elif total not in (None, "") and producto.get("precio_total_tipo") == "oferta":
            alertas.append("total_general_asignado_a_producto")
        en_pesos = str(producto.get("moneda") or "CLP").upper() == "CLP"
        if isinstance(unitario, (int, float)) and unitario > 0 and en_pesos:
            maximo_tipo = PRECIO_MAXIMO_POR_TIPO.get(subcategoria_producto(producto), PRECIO_UNITARIO_MAXIMO_CLP)
            if unitario < PRECIO_UNITARIO_MINIMO_CLP or unitario > min(PRECIO_UNITARIO_MAXIMO_CLP, maximo_tipo):
                alertas.append("precio_unitario_fuera_de_rango")
            if referencia_oferta and unitario > float(referencia_oferta) * 1.02:
                alertas.append("precio_unitario_supera_total_oferta")
            elif (referencia_oferta and len(productos) >= 2
                  and abs(unitario - float(referencia_oferta)) <= float(referencia_oferta) * 0.005):
                # Varios productos y uno cuesta exactamente toda la oferta: es el total general copiado.
                alertas.append("total_general_asignado_a_producto")
        if all(valor not in (None, "", 0) for valor in (cantidad, unitario, total)):
            calculado = float(cantidad) * float(unitario)
            if abs(calculado - float(total)) > max(2.0, abs(float(total)) * 0.01):
                alertas.append("cantidad_por_unitario_no_coincide_con_total")
        if producto.get("precio_sin_respaldo"):
            alertas.append("precio_no_aparece_en_documentos")
        if not fuentes:
            alertas.append("sin_respaldo_en_hallazgos_parciales")
        if str(producto.get("confianza") or "").lower() == "baja":
            alertas.append("confianza_modelo_baja")
        producto["alertas"] = alertas

    # El modelo a veces copia el precio de un item en otro de distinta categoria.
    for indice, producto in enumerate(productos):
        unitario = producto.get("precio_unitario")
        if not isinstance(unitario, (int, float)) or unitario < 100_000:
            continue
        for otro in productos[indice + 1:]:
            if (otro.get("precio_unitario") == unitario
                    and otro.get("categoria") != producto.get("categoria")):
                for afectado in (producto, otro):
                    afectado["alertas"].append("precio_repetido_en_otra_categoria")

    for indice, producto in enumerate(productos):
        for otro in productos[indice + 1:]:
            if not misma_identidad_producto(producto, otro):
                continue
            if lineas_distintas_coherentes(producto, otro):
                continue
            campos_conflictivos = [
                campo for campo in ("cantidad", "precio_unitario", "precio_total")
                if producto.get(campo) not in (None, "")
                and otro.get(campo) not in (None, "")
                and not numeros_iguales(producto[campo], otro[campo])
            ]
            for campo in campos_conflictivos:
                alerta = f"conflicto_entre_fuentes_{campo}"
                producto["alertas"].append(alerta)
                otro["alertas"].append(alerta)

    total_referencia = total_oferta_creible(total_oferta)
    if total_referencia and len(productos) >= 3:
        productos[:] = [
            p for p in productos
            if not (PATRON_DESCRIPCION_GENERICA.search(texto_normalizado(p.get("producto")))
                    and isinstance(p.get("precio_total"), (int, float)) and p["precio_total"] > 0
                    and any(abs(p["precio_total"] - objetivo) <= objetivo * 0.005
                            for objetivo in (total_referencia, total_referencia / TASA_IVA)))]
    totales = [producto.get("precio_total") for producto in productos]
    if (total_referencia and totales and all(isinstance(total, (int, float)) and total > 0 for total in totales)
            and all(linea_coherente(p) for p in productos)):
        # El total del cuadro de ofertas es neto. Si la suma de las lineas es exactamente ese total x 1,19,
        # las lineas se leyeron con IVA incluido (la oferta lo informo asi): se pasan a neto en vez de marcar
        # todo el proveedor como inconsistente.
        suma_lineas = sum(totales)
        con_iva = float(total_referencia) * TASA_IVA
        if abs(suma_lineas - con_iva) <= con_iva * 0.01:
            for producto in productos:
                producto["precio_total"] = round(producto["precio_total"] / TASA_IVA)
                producto["precio_unitario"] = round(producto["precio_unitario"] / TASA_IVA)
                producto["precios_convertidos_a_neto_por_total_oferta"] = True
            totales = [producto.get("precio_total") for producto in productos]
    if total_referencia and totales and all(isinstance(total, (int, float)) for total in totales):
        suma = sum(totales)
        if suma > float(total_referencia) * 1.02:
            sospechosos = [p for p in productos if any(
                clave in alerta for alerta in p["alertas"]
                for clave in ("invalida", "fuera_de_rango", "no_aparece", "repetido_en_otra", "no_coincide",
                              "conflicto_entre_fuentes", "total_general"))]
            # Un precio sacado de una ficha tecnica (impresion de una web) mientras otros precios vienen de
            # un anexo economico es el candidato natural a explicar el exceso.
            def precio_de_ficha(p):
                archivo = str(p.get("fuente_precio") or p.get("archivo_fuente") or "")
                return (isinstance(p.get("precio_unitario"), (int, float))
                        and tipo_documental(Path(archivo).name) == "tecnico")
            hay_economicos = any(isinstance(p.get("precio_unitario"), (int, float)) and not precio_de_ficha(p)
                                 for p in productos)
            if hay_economicos:
                sospechosos = sospechosos + [p for p in productos if precio_de_ficha(p) and p not in sospechosos]
            # Si ya hay productos con errores propios, ellos explican el exceso: no se castiga a los demas.
            for producto in (sospechosos or productos):
                producto["alertas"].append("suma_productos_supera_total_oferta")

    if total_referencia and len(productos) >= 3:
        for producto in productos:
            total = producto.get("precio_total")
            if (isinstance(total, (int, float)) and total > 0
                    and any(abs(total - objetivo) <= objetivo * 0.005 for objetivo in (total_referencia, total_referencia / TASA_IVA))
                    and "total_general_asignado_a_producto" not in producto["alertas"]):
                producto["alertas"].append("total_general_asignado_a_producto")

    for producto in productos:
        alertas = producto["alertas"]
        if any(
            "invalida" in alerta
            or "fuera_de_rango" in alerta
            or "no_aparece" in alerta
            or "repetido_en_otra" in alerta
            or "no_coincide" in alerta
            or "supera_total" in alerta
            or "conflicto_entre_fuentes" in alerta
            or "total_general" in alerta
            for alerta in alertas
        ):
            producto["estado_validacion"] = "inconsistente"
        elif any("no_identificad" in alerta for alerta in alertas):
            producto["estado_validacion"] = "incompleto"
        elif alertas:
            producto["estado_validacion"] = "revisar"
        else:
            producto["estado_validacion"] = "ok"
    return productos


PATRON_EQUIPO_ESTRICTO = re.compile(
    r"notebook|laptop|port[aá]til|computador|all.?in.?one|\baio\b|desktop|workstation|monitor|impresora|multifuncional", re.I)
PATRON_MARCA_O_FAMILIA = re.compile(
    r"(?<![A-Za-z0-9])(?:hp|lenovo|dell|asus|acer|apple|samsung|epson|canon|brother|msi|lg|xerox|kyocera|ricoh|"
    r"lexmark|viewsonic|aoc|philips|gigabyte|huawei|xiaomi|zebra|thinkstation|workstation|precision|"
    r"estaci[oó]n de trabajo)(?![A-Za-z0-9])", re.I)
PATRON_NOMBRE_HISTORIAL = re.compile(
    r"experiencia|comprobante|curr[i\u00ed]culum|referencias?|orden[_ .-]?de[_ .-]?compra|(?<![a-z])ocs?(?![a-z])|"
    r"ilovepdf|\d{3,}-\d{1,4}-(?:se|ag|cm|co)\d{2}(?!\d)", re.I)
PATRON_CONTENIDO_OC = re.compile(
    r"nombre\s+orden\s+de\s+compra|n[u\u00fa]mero\s+de\s+orden\s+de\s+compra|orden\s+de\s+compra\s*(?:n\W|:)", re.I)


def documento_de_historial(relativo, texto):
    """Experiencia del oferente: ordenes de compra de ventas pasadas, curriculum, comprobantes. Contienen
    equipos y precios reales, pero NO son lo que se oferta en esta licitacion. Se detecta por la ruta
    completa (carpetas dentro de zip/rar) o por el contenido."""
    return bool(PATRON_NOMBRE_HISTORIAL.search(str(relativo)) or PATRON_CONTENIDO_OC.search(texto[:4000]))


PATRON_ADMIN_RUIDO = re.compile(
    r"carnet|c[eé]dula|(?<![a-z])rut(?![a-z])|poder|escritura|patente|tributari|certificad|declaraci|"
    r"integridad|pacto|constituci[oó]n|f[-_ ]?30|antecedentes|inhabilidad|boleta|seriedad|firma|sello|"
    r"vigencia|anotaciones|domicilio|resoluci[oó]n", re.I)
PATRON_ADMIN_POSITIVO = re.compile(
    r"cotiz|econ[oó]m|oferta|precio|ficha|t[eé]cnic|especific|(?<![a-z])ft(?![a-z])|catalogo", re.I)


def es_ruido_administrativo(relativo):
    """Documento del anexo administrativo que claramente no es la oferta (carnet, RUT, certificados,
    declaraciones, programas de integridad...). No vale la pena hacerle OCR ni quedara "pendiente"."""
    nombre = str(relativo).replace("\\", "/").split("/")[-1]
    if "administrativo__" not in str(relativo).replace("\\", "/"):
        return False
    return bool(PATRON_ADMIN_RUIDO.search(nombre)) and not PATRON_ADMIN_POSITIVO.search(nombre)


PATRON_NOMBRE_RUIDO = re.compile(
    r"iso[_ .-]?\d|certificad|pacto|integridad|reciclaj|etiqueta|p[oó]liza|garant[ií]a|declaraci[oó]n|"
    r"experiencia|orden[_ .-]?de[_ .-]?compra|energy[_ .-]?star|epeat|blue[_ .-]?angel|gu[ií]a[_ .-]?de[_ .-]?compra|foro",
    re.I)


def anexo_tecnico_sin_equipos(texto, nombre=""):
    """True si un anexo tecnico no puede aportar equipos: el texto no menciona ningun equipo del
    alcance, o su nombre indica un certificado/politica y no declara marca ni modelo. Evita gastar
    llamadas del modelo en certificados ISO, pactos de integridad o documentos vacios."""
    if not texto.strip():
        return True
    muestra = texto[:80000]
    if not PATRON_PALABRA_PRODUCTO.search(muestra) and not PATRON_MARCA_O_FAMILIA.search(muestra):
        return True  # ni un equipo, ni una marca/familia conocida: no hay nada que extraer
    if PATRON_NOMBRE_HISTORIAL.search(nombre):
        return True  # experiencia/ordenes de compra pasadas: son ventas anteriores, no lo ofertado ahora
    if not PATRON_NOMBRE_RUIDO.search(nombre):
        return False
    if PATRON_IDENTIDAD_TECNICA.search(muestra):
        return False
    # Un "anexo de garantias" puede listar los equipos ofertados en una tabla: se conserva.
    return not any(linea.startswith("FILA") and PATRON_EQUIPO_ESTRICTO.search(linea)
                   for linea in muestra.splitlines())


def extraer_parciales_archivo(path, texto, proveedor, rut, total_oferta, args):
    relativo = str(path)
    tipo_doc = tipo_documental(path.name)
    fragmentos = dividir_texto(compactar_texto(texto), args.max_chars_archivo)
    productos = []
    estados = []
    respuestas = []
    prompts = []
    observaciones = []
    consumos = []
    limite_tecnico = getattr(args, "max_fragmentos_tecnico", 3)
    if tipo_doc != "economico" and limite_tecnico and len(fragmentos) > limite_tecnico:
        # Manuales, rankings o catalogos enormes: la marca y el modelo estan al inicio y el precio viene del
        # anexo economico (que nunca se recorta); leer 13 fragmentos solo gasta tiempo del modelo.
        observaciones.append(f"anexo tecnico largo: solo se leyeron los primeros {limite_tecnico} de {len(fragmentos)} fragmentos")
        fragmentos = fragmentos[:limite_tecnico]
    for numero, fragmento in enumerate(fragmentos, 1):
        productos_pista, precios_pista = pistas(fragmento)
        texto_fragmento = f"[FRAGMENTO {numero} DE {len(fragmentos)}]\n{fragmento}"
        prompt = (PROMPT_ARCHIVO_COMPACTO if getattr(args, "prompt_compacto", False) else PROMPT_ARCHIVO).format(
            proveedor=proveedor,
            rut=rut,
            archivo=relativo,
            tipo_documental=tipo_doc,
            total_oferta=total_oferta or "no informado",
            pistas_producto="\n".join(productos_pista) or "(ninguna)",
            pistas_precio="\n".join(precios_pista) or "(ninguna)",
            texto=texto_fragmento
        )
        datos, estado, cruda, consumo = consultar_modelo(
            prompt, args.modelo, args.timeout, args.num_ctx, args.backend, args.max_tokens,
            args.reintentos_modelo, args.espera_reintento
        )
        estados.append(estado)
        respuestas.append(f"[FRAGMENTO {numero}]\n{cruda}")
        prompts.append(prompt)
        consumos.append(consumo)
        if isinstance(datos, dict):
            if datos.get("observaciones"):
                observaciones.append(str(datos["observaciones"]))
            for producto in datos.get("productos", []):
                if isinstance(producto, dict):
                    if getattr(args, "prompt_compacto", False) and producto.get("precio_total") not in (None, ""):
                        producto.setdefault("precio_total_tipo", "linea")
                    limpio = normalizar_producto(producto, proveedor, rut, relativo, tipo_doc)
                    if limpio:
                        productos.append(limpio)
        if args.pausa_archivo > 0:
            time.sleep(args.pausa_archivo)
    estado_final = "ok" if all(estado in ESTADOS_IA_OK for estado in estados) else "; ".join(estados)
    return (
        productos,
        estado_final,
        "\n\n".join(respuestas),
        "\n\n".join(prompts),
        "; ".join(observaciones) or None,
        combinar_consumos(consumos, args.modelo),
    )


# ---------------------------------------------------------------------------
# MODO VISION (--vision)
# ---------------------------------------------------------------------------
def imagen_desde_pixmap(pixmap):
    """JPEG calidad 85: una pagina escaneada pesa ~6 veces menos que en PNG y el modelo
    la lee igual. PNG solo si la version de PyMuPDF no soporta JPEG."""
    try:
        return {"mime": "image/jpeg", "base64": base64.b64encode(pixmap.tobytes("jpg", jpg_quality=85)).decode("ascii")}
    except Exception:
        return {"mime": "image/png", "base64": base64.b64encode(pixmap.tobytes("png")).decode("ascii")}


def renderizar_paginas_pdf(path, indices, dpi):
    """Devuelve {indice: {"mime", "base64"}} con las paginas pedidas como imagen JPEG."""
    imagenes = {}
    escala = dpi / 72
    with FITZ_LOCK:
        doc = fitz.open(ruta_sistema(path))
        try:
            for index in indices:
                if 0 <= index < len(doc):
                    pixmap = doc[index].get_pixmap(matrix=fitz.Matrix(escala, escala), alpha=False)
                    imagenes[index] = imagen_desde_pixmap(pixmap)
        finally:
            doc.close()
    return imagenes


def cargar_imagen_archivo(path, lado_maximo=1600):
    """Carga un JPG/PNG/WEBP adjunto. Reduce fotos muy grandes para no gastar contexto."""
    with open(ruta_sistema(path), "rb") as archivo:
        datos = archivo.read()
    mime = EXTENSIONES_IMAGEN[path.suffix.lower()]
    if Image is not None:
        try:
            with Image.open(io.BytesIO(datos)) as original:
                if max(original.size) > lado_maximo or mime == "image/webp":
                    imagen = original.convert("RGB")
                    imagen.thumbnail((lado_maximo, lado_maximo))
                    buffer = io.BytesIO()
                    imagen.save(buffer, format="JPEG", quality=85)
                    datos, mime = buffer.getvalue(), "image/jpeg"
        except Exception:
            pass
    return {"mime": mime, "base64": base64.b64encode(datos).decode("ascii")}


def sumar_consumos(consumos, modelo):
    """Como combinar_consumos, pero suma llamadas de consumos ya combinados."""
    consumos = [consumo for consumo in consumos if consumo]
    return {
        "modelo": next((fila.get("modelo") for fila in consumos if fila.get("modelo")), modelo),
        "prompt_tokens": sum(fila.get("prompt_tokens") or 0 for fila in consumos),
        "completion_tokens": sum(fila.get("completion_tokens") or 0 for fila in consumos),
        "total_tokens": sum(fila.get("total_tokens") or 0 for fila in consumos),
        "llamadas": sum(int(fila.get("llamadas") or 0) for fila in consumos),
    }


def extraer_parciales_vision(path, paginas, proveedor, rut, total_oferta, args):
    """Envia paginas como imagen (+ su texto) en grupos de --vision-paginas-por-llamada.

    paginas: lista de dicts {"pagina": numero_1_based, "imagen": {...}, "texto": str}
    Devuelve la misma tupla que extraer_parciales_archivo.
    """
    relativo = str(path)
    tipo_doc = tipo_documental(path.name)
    por_llamada = max(1, args.vision_paginas_por_llamada)
    grupos = [paginas[inicio:inicio + por_llamada] for inicio in range(0, len(paginas), por_llamada)]
    productos = []
    estados = []
    respuestas = []
    prompts = []
    observaciones = []
    consumos = []
    for numero, grupo in enumerate(grupos, 1):
        numeros = ", ".join(str(pagina["pagina"]) for pagina in grupo)
        textos = "\n\n".join(pagina["texto"] for pagina in grupo if pagina["texto"].strip())
        texto_prompt = (
            f"{REGLAS_VISION}\n\n"
            f"[GRUPO {numero} DE {len(grupos)}] IMAGENES ADJUNTAS, EN ESTE ORDEN: paginas {numeros}\n"
            "TEXTO EXTRAIDO AUTOMATICAMENTE DE ESAS PAGINAS "
            "(puede venir desordenado; usalo solo para confirmar cifras y nombres):\n"
            f"{textos or '(sin texto extraible: paginas escaneadas o imagen)'}"
        )
        prompt = PROMPT_ARCHIVO.format(
            proveedor=proveedor,
            rut=rut,
            archivo=relativo,
            tipo_documental=tipo_doc,
            total_oferta=total_oferta or "no informado",
            pistas_producto="(usar las imagenes adjuntas)",
            pistas_precio="(usar las imagenes adjuntas)",
            texto=texto_prompt,
        )
        datos, estado, cruda, consumo = consultar_modelo(
            prompt, args.modelo, args.timeout, args.num_ctx, args.backend, args.max_tokens,
            args.reintentos_modelo, args.espera_reintento,
            imagenes=[pagina["imagen"] for pagina in grupo],
        )
        estados.append(estado)
        respuestas.append(f"[VISION GRUPO {numero} - paginas {numeros}]\n{cruda}")
        prompts.append(f"[VISION - {len(grupo)} imagen(es): paginas {numeros}]\n{prompt}")
        consumos.append(consumo)
        if isinstance(datos, dict):
            if datos.get("observaciones"):
                observaciones.append(str(datos["observaciones"]))
            for producto in datos.get("productos", []):
                if isinstance(producto, dict):
                    limpio = normalizar_producto(producto, proveedor, rut, relativo, tipo_doc)
                    if limpio:
                        limpio["lectura"] = "vision"
                        productos.append(limpio)
        if args.pausa_archivo > 0:
            time.sleep(args.pausa_archivo)
    estado_final = "ok" if all(estado in ESTADOS_IA_OK for estado in estados) else "; ".join(estados)
    return (
        productos,
        estado_final,
        "\n\n".join(respuestas),
        "\n\n".join(prompts),
        "; ".join(observaciones) or None,
        combinar_consumos(consumos, args.modelo),
    )


def extraer_parciales_mixto(path, relativo, detalle, proveedor, rut, total_oferta, args):
    """Archivo con paginas para vision: texto normal para las paginas faciles
    y vision para las dificiles. Para imagenes sueltas (JPG/PNG) solo vision."""
    if path.suffix.lower() in EXTENSIONES_IMAGEN:
        paginas = [{"pagina": 1, "imagen": cargar_imagen_archivo(path), "texto": ""}]
        texto_resto = ""
    else:
        texto_paginas = detalle.get("texto_paginas", {})
        indices_vision = sorted(detalle.get("paginas_vision", {}))
        imagenes = renderizar_paginas_pdf(path, indices_vision, args.dpi_vision)
        paginas = [
            {"pagina": index + 1, "imagen": imagenes[index], "texto": texto_paginas.get(index, "")}
            for index in indices_vision if index in imagenes
        ]
        texto_resto = "\n".join(
            texto_paginas[index] for index in sorted(texto_paginas) if index not in set(indices_vision)
        ).strip()

    resultados = []
    if len(texto_resto) > 40:
        resultados.append(extraer_parciales_archivo(Path(relativo), texto_resto, proveedor, rut, total_oferta, args))
    if paginas:
        resultados.append(extraer_parciales_vision(Path(relativo), paginas, proveedor, rut, total_oferta, args))
    if not resultados:
        return [], "no_ejecutada", "", "", None, {}

    productos = [producto for resultado in resultados for producto in resultado[0]]
    estados = [resultado[1] for resultado in resultados]
    estado_final = "ok" if all(estado in ESTADOS_IA_OK for estado in estados) else "; ".join(estados)
    observaciones = "; ".join(resultado[4] for resultado in resultados if resultado[4]) or None
    return (
        productos,
        estado_final,
        "\n\n".join(resultado[2] for resultado in resultados),
        "\n\n".join(resultado[3] for resultado in resultados),
        observaciones,
        sumar_consumos([resultado[5] for resultado in resultados], args.modelo),
    )


def dividir_registros_json(registros, limite):
    lotes = []
    actual = []
    largo = 2
    for registro in registros:
        serializado = json.dumps(registro, ensure_ascii=False, separators=(",", ":"))
        adicional = len(serializado) + (1 if actual else 0)
        if actual and largo + adicional > limite:
            lotes.append(actual)
            actual = []
            largo = 2
        actual.append(registro)
        largo += adicional
    if actual:
        lotes.append(actual)
    return lotes


def normalizar_consolidados(datos, respaldo, proveedor, rut):
    consolidados = []
    archivos_reales = sorted({
        str(archivo) for p in respaldo
        for archivo in (p.get("archivo_fuente"), p.get("fuente_producto"), p.get("fuente_precio"),
                        *(p.get("fuentes_respaldo") or []))
        if archivo and archivo != "consolidacion"})
    if isinstance(datos, dict) and isinstance(datos.get("productos"), list):
        for producto in datos["productos"]:
            if not isinstance(producto, dict):
                continue
            limpio = normalizar_producto(producto, proveedor, rut, "consolidacion", "consolidado")
            if limpio:
                # Solo se aceptan archivos que existen entre los hallazgos parciales: el modelo
                # puede abreviar o inventar nombres y la trazabilidad no puede depender de ello.
                limpio["fuente_producto"] = archivo_declarado(producto.get("fuente_producto"), archivos_reales)
                limpio["fuente_precio"] = archivo_declarado(producto.get("fuente_precio"), archivos_reales)
                consolidados.append(limpio)
    return consolidados or respaldo


def consolidar_parciales(parciales, proveedor, rut, total_oferta, args):
    respaldo_marcas = [dict(producto) for producto in parciales]
    parciales = deduplicar_productos(filtrar_productos_relevantes(parciales))
    if not parciales or args.sin_consolidar:
        parciales = completar_marcas_desde_parciales(parciales, respaldo_marcas)
        return validar_productos(parciales, parciales, total_oferta), "omitida", "", "", {}

    lotes = dividir_registros_json(parciales, args.max_chars_consolidacion)
    consolidados = []
    estados = []
    respuestas = []
    prompts = []
    consumos = []
    for numero, lote in enumerate(lotes, 1):
        resumen = json.dumps(lote, ensure_ascii=False, separators=(",", ":"))
        prompt = PROMPT_CONSOLIDAR.format(
            proveedor=proveedor,
            rut=rut,
            total_oferta=total_oferta or "no informado",
            parciales=f"LOTE {numero} DE {len(lotes)}\n{resumen}"
        )
        datos, estado, cruda, consumo = consultar_modelo(
            prompt, args.modelo, args.timeout, args.num_ctx, args.backend, args.max_tokens,
            args.reintentos_modelo, args.espera_reintento
        )
        consolidados.extend(normalizar_consolidados(datos, lote, proveedor, rut))
        estados.append(estado)
        respuestas.append(f"[LOTE {numero}]\n{cruda}")
        prompts.append(prompt)
        consumos.append(consumo)
        if args.pausa_archivo > 0:
            time.sleep(args.pausa_archivo)

    consolidados = deduplicar_productos(filtrar_productos_relevantes(consolidados))
    if len(lotes) > 1:
        resumen_final = json.dumps(consolidados, ensure_ascii=False, separators=(",", ":"))
        if len(resumen_final) <= args.max_chars_consolidacion:
            prompt = PROMPT_CONSOLIDAR.format(
                proveedor=proveedor,
                rut=rut,
                total_oferta=total_oferta or "no informado",
                parciales=f"CONSOLIDACION FINAL DE {len(lotes)} LOTES\n{resumen_final}"
            )
            datos, estado, cruda, consumo = consultar_modelo(
                prompt, args.modelo, args.timeout, args.num_ctx, args.backend, args.max_tokens,
                args.reintentos_modelo, args.espera_reintento
            )
            consolidados = normalizar_consolidados(datos, consolidados, proveedor, rut)
            estados.append(estado)
            respuestas.append(f"[CONSOLIDACION FINAL]\n{cruda}")
            prompts.append(prompt)
            consumos.append(consumo)

    estado_final = "ok" if all(estado in ESTADOS_IA_OK for estado in estados) else "; ".join(estados)
    consolidados = completar_marcas_desde_parciales(consolidados, respaldo_marcas)
    return (
        validar_productos(consolidados, parciales, total_oferta),
        estado_final,
        "\n\n".join(respuestas),
        "\n\n".join(prompts),
        combinar_consumos(consumos, args.modelo),
    )


def reemplazar_resultado(resultados, nuevo):
    """Reemplaza en la lista el resultado del mismo RUT (o lo agrega si no existe)."""
    rut = str(nuevo.get("rut") or "").strip()
    for indice, resultado in enumerate(resultados):
        if str(resultado.get("rut") or "").strip() == rut:
            resultados[indice] = nuevo
            return
    resultados.append(nuevo)


def listar_archivos_oferta(carpeta):
    """Enumera tambien archivos cuyo nombre completo supera los 260 caracteres."""
    carpeta = Path(carpeta)
    raiz_fisica = ruta_sistema(carpeta)
    archivos = []
    def error_en_recorrido(error):
        raise error

    for directorio, subdirectorios, nombres in os.walk(raiz_fisica, onerror=error_en_recorrido):
        subdirectorios[:] = sorted(d for d in subdirectorios if d != "_debug_ia")
        relativa = Path(os.path.relpath(directorio, raiz_fisica))
        for nombre in sorted(nombres):
            if nombre in {"oferta.json", "extraccion_ia.json", ".extraido_ok", "desktop.ini"}:
                continue
            path = carpeta / relativa / nombre
            try:
                if os.path.isfile(ruta_sistema(path)):
                    archivos.append(path)
                elif not os.path.exists(ruta_sistema(path)):
                    raise OSError(f"Archivo encontrado pero no accesible: {path}")
            except OSError as exc:
                raise OSError(f"No se pudo enumerar {path}: {exc}") from exc
    return archivos


# ---------------------------------------------------------------------------
# REPROCESO SELECTIVO: solo los archivos que lo necesitan
# ---------------------------------------------------------------------------
def archivo_fallo_modelo(registro):
    estado = str(registro.get("estado_ia") or "")
    return estado not in ESTADOS_IA_NEUTROS and estado not in ESTADOS_IA_OK


def archivo_pendiente_ocr(registro, previo):
    # Si el OCR ya se intento (y Tesseract funcionaba) el archivo no tiene texto recuperable: rehacerlo en cada
    # corrida solo gasta tiempo. Se reintenta cuando no hubo OCR (checkpoints viejos o Tesseract ausente).
    return registro.get("archivo") in set(previo.get("necesita_ocr") or []) and not previo.get("ocr_intentado")


def revision_necesaria(resultado, carpeta, args):
    """Chequeo rapido (sin leer PDFs) de si un proveedor ya procesado puede
    necesitar trabajo. El detalle de QUE archivos rehacer lo decide
    planificar_reproceso dentro del hilo de trabajo."""
    if resultado_requiere_reproceso(resultado, args.ocr, args.modelo, args.backend, args.vision,
                                    modo_extraccion(args), args.vision_solo_escaneadas):
        return True
    # Archivos que aparecieron despues (ej. anexos tecnicos descargados mas tarde).
    registrados = {registro.get("archivo") for registro in resultado.get("resultados_archivos", [])}
    return any(
        str(path.relative_to(carpeta)) not in registrados
        for path in listar_archivos_oferta(carpeta)
    )


def planificar_reproceso(carpeta, previo, args, seven_zip):
    """Decide que hacer con un proveedor ya procesado.

    Devuelve (modo, archivos):
      - ("completo", None): cambio de modelo/backend; se rehace todo.
      - ("parcial", {archivos}): se rehacen solo esos archivos y se reutiliza el
        resto; luego se vuelve a consolidar el proveedor.
      - ("consolidar", set()): solo fallo la consolidacion.
      - ("nada", set()): no hay nada que rehacer.
    """
    cambio_modelo = bool(previo.get("modelo") and str(previo.get("modelo")) != str(args.modelo))
    cambio_backend = bool(previo.get("backend") and str(previo.get("backend")) != str(args.backend))
    cambio_modo = (previo.get("modo") or "por_archivo") != modo_extraccion(args)
    if cambio_modelo or cambio_backend or cambio_modo:
        return "completo", None

    descomprimir(carpeta, seven_zip)
    previos = {registro.get("archivo"): registro for registro in previo.get("resultados_archivos", [])}
    revisar_vision = args.vision and not previo.get("vision") and not args.vision_solo_escaneadas
    rehacer = {}
    for path in listar_archivos_oferta(carpeta):
        verificar_detencion()
        relativo = str(path.relative_to(carpeta))
        registro = previos.get(relativo)
        extension = path.suffix.lower()
        if registro is None:
            rehacer[relativo] = "archivo_nuevo"
        elif archivo_fallo_modelo(registro):
            rehacer[relativo] = "fallo_modelo"
        elif (args.ocr or args.vision) and archivo_pendiente_ocr(registro, previo):
            rehacer[relativo] = "pendiente_ocr"
        elif revisar_vision and not args.solo_texto:
            if extension in EXTENSIONES_IMAGEN:
                rehacer[relativo] = "imagen"
            elif extension == ".pdf" and registro.get("estado_lectura") != "texto_vision":
                # Lectura rapida sin modelo: solo clasifica las paginas.
                detalle = {}
                extraer_archivo(path, args, detalle)
                if detalle.get("paginas_vision"):
                    rehacer[relativo] = "paginas_para_vision"

    if rehacer:
        return "parcial", rehacer
    estado_consolidacion = str(previo.get("estado_consolidacion") or "")
    if estado_consolidacion not in ESTADOS_IA_OK | {"", "omitida", "no_ejecutada"}:
        return "consolidar", {}
    return "nada", {}


def modo_extraccion(args):
    return "por_proveedor" if getattr(args, "por_proveedor", False) else "por_archivo"


# ---------------------------------------------------------------------------
# MODO POR PROVEEDOR (--por-proveedor)
# ---------------------------------------------------------------------------
PATRON_PALABRA_PRODUCTO = re.compile(
    PATRON_EQUIPO_RELEVANTE.pattern + r"|monitor|impresora|multifuncional|all.?in.?one", re.I)
PATRON_MARCA_PAGINA = re.compile(r"\b(?:" + "|".join(re.escape(m) for m in MARCAS) + r")\b", re.I)
PATRON_IDENTIDAD_TECNICA = re.compile(r"\b(?:marca|modelo)\s*(?::|\||=)|\b(?:brand|model)\s*:", re.I)
PATRON_ESPECIFICACION_TECNICA = re.compile(
    r"procesador|\bram\b|\bssd\b|\bhdd\b|resoluci[oó]n|pulgadas|\b(?:ddr[345]|hdmi|displayport)\b", re.I)


def compactar_texto(texto):
    """Quita el relleno de la extraccion con diseno: espacios al final de linea, rachas
    largas de espacios y lineas vacias repetidas. Conserva la separacion entre columnas."""
    lineas = [re.sub(r" {4,}", "   ", linea.rstrip()) for linea in texto.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lineas)).strip()


def pagina_relevante(texto):
    """Conserva identidad y fichas tecnicas sin exigir una marca de nuestra lista."""
    if PATRON_MONTO.search(texto):
        return True
    if PATRON_PALABRA_PRODUCTO.search(texto) or PATRON_IDENTIDAD_TECNICA.search(texto):
        return True
    especificaciones = {m.group(0).lower() for m in PATRON_ESPECIFICACION_TECNICA.finditer(texto)}
    return len(especificaciones) >= 2


def dividir_bloque_proveedor(bloque, limite):
    """Divide texto largo por lineas, conservando filas y encabezados de tablas.

    Una imagen se mantiene con su pagina completa. Una fila individual que supera
    el limite se mantiene entera: el limite es un objetivo, nunca un recorte.
    """
    if bloque.get("imagen") or len(bloque["texto"]) <= limite:
        return [bloque]
    partes, actual, encabezados = [], [], []
    largo = 0
    for linea in bloque["texto"].splitlines():
        if actual and largo + len(linea) + 1 > limite:
            partes.append({**bloque, "texto": "\n".join(actual)})
            actual = list(encabezados)
            largo = sum(len(valor) + 1 for valor in actual)
        actual.append(linea)
        largo += len(linea) + 1
        if linea.startswith(("[HOJA ", "[PAGINA ")):
            encabezados = [linea]
        elif linea.startswith("[TABLA "):
            encabezados = [valor for valor in encabezados if not valor.startswith(("[TABLA ", "FILA "))]
            encabezados.append(linea)
        elif linea.startswith("FILA 1:"):
            encabezados = [valor for valor in encabezados if not valor.startswith("FILA ")]
            # Repite solo encabezados razonablemente cortos para no duplicar una fila gigante.
            if len(linea) < min(1000, limite // 4):
                encabezados.append(linea)
    if actual:
        partes.append({**bloque, "texto": "\n".join(actual)})
    return partes


def preparar_bloques_proveedor(carpeta, archivos, args):
    """Lee todos los anexos y devuelve (bloques, registros, necesita_ocr, no_soportados, marcas).
    Cada bloque es una pagina (o un archivo sin paginas) con su texto y, si corresponde, su imagen."""
    bloques, registros, necesita_ocr, no_soportados = [], [], [], []
    descartados = []  # paginas con texto que el filtro no envio (respaldo)
    marcas = set()
    huellas = {}
    for path in sorted(archivos, key=prioridad):
        verificar_detencion()
        relativo = str(path.relative_to(carpeta))
        extension = path.suffix.lower()
        registro = {"archivo": relativo, "tipo": extension, "tipo_documental": tipo_documental(path.name),
                    "estado_lectura": "", "caracteres": 0, "estado_ia": "no_ejecutada",
                    "productos_encontrados": 0, "observaciones": None, "paginas_enviadas": []}
        registros.append(registro)
        # El mismo PDF subido como economico y como tecnico: se lee una sola vez (el economico va primero).
        try:
            huella = hashlib.md5(Path(ruta_sistema(path)).read_bytes()).hexdigest()
        except OSError:
            huella = None
        if huella and huella in huellas:
            registro.update({"estado_lectura": "duplicado", "estado_ia": "omitido_duplicado",
                             "observaciones": f"mismo contenido que {huellas[huella]}"})
            continue
        if huella:
            huellas[huella] = relativo

        if extension in EXTENSIONES_IMAGEN:
            registro["estado_lectura"] = "imagen"
            if args.vision:
                bloques.append({"archivo": relativo, "pagina": 1, "texto": "", "imagen": cargar_imagen_archivo(path)})
                registro["paginas_enviadas"] = [1]
            else:
                no_soportados.append(relativo)
            continue

        if extension == ".pdf":
            detalle = {}
            texto, estado = extraer_archivo(path, args, detalle)
            registro["estado_lectura"] = estado
            registro["caracteres"] = len(texto)
            if documento_de_historial(relativo, texto):
                registro["observaciones"] = "documento de experiencia/ordenes de compra pasadas: no es lo ofertado"
                continue
            texto_paginas = detalle.get("texto_paginas") or ({0: texto} if texto.strip() else {})
            paginas_vision = detalle.get("paginas_vision") or {}
            if detalle.get("paginas_vacias"):
                registro["paginas_vacias"] = [indice + 1 for indice in detalle["paginas_vacias"]]
            if paginas_vision and args.vision:
                registro["paginas_vision"] = {str(i + 1): m for i, m in sorted(paginas_vision.items())}
            escaneadas = sorted(i for i, motivo in paginas_vision.items() if motivo == "sin_texto")
            imagenes = {}
            if args.vision and paginas_vision and not args.vision_solo_escaneadas:
                imagenes = renderizar_paginas_pdf(path, sorted(paginas_vision), args.dpi_vision)
            else:
                sin_leer = list(escaneadas)
                ruido_admin = es_ruido_administrativo(relativo)
                if escaneadas and args.ocr and not ruido_admin:
                    paginas_ocr = {}
                    _, estado_ocr = extraer_pdf_ocr(path, args.max_paginas, args.tesseract_cmd, args.idioma_ocr,
                                                    args.dpi_ocr, paginas=set(escaneadas), texto_paginas=paginas_ocr)
                    if estado_ocr.startswith("error_") or estado_ocr == "ocr_dependencia_no_disponible":
                        registro["estado_ocr"] = estado_ocr
                    texto_paginas.update(paginas_ocr)
                    sin_leer = [indice for indice in escaneadas
                                if len(paginas_ocr.get(indice, "").strip()) <= 40]
                if sin_leer and args.vision:
                    imagenes = renderizar_paginas_pdf(path, sin_leer, args.dpi_vision)
                    registro["paginas_vision"] = {str(i + 1): "sin_texto" for i in sin_leer}
                elif sin_leer and not ruido_admin and relativo not in necesita_ocr:
                    necesita_ocr.append(relativo)
            for indice in sorted(set(texto_paginas) | set(imagenes)):
                texto_pagina = texto_paginas.get(indice, "")
                marcas.update(coincidencia.group(0) for coincidencia in PATRON_MARCA_PAGINA.finditer(texto_pagina))
                bloque = {"archivo": relativo, "pagina": indice + 1,
                          "texto": compactar_texto(texto_pagina),
                          "imagen": imagenes.get(indice)}
                if indice in imagenes or pagina_relevante(texto_pagina):
                    bloques.append(bloque)
                    registro["paginas_enviadas"].append(indice + 1)
                elif bloque["texto"].strip():
                    descartados.append(bloque)
            continue

        texto, estado = extraer_archivo(path, args)
        registro["estado_lectura"] = estado
        registro["caracteres"] = len(texto)
        if estado in ("no_soportado", "xls_legacy_no_soportado"):
            no_soportados.append(relativo)
            continue
        marcas.update(coincidencia.group(0) for coincidencia in PATRON_MARCA_PAGINA.finditer(texto))
        if texto.strip():
            bloque = {"archivo": relativo, "pagina": None,
                      "texto": compactar_texto(texto), "imagen": None}
            if pagina_relevante(texto):
                bloques.append(bloque)
                registro["paginas_enviadas"] = ["todo"]
            else:
                descartados.append(bloque)

    if not bloques and descartados:
        # Si el filtro descarto todo, revisar todos los bloques disponibles.
        # agrupar_bloques los reparte en llamadas sin perder el final del documento.
        por_archivo = {registro["archivo"]: registro for registro in registros}
        for bloque in descartados:
            bloques.append(bloque)
            registro = por_archivo[bloque["archivo"]]
            registro["paginas_enviadas"].append(bloque["pagina"] or "todo")
            registro["respaldo_sin_filtro"] = True
    return bloques, registros, necesita_ocr, no_soportados, marcas


def agrupar_bloques(bloques, max_chars, max_imagenes):
    """Reparte los bloques en llamadas respetando el limite de texto y de imagenes.
    Casi todos los proveedores caben en una sola llamada."""
    grupos, actual, chars, imagenes = [], [], 0, 0
    bloques = [parte for bloque in bloques
               for parte in dividir_bloque_proveedor(bloque, max(1, max_chars - 80))]
    for bloque in bloques:
        largo = len(bloque["texto"]) + 80
        con_imagen = 1 if bloque.get("imagen") else 0
        if actual and (chars + largo > max_chars or imagenes + con_imagen > max_imagenes):
            grupos.append(actual)
            actual, chars, imagenes = [], 0, 0
        actual.append(bloque)
        chars += largo
        imagenes += con_imagen
    if actual:
        grupos.append(actual)
    return grupos


def archivo_declarado(nombre, archivos_llamada):
    """Asocia el 'archivo' que devuelve el modelo con un archivo real del proveedor."""
    nombre = str(nombre or "").strip().lower()
    if not nombre:
        return None
    for relativo in archivos_llamada:
        if nombre == relativo.lower() or nombre == Path(relativo).name.lower():
            return relativo
    for relativo in archivos_llamada:
        base = Path(relativo).stem.lower()
        if base and (base in nombre or nombre in relativo.lower()):
            return relativo
    return None


def procesar_oferta_por_proveedor(carpeta, args, seven_zip):
    """Una llamada (o pocas) por proveedor con las paginas relevantes de todos sus anexos."""
    inicio = time.perf_counter()
    info = cargar_json(carpeta / "oferta.json")
    proveedor = info.get("proveedor") or carpeta.name
    rut = info.get("rut") or carpeta.name.split("__", 1)[0]
    total_oferta = info.get("total") or info.get("total_oferta") or ""
    archivos_comprimidos = descomprimir(carpeta, seven_zip)
    inicio_lectura = time.perf_counter()
    bloques, registros, necesita_ocr, no_soportados, marcas = preparar_bloques_proveedor(
        carpeta, listar_archivos_oferta(carpeta), args)
    segundos_lectura = time.perf_counter() - inicio_lectura
    por_archivo = {registro["archivo"]: registro for registro in registros}

    debug_dir = carpeta / "_debug_ia"
    if args.debug:
        Path(ruta_sistema(debug_dir)).mkdir(exist_ok=True)

    parciales, estados, llamadas = [], [], 0
    segundos_modelo = 0.0
    grupos = [] if args.solo_texto else agrupar_bloques(bloques, args.max_chars_proveedor, args.max_imagenes_proveedor)
    print(f"  {proveedor[:34]}: lectura={segundos_lectura:.1f}s | "
          f"bloques relevantes={len(bloques)} | llamadas de extraccion={len(grupos)}", flush=True)
    for numero, grupo in enumerate(grupos, 1):
        verificar_detencion()
        partes, imagenes, lista_imagenes = [], [], []
        for bloque in grupo:
            etiqueta = f"=== ARCHIVO: {bloque['archivo']} ({tipo_documental(bloque['archivo'])})" + (
                f" | PAGINA {bloque['pagina']}" if bloque["pagina"] else "") + " ==="
            if bloque.get("imagen"):
                imagenes.append(bloque["imagen"])
                lista_imagenes.append(f"{len(imagenes)}) {bloque['archivo']} pagina {bloque['pagina']}")
                etiqueta += f" [IMAGEN {len(imagenes)}]"
            partes.append(f"{etiqueta}\n{bloque['texto'] or '(sin texto: ver imagen)'}")
        nota = NOTA_IMAGENES_PROVEEDOR.format(lista_imagenes="\n".join(lista_imagenes)) if imagenes else ""
        prompt = PROMPT_PROVEEDOR.format(proveedor=proveedor, total_oferta=total_oferta or "no informado",
                                         nota_imagenes=nota, documentos="\n\n".join(partes))
        inicio_modelo = time.perf_counter()
        datos, estado, cruda, consumo = consultar_modelo(
            prompt, args.modelo, args.timeout, args.num_ctx, args.backend, args.max_tokens,
            args.reintentos_modelo, args.espera_reintento, imagenes=imagenes or None)
        segundos_modelo += time.perf_counter() - inicio_modelo
        llamadas += 1
        estados.append(estado)
        archivos_llamada = list(dict.fromkeys(bloque["archivo"] for bloque in grupo))
        # el consumo se asigna al primer archivo de la llamada; los demas quedan con 0 llamadas
        for indice, relativo in enumerate(archivos_llamada):
            registro = por_archivo[relativo]
            registro["estado_ia"] = estado if registro["estado_ia"] in ("no_ejecutada", *ESTADOS_IA_OK) else registro["estado_ia"]
            anterior = registro.get("consumo_local") or {"llamadas": 0}
            nuevo = {**consumo, "llamadas": 1} if indice == 0 else {"llamadas": 0}
            registro["consumo_local"] = combinar_consumos([anterior, nuevo], args.modelo)
            registro["consumo_local"]["llamadas"] = anterior.get("llamadas", 0) + nuevo["llamadas"]
            registro["llamada_proveedor"] = numero
        if args.debug:
            escribir_texto(debug_dir / f"proveedor_{numero:02d}__prompt.txt", prompt)
            escribir_texto(debug_dir / f"proveedor_{numero:02d}__respuesta.txt", cruda)
        if isinstance(datos, dict):
            for producto in datos.get("productos", []):
                if not isinstance(producto, dict):
                    continue
                archivo_producto = archivo_declarado(producto.get("archivo_producto"), archivos_llamada)
                archivo_precio = archivo_declarado(producto.get("archivo_precio"), archivos_llamada)
                relativo = (
                    archivo_declarado(producto.get("archivo"), archivos_llamada)
                    or archivo_precio
                    or archivo_producto
                    or archivos_llamada[0]
                )
                if producto.get("precio_total") not in (None, "") and not producto.get("precio_total_tipo"):
                    producto["precio_total_tipo"] = "linea"
                limpio = normalizar_producto(producto, proveedor, rut, relativo, tipo_documental(relativo))
                if limpio:
                    limpio["fuente_producto"] = archivo_producto
                    limpio["fuente_precio"] = archivo_precio
                    parciales.append(limpio)
                    por_archivo[relativo]["productos_encontrados"] += 1
        if args.pausa_archivo > 0:
            time.sleep(args.pausa_archivo)

    relevantes = filtrar_productos_relevantes(parciales)
    if len(grupos) > 1 and not args.sin_consolidar and parciales:
        inicio_modelo = time.perf_counter()
        productos, estado_consolidacion, _, _, consumo_consolidacion = consolidar_parciales(
            parciales, proveedor, rut, total_oferta, args)
        segundos_modelo += time.perf_counter() - inicio_modelo
        llamadas += consumo_consolidacion.get("llamadas", 0)
    else:
        productos, estado_consolidacion, consumo_consolidacion = relevantes, "omitida", {}
    productos = completar_marcas_desde_parciales(productos, parciales)
    productos = reconciliar_precios_con_anexos(productos, bloques, total_oferta)
    productos = rescatar_filas_economicas(productos, relevantes, bloques, proveedor, rut)
    productos = cruzar_filas_agregadas(productos, bloques)
    productos = completar_marcas_por_nombre_de_archivo(productos, [r["archivo"] for r in registros])
    productos = completar_marcas_por_linea(productos, bloques)

    return {
        "proveedor": proveedor,
        "rut": rut,
        "modelo": args.modelo,
        "backend": args.backend,
        "vision": bool(args.vision),
        "modo": "por_proveedor",
        "estado_oferta": info.get("estado", ""),
        "total_oferta": total_oferta,
        "productos": productos,
        "productos_parciales": relevantes,
        "resultados_archivos": registros,
        "estado_consolidacion": estado_consolidacion,
        "consumo_consolidacion": consumo_consolidacion,
        "archivos_comprimidos": archivos_comprimidos,
        "necesita_ocr": necesita_ocr,
        "ocr_intentado": bool(getattr(args, "ocr", False) and OCR_DISPONIBLE),
        "no_soportados": no_soportados,
        "marcas_detectadas_texto": sorted(marcas),
        "llamadas_modelo": llamadas,
        "segundos": round(time.perf_counter() - inicio, 1),
        "segundos_lectura": round(segundos_lectura, 1),
        "segundos_modelo": round(segundos_modelo, 1),
    }


def bloques_economicos_completos(carpeta, archivos, bloques_leidos, args):
    """Texto de todos los anexos economicos legibles: los ya leidos en esta corrida y, si hubo
    reutilizacion de checkpoints, los demas (lectura local, sin modelo)."""
    bloques = list(bloques_leidos)
    leidos = {bloque["archivo"] for bloque in bloques}
    for path in archivos:
        relativo = str(path.relative_to(carpeta))
        if (relativo in leidos or tipo_documental(path.name) != "economico"
                or path.suffix.lower() not in (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".txt")):
            continue
        try:
            texto, estado = extraer_archivo(path, args)
        except Exception:
            continue
        if estado.startswith("texto") and texto.strip():
            bloques.append({"archivo": relativo, "pagina": None, "texto": texto})
    return bloques


def procesar_oferta(carpeta, args, seven_zip, previo=None, archivos_reprocesar=None):
    """Procesa un proveedor.

    Si se entregan `previo` y `archivos_reprocesar`, solo se leen y consultan al
    modelo esos archivos; los demas reutilizan su registro y productos parciales
    del resultado anterior. La consolidacion se vuelve a ejecutar siempre.
    """
    inicio_oferta = time.perf_counter()
    info = cargar_json(carpeta / "oferta.json")
    proveedor = info.get("proveedor") or carpeta.name
    rut = info.get("rut") or carpeta.name.split("__", 1)[0]
    total_oferta = info.get("total") or info.get("total_oferta") or ""
    archivos_comprimidos = descomprimir(carpeta, seven_zip)

    archivos = listar_archivos_oferta(carpeta)

    parciales = []
    resultados_archivos = []
    necesita_ocr = []
    no_soportados = []
    marcas_texto = set()

    bloques_economicos = []
    numeros_documentos = set()
    reutilizar = previo is not None and archivos_reprocesar is not None
    registros_previos = {}
    parciales_previos = {}
    if reutilizar:
        registros_previos = {registro.get("archivo"): registro for registro in previo.get("resultados_archivos", [])}
        for producto in previo.get("productos_parciales", []):
            parciales_previos.setdefault(producto.get("archivo_fuente"), []).append(producto)
        marcas_texto.update(previo.get("marcas_detectadas_texto") or [])

    debug_dir = carpeta / "_debug_ia"
    if args.debug:
        Path(ruta_sistema(debug_dir)).mkdir(exist_ok=True)

    for indice, path in enumerate(sorted(archivos, key=prioridad), 1):
        verificar_detencion()
        relativo = str(path.relative_to(carpeta))
        if reutilizar and relativo not in archivos_reprocesar and relativo in registros_previos:
            registro = dict(registros_previos[relativo])
            registro["reutilizado"] = True
            parciales.extend(parciales_previos.get(relativo, []))
            if relativo in set(previo.get("necesita_ocr") or []):
                necesita_ocr.append(relativo)
            if relativo in set(previo.get("no_soportados") or []):
                no_soportados.append(relativo)
            resultados_archivos.append(registro)
            continue

        detalle = {} if args.vision and path.suffix.lower() == ".pdf" else None
        texto, estado_lectura = extraer_archivo(path, args, detalle)
        if estado_lectura.startswith("texto") and tipo_documental(path.name) == "economico":
            bloques_economicos.append({"archivo": relativo, "pagina": None, "texto": texto})
        if texto:
            numeros_documentos |= numeros_del_texto(texto)
        if detalle and args.vision_solo_escaneadas:
            detalle["paginas_vision"] = {i: m for i, m in (detalle.get("paginas_vision") or {}).items()
                                         if m == "sin_texto"}
        for marca in MARCAS:
            if re.search(rf"\b{re.escape(marca)}\b", texto, re.I):
                marcas_texto.add(marca)

        registro = {
            "archivo": relativo,
            "tipo": path.suffix.lower(),
            "tipo_documental": tipo_documental(path.name),
            "estado_lectura": estado_lectura,
            "caracteres": len(texto),
            "estado_ia": "no_ejecutada",
            "productos_encontrados": 0,
            "observaciones": None
        }
        paginas_vision = (detalle or {}).get("paginas_vision") or {}
        if paginas_vision:
            registro["paginas_vision"] = {str(index + 1): motivo for index, motivo in sorted(paginas_vision.items())}

        if "ocr_pendiente" in estado_lectura:
            necesita_ocr.append(relativo)
        if estado_lectura in ("necesita_ocr", "ocr_dependencia_no_disponible") or estado_lectura.startswith("error_ocr"):
            if relativo not in necesita_ocr:
                necesita_ocr.append(relativo)
        elif estado_lectura in ("no_soportado", "xls_legacy_no_soportado"):
            no_soportados.append(relativo)
        elif (estado_lectura.startswith("texto") and texto and not args.solo_texto
              and not args.sin_filtro_tecnicos and documento_de_historial(relativo, texto)):
            registro["estado_ia"] = "omitido_sin_equipos"
            registro["observaciones"] = "documento de experiencia/ordenes de compra pasadas: no es lo ofertado"
        elif estado_lectura in ("texto_vision", "imagen") and not args.solo_texto:
            productos, estado_ia, cruda, prompt, observaciones, consumo = extraer_parciales_mixto(
                path, relativo, detalle or {}, proveedor, rut, total_oferta, args
            )
            parciales.extend(productos)
            registro["estado_ia"] = estado_ia
            registro["productos_encontrados"] = len(productos)
            registro["observaciones"] = observaciones
            registro["consumo_local"] = consumo
            registro["fragmentos"] = consumo.get("llamadas", 1)
            if args.debug:
                base = f"{indice:02d}__{re.sub(r'[^\w.-]', '_', path.stem)[:70]}"
                escribir_texto(debug_dir / f"{base}__prompt.txt", prompt)
                escribir_texto(debug_dir / f"{base}__respuesta.txt", cruda)
        elif (estado_lectura.startswith("texto") and texto and not args.solo_texto
              and tipo_documental(path.name) != "economico" and not args.sin_filtro_tecnicos
              and anexo_tecnico_sin_equipos(texto, relativo)):
            registro["estado_ia"] = "omitido_sin_equipos"
            registro["observaciones"] = "anexo tecnico sin especificaciones, marca/modelo ni montos de equipos"
        elif estado_lectura.startswith("texto") and texto and not args.solo_texto:
            productos, estado_ia, cruda, prompt, observaciones, consumo = extraer_parciales_archivo(
                Path(relativo), texto, proveedor, rut, total_oferta, args
            )
            parciales.extend(productos)
            registro["estado_ia"] = estado_ia
            registro["productos_encontrados"] = len(productos)
            registro["observaciones"] = observaciones
            registro["consumo_local"] = consumo
            registro["fragmentos"] = consumo.get("llamadas", 1)
            if args.debug:
                base = f"{indice:02d}__{re.sub(r'[^\w.-]', '_', path.stem)[:70]}"
                escribir_texto(debug_dir / f"{base}__prompt.txt", prompt)
                escribir_texto(debug_dir / f"{base}__respuesta.txt", cruda)
        elif args.solo_texto and estado_lectura.startswith("texto"):
            registro["estado_ia"] = "solo_texto"

        resultados_archivos.append(registro)

    productos_finales = []
    estado_consolidacion = "no_ejecutada"
    consumo_consolidacion = {}
    if not args.solo_texto:
        productos_finales, estado_consolidacion, cruda_consolidacion, prompt_consolidacion, consumo_consolidacion = consolidar_parciales(
            parciales, proveedor, rut, total_oferta, args
        )
        if args.debug and prompt_consolidacion:
            escribir_texto(debug_dir / "99__consolidacion__prompt.txt", prompt_consolidacion)
            escribir_texto(debug_dir / "99__consolidacion__respuesta.txt", cruda_consolidacion)

    if productos_finales and not args.solo_texto:
        bloques_completos = bloques_economicos_completos(carpeta, archivos, bloques_economicos, args)
        productos_finales = reconciliar_precios_con_anexos(productos_finales, bloques_completos, total_oferta, anular=False)
        productos_finales = rescatar_filas_economicas(
            productos_finales, filtrar_productos_relevantes(parciales), bloques_completos, proveedor, rut)
        productos_finales = cruzar_filas_agregadas(productos_finales, bloques_completos)
        productos_finales = completar_marcas_por_nombre_de_archivo(
            productos_finales, [str(a.relative_to(carpeta)) for a in archivos])
        productos_finales = completar_marcas_por_linea(productos_finales, bloques_completos)
        if not reutilizar:  # con checkpoints reutilizados faltan los textos de los archivos no releidos
            for bloque in bloques_completos:
                numeros_documentos |= numeros_del_texto(bloque["texto"])
            productos_finales = marcar_precios_sin_respaldo(productos_finales, numeros_documentos, bloques_completos)

    return {
        "proveedor": proveedor,
        "rut": rut,
        "modelo": args.modelo,
        "backend": args.backend,
        "vision": bool(args.vision),
        "modo": "por_archivo",
        "segundos": round(time.perf_counter() - inicio_oferta, 1),
        "estado_oferta": info.get("estado", ""),
        "total_oferta": total_oferta,
        "productos": productos_finales,
        "productos_parciales": filtrar_productos_relevantes(parciales),
        "resultados_archivos": resultados_archivos,
        "estado_consolidacion": estado_consolidacion,
        "consumo_consolidacion": consumo_consolidacion,
        "archivos_comprimidos": archivos_comprimidos,
        "necesita_ocr": necesita_ocr,
        "ocr_intentado": bool(getattr(args, "ocr", False) and OCR_DISPONIBLE),
        "no_soportados": no_soportados,
        "marcas_detectadas_texto": sorted(marcas_texto)
    }


def diagnosticar_vision(licitaciones, args, ruta_csv):
    """Clasifica las paginas como lo haria --vision, sin llamar al modelo ni tocar
    checkpoints. Sirve para calibrar --vision-min-filas con documentos reales."""
    filas = []
    for licitacion in licitaciones:
        oferentes = sorted(
            carpeta for carpeta in listar_subdirectorios(licitacion)
            if not carpeta.name.startswith("_")
        )
        if args.proveedor:
            oferentes = [carpeta for carpeta in oferentes if carpeta.name == args.proveedor]
        if args.limite_proveedores:
            oferentes = oferentes[:args.limite_proveedores]
        for carpeta in oferentes:
            for path in sorted(listar_archivos_oferta(carpeta), key=prioridad):
                extension = path.suffix.lower()
                if extension not in EXTENSIONES_IMAGEN and extension != ".pdf":
                    continue
                fila = {
                    "licitacion": licitacion.name,
                    "proveedor": carpeta.name,
                    "archivo": str(path.relative_to(carpeta)),
                    "tipo_documental": tipo_documental(path.name),
                    "paginas_leidas": 1,
                    "paginas_vision": "",
                    "motivos": "",
                }
                if extension in EXTENSIONES_IMAGEN:
                    fila.update(paginas_vision="1", motivos="imagen")
                else:
                    detalle = {}
                    extraer_archivo(path, args, detalle)
                    paginas = detalle.get("paginas_vision", {})
                    fila["paginas_leidas"] = len(detalle.get("texto_paginas", {}))
                    fila["paginas_vision"] = ", ".join(str(indice + 1) for indice in sorted(paginas))
                    fila["motivos"] = ", ".join(paginas[indice] for indice in sorted(paginas))
                filas.append(fila)
                print(f"  {licitacion.name} | {carpeta.name[:30]:30} | {fila['archivo'][:40]:40} | "
                      f"vision: {fila['paginas_vision'] or '-'} {('(' + fila['motivos'] + ')') if fila['motivos'] else ''}")

    Path(ruta_sistema(ruta_csv.parent)).mkdir(parents=True, exist_ok=True)
    campos = ["licitacion", "proveedor", "archivo", "tipo_documental", "paginas_leidas", "paginas_vision", "motivos"]
    with open(ruta_sistema(ruta_csv), "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=campos)
        escritor.writeheader()
        escritor.writerows(filas)

    con_vision = [fila for fila in filas if fila["paginas_vision"]]
    motivos = {}
    for fila in con_vision:
        for motivo in filter(None, fila["motivos"].split(", ")):
            motivos[motivo] = motivos.get(motivo, 0) + 1
    paginas_imagen = sum(len(fila["paginas_vision"].split(", ")) for fila in con_vision)
    por_llamada = max(1, args.vision_paginas_por_llamada)
    llamadas = sum(
        -(-len(fila["paginas_vision"].split(", ")) // por_llamada) for fila in con_vision
    )
    print("\n" + "=" * 74)
    print(f"DIAGNOSTICO VISION (--vision-min-filas {args.vision_min_filas})")
    print(f"Archivos PDF/imagen revisados: {len(filas)}")
    print(f"Archivos que usarian vision:   {len(con_vision)}")
    print(f"Paginas como imagen:           {paginas_imagen}  (~{llamadas} llamadas con imagen)")
    for motivo, cantidad in sorted(motivos.items(), key=lambda item: -item[1]):
        print(f"  {motivo:18} {cantidad} pagina(s)")
    print(f"Detalle: {ruta_csv}")
    print("=" * 74)


def listar_subdirectorios(carpeta):
    with os.scandir(ruta_sistema(carpeta)) as entradas:
        nombres = sorted(entrada.name for entrada in entradas if entrada.is_dir(follow_symlinks=False))
    return [Path(carpeta) / nombre for nombre in nombres]


def detectar_licitaciones(raiz):
    """Acepta una licitacion, un mes o un arbol de meses/años sin recorrer anexos."""
    pendientes = [Path(raiz)]
    encontradas = []
    while pendientes:
        carpeta = pendientes.pop()
        hijos = listar_subdirectorios(carpeta)
        if any(os.path.isfile(ruta_sistema(hijo / "oferta.json")) for hijo in hijos):
            encontradas.append(carpeta)
        else:
            pendientes.extend(reversed([hijo for hijo in hijos if not hijo.name.startswith("_")]))
    return sorted(encontradas, key=lambda path: str(path).lower())


def escribir_log(ruta, evento, **datos):
    registro = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "evento": evento,
        **datos,
    }
    Path(ruta_sistema(ruta.parent)).mkdir(parents=True, exist_ok=True)
    with LOG_LOCK, open(ruta_sistema(ruta), "a", encoding="utf-8") as archivo:
        archivo.write(json.dumps(registro, ensure_ascii=False) + "\n")


def resultado_requiere_reproceso(resultado, usar_ocr=False, modelo=None, backend=None, usar_vision=False,
                                 modo=None, solo_escaneadas=False):
    estados = [
        str(archivo.get("estado_ia") or "")
        for archivo in resultado.get("resultados_archivos", [])
        if archivo.get("estado_ia") not in ESTADOS_IA_NEUTROS
    ]
    estado_consolidacion = str(resultado.get("estado_consolidacion") or "")
    fallo_modelo = any(estado not in ESTADOS_IA_OK for estado in estados)
    fallo_consolidacion = estado_consolidacion not in ESTADOS_IA_OK | {"", "omitida", "no_ejecutada"}
    cambio_modelo = bool(
        modelo
        and resultado.get("modelo")
        and str(resultado.get("modelo")) != str(modelo)
    )
    cambio_backend = bool(
        backend
        and resultado.get("backend")
        and str(resultado.get("backend")) != str(backend)
    )
    return (
        fallo_modelo
        or fallo_consolidacion
        or cambio_modelo
        or cambio_backend
        or (usar_ocr and bool(resultado.get("necesita_ocr")) and not resultado.get("ocr_intentado"))
        # Al activar --vision se reprocesan los proveedores leidos solo con texto.
        or (usar_vision and (bool(resultado.get("necesita_ocr")) if solo_escaneadas
                             else not resultado.get("vision")))
        or bool(modo and (resultado.get("modo") or "por_archivo") != modo)
    )


def registrar_resultado(codigo, metadata, resultado, productos, consumos, resumenes, archivos_ocr):
    resultado["productos"] = validar_productos(
        resultado.get("productos", []),
        resultado.get("productos_parciales", []),
        resultado.get("total_oferta"),
    )
    RESULTADOS_PROVEEDOR[(codigo, str(resultado.get("rut") or "").strip())] = resultado
    contexto = {
        "codigo": codigo,
        "nombre_licitacion": metadata.get("nombre_licitacion", ""),
        "fecha_publicacion": metadata.get("fecha_publicacion", ""),
        "estado_licitacion": metadata.get("estado_licitacion", ""),
        "organismo": metadata.get("organismo", ""),
    }
    for producto in resultado["productos"]:
        productos.append({**contexto, **producto})

    llamadas = 0
    errores_modelo = 0
    for archivo in resultado.get("resultados_archivos", []):
        estado = str(archivo.get("estado_ia") or "no_ejecutada")
        if estado in ESTADOS_IA_NEUTROS:
            continue
        consumo = archivo.get("consumo_local") or {}
        llamadas_archivo = int(consumo["llamadas"]) if consumo.get("llamadas") is not None else 1
        llamadas += llamadas_archivo
        if estado not in ESTADOS_IA_OK:
            errores_modelo += 1
        consumos.append({
            **contexto,
            "proveedor": resultado.get("proveedor", ""),
            "rut": resultado.get("rut", ""),
            "tipo_llamada": "archivo",
            "archivo": archivo.get("archivo", ""),
            "estado": estado,
            **consumo,
        })

    estado_consolidacion = str(resultado.get("estado_consolidacion") or "no_ejecutada")
    if estado_consolidacion not in ("no_ejecutada", "omitida"):
        consumo = resultado.get("consumo_consolidacion") or {}
        llamadas_consolidacion = int(consumo.get("llamadas") or 1)
        llamadas += llamadas_consolidacion
        if estado_consolidacion not in ESTADOS_IA_OK:
            errores_modelo += 1
        consumos.append({
            **contexto,
            "proveedor": resultado.get("proveedor", ""),
            "rut": resultado.get("rut", ""),
            "tipo_llamada": "consolidacion",
            "archivo": "",
            "estado": estado_consolidacion,
            **consumo,
        })

    pendientes_ocr = resultado.get("necesita_ocr", [])
    archivos_ocr.extend(
        {
            "licitacion": codigo,
            "proveedor": resultado.get("proveedor", ""),
            "archivo": archivo,
        }
        for archivo in pendientes_ocr
    )
    productos_validos = sum(
        producto.get("estado_validacion") == "ok" for producto in resultado["productos"]
    )
    resumenes.append({
        **contexto,
        "proveedor": resultado.get("proveedor", ""),
        "rut": resultado.get("rut", ""),
        "estado_oferta": resultado.get("estado_oferta", ""),
        "productos": len(resultado["productos"]),
        "productos_validos": productos_validos,
        "productos_revisar": len(resultado["productos"]) - productos_validos,
        "llamadas": llamadas,
        "errores_modelo": errores_modelo,
        "ocr": len(pendientes_ocr),
        "no_soportados": len(resultado.get("no_soportados", [])),
    })


def generar_excel(productos, consumos, resumenes, ruta):
    columnas = [
        "ID licitación", "Nombre licitación", "Fecha licitación", "Proveedor", "Producto ofertado",
        "Tipo de producto", "Marca producto", "Cantidad", "Precio unitario", "Precio total",
        "Fuente de información", "Estado validación", "Alertas",
        "Cobertura documental", "Motivo cobertura",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    for columna, titulo in enumerate(columnas, 1):
        celda = ws.cell(1, columna, titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")

    def agregar_fuente(fuentes, tipo, archivo, pagina=None):
        if not archivo or archivo == "consolidacion":
            return
        texto = f"{tipo}: {archivo}"
        if pagina not in (None, ""):
            texto += f" (p. {pagina})"
        if texto not in fuentes:
            fuentes.append(texto)

    for indice, producto in enumerate(productos, 2):
        fuentes = []
        archivo_base = producto.get("archivo_fuente")
        archivo_producto = producto.get("fuente_producto")
        archivo_precio = producto.get("fuente_precio")
        if archivo_producto or archivo_precio:
            agregar_fuente(
                fuentes, "Producto", archivo_producto or archivo_base,
                producto.get("pagina_producto") or producto.get("pagina"),
            )
            if producto.get("precio_unitario") not in (None, "") or producto.get("precio_total") not in (None, ""):
                agregar_fuente(
                    fuentes, "Precio", archivo_precio or archivo_base,
                    producto.get("pagina_precio") or producto.get("pagina"),
                )
        else:
            agregar_fuente(fuentes, "Fuente", archivo_base, producto.get("pagina"))
        agregar_fuente(fuentes, "Marca", producto.get("fuente_marca"), producto.get("pagina_marca"))
        fuentes_primarias = {archivo_base, archivo_producto, archivo_precio, producto.get("fuente_marca")}
        for archivo in producto.get("fuentes_respaldo") or []:
            if archivo not in fuentes_primarias:
                agregar_fuente(fuentes, "Respaldo", archivo)

        valores = (
            producto.get("codigo"),
            producto.get("nombre_licitacion"),
            producto.get("fecha_publicacion"),
            producto.get("proveedor"),
            producto.get("producto"),
            subcategoria_producto(producto),
            producto.get("marca"),
            producto.get("cantidad"),
            producto.get("precio_unitario"),
            producto.get("precio_total"),
            "; ".join(fuentes),
            producto.get("estado_validacion"),
            "; ".join(producto.get("alertas") or []),
            producto.get("cobertura_documental"),
            producto.get("cobertura_motivo"),
        )
        for columna, valor in enumerate(valores, 1):
            ws.cell(indice, columna, valor)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta_sistema(ruta))


def main():
    global DEBUG
    parser = argparse.ArgumentParser(description="Extraccion por archivo y consolidacion ligera por proveedor")
    parser.add_argument("--dir", default="ofertas")
    parser.add_argument("--modelo", default="llama3.2")
    parser.add_argument("--backend", choices=("ollama", "lmstudio"), default="ollama",
                        help="Motor local: ollama o lmstudio")
    parser.add_argument("--solo-texto", action="store_true")
    parser.add_argument("--max-fragmentos-tecnico", type=int, default=3,
                        help="maximo de fragmentos que se leen de un anexo tecnico largo (0 = todos)")
    parser.add_argument("--rapido", action="store_true",
                        help="preset recomendado: --prompt-compacto y --sin-consolidar (las reglas de codigo "
                             "fusionan fichas, IVA, duplicados y marcas; ~2x mas rapido que el modo completo)")
    parser.add_argument("--prompt-compacto", action="store_true",
                        help="salida corta del modelo (sin evidencia literal ni confianza): mas rapido")
    parser.add_argument("--sin-filtro-tecnicos", action="store_true",
                        help="envia al modelo todos los anexos tecnicos, aun sin senales de equipos")
    parser.add_argument("--sin-consolidar", action="store_true")
    parser.add_argument("--sin-cobertura", action="store_true",
                        help="no calcula las columnas de cobertura documental del Excel (DATO_NO_EXISTE, "
                             "NO_ENCONTRADO, PENDIENTE_OCR)")
    parser.add_argument("--workers-cobertura", type=int, default=2,
                        help="procesos para releer los documentos al calcular la cobertura (default: 2)")
    parser.add_argument("--excel")
    parser.add_argument("--rehacer", action="store_true")
    parser.add_argument("--debug", action="store_true")
    parser.add_argument("--seven-zip")
    parser.add_argument("--max-paginas", type=int, default=20)
    parser.add_argument("--max-filas-excel", type=int, default=500)
    parser.add_argument("--max-columnas-excel", type=int, default=40)
    parser.add_argument("--espera-backend-min", type=int, default=30,
                        help="minutos que se espera a que LM Studio vuelva si deja de responder (default: 30)")
    parser.add_argument("--permitir-sin-ocr", action="store_true",
                        help="continua aunque Tesseract no funcione (los PDF escaneados quedaran sin leer)")
    parser.add_argument("--ocr", action="store_true",
                        help="Aplica Tesseract solo a PDF sin texto extraible")
    parser.add_argument("--tesseract-cmd",
                        help="Ruta a tesseract.exe si no esta disponible en PATH")
    parser.add_argument("--idioma-ocr", default="spa+eng")
    parser.add_argument("--dpi-ocr", type=int, default=180)
    parser.add_argument("--max-chars-archivo", type=int, default=7000)
    parser.add_argument("--max-chars-consolidacion", type=int, default=12000)
    parser.add_argument("--num-ctx", type=int, default=8192)
    parser.add_argument("--timeout", type=int, default=300)
    parser.add_argument("--reintentos-modelo", type=int, default=2)
    parser.add_argument("--espera-reintento", type=int, default=5)
    parser.add_argument("--log", help="Ruta del log JSONL; por defecto se guarda junto al Excel")
    parser.add_argument("--limite-licitaciones", type=int,
                        help="Procesa como maximo esta cantidad de licitaciones")
    parser.add_argument("--limite-proveedores", type=int,
                        help="Procesa como maximo esta cantidad de proveedores por licitacion")
    parser.add_argument("--pausa-licitacion", type=int, default=0,
                        help="Pausa en segundos despues de cada licitacion")
    parser.add_argument("--pausa-archivo", type=int, default=0,
                        help="Pausa en segundos entre llamadas al modelo")
    parser.add_argument("--paralelo", type=int, default=1,
                        help="Proveedores procesados a la vez. Debe coincidir con "
                             "'Max Concurrent Predictions' del modelo en LM Studio (ej. 4).")
    parser.add_argument("--vision", action="store_true",
                        help="Envia como imagen las paginas PDF con tablas sin bordes/colapsadas "
                             "o escaneadas, y los JPG/PNG adjuntos. Requiere modelo con vision.")
    parser.add_argument("--vision-solo-escaneadas", action="store_true",
                        help="Con --vision: solo las paginas sin texto van como imagen (con --ocr, solo las que el "
                             "OCR no pudo leer), y solo se reprocesan proveedores con OCR pendiente.")
    parser.add_argument("--dpi-vision", type=int, default=150,
                        help="Resolucion de las paginas enviadas como imagen (150 suele bastar).")
    parser.add_argument("--vision-paginas-por-llamada", type=int, default=3,
                        help="Paginas-imagen por llamada al modelo en modo --vision.")
    parser.add_argument("--rampa", type=float, default=5,
                        help="Segundos entre el inicio de los primeros proveedores en paralelo, para no "
                             "saturar LM Studio/LM Link al arrancar (0 = todos a la vez).")
    parser.add_argument("--sin-calentamiento", action="store_true",
                        help="No enviar la peticion de prueba inicial al modelo.")
    parser.add_argument("--por-proveedor", action="store_true",
                        help="Una llamada por proveedor con las paginas relevantes de todos sus anexos "
                             "(mas rapido). Sin esto: una llamada por fragmento de cada archivo + consolidacion.")
    parser.add_argument("--max-chars-proveedor", type=int, default=24000,
                        help="Texto maximo por llamada en --por-proveedor; si se excede se usan mas llamadas.")
    parser.add_argument("--max-imagenes-proveedor", type=int, default=4,
                        help="Paginas-imagen maximas por llamada en --por-proveedor (con --vision).")
    parser.add_argument("--vision-min-filas", type=int, default=2,
                        help="Filas con precio necesarias para tratar una pagina sin tabla "
                             "detectada como tabla sin bordes (mas alto = menos vision).")
    parser.add_argument("--diagnostico-vision", action="store_true",
                        help="Solo clasifica paginas (sin modelo ni checkpoints) y genera "
                             "diagnostico_vision.csv para calibrar la regla de vision.")
    parser.add_argument("--max-tokens", type=int, default=4096,
                        help="Maximo de tokens generados por llamada LM Studio")
    parser.add_argument("--proveedor",
                        help="Procesa solo el proveedor cuyo nombre coincide con la carpeta")
    parser.add_argument("--metadata-csv",
                        help="CSV opcional con codigo, nombre y fecha_publicacion; prevalece sobre el catalogo y los CSV mensuales")
    args = parser.parse_args()
    if args.rapido:
        args.prompt_compacto = True
        args.sin_consolidar = True
    DEBUG = args.debug
    global ESPERA_BACKEND_MINUTOS
    ESPERA_BACKEND_MINUTOS = max(0, args.espera_backend_min)

    if args.ocr and not args.solo_texto:
        ocr_ok, ocr_mensaje = verificar_ocr(args.tesseract_cmd, args.idioma_ocr)
        global OCR_DISPONIBLE
        OCR_DISPONIBLE = ocr_ok
        if ocr_ok:
            print(f"OCR disponible: {ocr_mensaje}")
        elif args.permitir_sin_ocr:
            print(f"AVISO: el OCR no funciona y se continua sin el: {ocr_mensaje}")
        else:
            sys.exit(f"ERROR: se pidio --ocr pero {ocr_mensaje}." + chr(10) +
                     "Sin OCR los PDF escaneados no se leen y se pierden productos. Corrige la instalacion "
                     "o usa --permitir-sin-ocr para continuar de todos modos.")

    raiz = Path(args.dir)
    if not os.path.isdir(ruta_sistema(raiz)):
        sys.exit(f"No existe: {raiz}")

    seven_zip = encontrar_7zip(args.seven_zip)
    licitaciones = detectar_licitaciones(raiz)
    if args.limite_licitaciones is not None:
        if args.limite_licitaciones < 1:
            sys.exit("--limite-licitaciones debe ser mayor que cero")
        licitaciones = licitaciones[:args.limite_licitaciones]
    metadata_licitaciones = cargar_metadata_para_corrida(licitaciones, args.metadata_csv)
    todos = []
    consumos = []
    resumenes = []
    archivos_ocr = []
    ruta_excel = Path(args.excel) if args.excel else raiz / "resultado_productos.xlsx"
    ruta_log = Path(args.log) if args.log else ruta_excel.with_suffix(".log.jsonl")

    if args.diagnostico_vision:
        diagnosticar_vision(licitaciones, args, ruta_excel.with_name("diagnostico_vision.csv"))
        return

    print("=" * 74)
    print(f"EXTRACCION IA v6 | backend={args.backend} | modelo={args.modelo} | licitaciones={len(licitaciones)}")
    print(f"7-Zip: {seven_zip or 'NO encontrado'} | solo_texto={args.solo_texto} "
          f"| paralelo={args.paralelo} | vision={args.vision}")
    print("Logica: fragmentos + vision selectiva + consolidacion y validacion auditable")
    print("=" * 74)
    escribir_log(
        ruta_log,
        "inicio_corrida",
        directorio=str(raiz),
        backend=args.backend,
        modelo=args.modelo,
        licitaciones=len(licitaciones),
        paralelo=args.paralelo,
        vision=args.vision,
        modo=modo_extraccion(args),
    )

    detener = False
    interrumpido = False
    if args.paralelo < 1:
        sys.exit("--paralelo debe ser mayor o igual a 1")
    DETENER.clear()

    # --- Fase 1: revisar checkpoints y armar la cola de proveedores pendientes ---
    # IMPORTANTE: los resultados "por revisar" se mantienen en la lista del checkpoint
    # hasta que su reproceso termine bien. Si el reproceso falla o se interrumpe, el
    # resultado anterior sigue intacto en extraccion_ia.json.
    estado_licitaciones = {}
    tareas = []
    for licitacion in licitaciones:
        oferentes = [
            carpeta for carpeta in listar_subdirectorios(licitacion)
            if not carpeta.name.startswith("_")
        ]
        if args.proveedor:
            oferentes = [carpeta for carpeta in oferentes if carpeta.name == args.proveedor]
            if not oferentes:
                sys.exit(
                    f"No se encontro el proveedor '{args.proveedor}' en {licitacion}"
                )
        if not oferentes:
            continue

        salida = licitacion / "extraccion_ia.json"
        resultados = []
        previos_por_carpeta = {}
        if os.path.isfile(ruta_sistema(salida)) and not args.rehacer:
            try:
                with open(ruta_sistema(salida), encoding="utf-8") as archivo:
                    resultados = json.load(archivo)
            except (json.JSONDecodeError, OSError):
                resultados = []
            carpeta_por_rut = {
                str(cargar_json(carpeta / "oferta.json").get("rut") or "").strip(): carpeta
                for carpeta in oferentes
            }
            # Proveedores ya procesados que pueden necesitar trabajo (fallos, OCR
            # pendiente, archivos nuevos o revision para vision). No se rehacen
            # completos: el hilo de trabajo decide que archivos repetir.
            for resultado in resultados:
                carpeta = carpeta_por_rut.get(str(resultado.get("rut") or "").strip())
                if carpeta is not None and revision_necesaria(resultado, carpeta, args):
                    previos_por_carpeta[carpeta] = resultado
                else:
                    registrar_resultado(
                        licitacion.name,
                        metadata_licitaciones.get(licitacion.name, {}),
                        resultado,
                        todos,
                        consumos,
                        resumenes,
                        archivos_ocr,
                    )
            ruts_guardados = {
                str(resultado.get("rut") or "").strip() for resultado in resultados
            }
            oferentes = [
                carpeta for carpeta in oferentes
                if carpeta in previos_por_carpeta
                or str(cargar_json(carpeta / "oferta.json").get("rut") or "").strip()
                not in ruts_guardados
            ]
            if not oferentes:
                print(f"{licitacion.name}: ya procesada ({len(resultados)} proveedores)")
                continue
            nuevos = len(oferentes) - len(previos_por_carpeta)
            print(f"{licitacion.name}: reanudando; {nuevos} proveedores sin procesar y "
                  f"{len(previos_por_carpeta)} por revisar (solo se repiten los archivos necesarios)")

        if args.limite_proveedores is not None:
            if args.limite_proveedores < 1:
                sys.exit("--limite-proveedores debe ser mayor que cero")
            oferentes = oferentes[:args.limite_proveedores]
        # Los "por revisar" que quedaron fuera de la cola (por --limite-proveedores)
        # se registran tal cual para que igual aparezcan en el Excel.
        for carpeta, previo in previos_por_carpeta.items():
            if carpeta not in oferentes:
                registrar_resultado(
                    licitacion.name,
                    metadata_licitaciones.get(licitacion.name, {}),
                    previo,
                    todos,
                    consumos,
                    resumenes,
                    archivos_ocr,
                )

        print(f"{licitacion.name}: {len(oferentes)} proveedores en cola")
        estado_licitaciones[licitacion.name] = {"salida": salida, "resultados": resultados}
        tareas.extend((licitacion, carpeta, previos_por_carpeta.get(carpeta)) for carpeta in oferentes)

    # --- Fase 2: procesar la cola, hasta --paralelo proveedores a la vez ---
    # Cada hilo procesa un proveedor completo (lectura + llamadas al modelo).
    # Los resultados, checkpoints y el Excel se manejan solo desde este hilo principal,
    # asi que el formato de extraccion_ia.json y la reanudacion no cambian.
    if args.pausa_licitacion > 0 and args.paralelo > 1:
        print("Aviso: --pausa-licitacion se ignora cuando --paralelo es mayor que 1.")
    ultima_licitacion = {"nombre": None}

    arranque = {"siguiente": 0}
    lock_arranque = threading.Lock()

    def trabajar(licitacion, carpeta, previo):
        verificar_detencion()
        with lock_arranque:
            orden = arranque["siguiente"]
            arranque["siguiente"] += 1
        if orden < args.paralelo and args.rampa > 0:
            DETENER.wait(orden * args.rampa)  # 0 s, rampa, 2*rampa... (interrumpible con Ctrl+C)
            verificar_detencion()
        if args.paralelo == 1 and args.pausa_licitacion > 0:
            if ultima_licitacion["nombre"] not in (None, licitacion.name):
                print(f"Pausa de {args.pausa_licitacion} segundos para enfriar el equipo...")
                DETENER.wait(args.pausa_licitacion)
                verificar_detencion()
            ultima_licitacion["nombre"] = licitacion.name
        info_carpeta = cargar_json(carpeta / "oferta.json")
        escribir_log(
            ruta_log,
            "inicio_proveedor",
            licitacion=licitacion.name,
            proveedor=info_carpeta.get("proveedor", carpeta.name),
            rut=info_carpeta.get("rut", ""),
        )
        procesar = procesar_oferta_por_proveedor if args.por_proveedor else procesar_oferta
        if previo is None:
            return procesar(carpeta, args, seven_zip)
        modo, archivos = planificar_reproceso(carpeta, previo, args, seven_zip)
        if modo == "completo" or (args.por_proveedor and modo in ("parcial", "consolidar")):
            # en modo por proveedor cualquier cambio rehace el proveedor: son 1 o 2 llamadas
            resultado = procesar(carpeta, args, seven_zip)
        elif modo == "nada":
            resultado = dict(previo)
            resultado["productos"] = list(previo.get("productos", []))
            resultado["vision"] = bool(previo.get("vision") or args.vision)
        else:
            resultado = procesar_oferta(carpeta, args, seven_zip, previo=previo, archivos_reprocesar=set(archivos))
        resultado["reproceso"] = {"modo": modo, "archivos": archivos or {}}
        return resultado

    total_tareas = len(tareas)
    completados = 0
    terminados_ok = set()

    def manejar(futuro, licitacion, carpeta):
        """Procesa el resultado de un proveedor en el hilo principal."""
        nonlocal detener, completados
        if futuro.cancelled():
            return
        info_carpeta = cargar_json(carpeta / "oferta.json")
        try:
            resultado = futuro.result()
        except CorridaDetenida:
            return  # abandonado por Ctrl+C o por caida del backend; el previo se conserva
        except RuntimeError as exc:
            escribir_log(
                ruta_log,
                "error_proveedor",
                licitacion=licitacion.name,
                proveedor=info_carpeta.get("proveedor", carpeta.name),
                rut=info_carpeta.get("rut", ""),
                error=str(exc),
            )
            if not detener:
                print(f"\nERROR FATAL: {exc}")
                print(f"Se detiene la corrida porque {args.backend} dejo de responder.")
                print("Se cancelan los proveedores en cola; los que estaban en curso se abandonan.")
                detener = True
                DETENER.set()
                for pendiente in futuros:
                    pendiente.cancel()
            return
        except Exception as exc:
            print(f"\nERROR EN PROVEEDOR {licitacion.name}/{carpeta.name}: {exc}")
            escribir_log(
                ruta_log,
                "error_proveedor",
                licitacion=licitacion.name,
                proveedor=info_carpeta.get("proveedor", carpeta.name),
                rut=info_carpeta.get("rut", ""),
                error=str(exc),
            )
            return

        completados += 1
        terminados_ok.add(carpeta)
        estado_licitacion = estado_licitaciones[licitacion.name]
        resultados = estado_licitacion["resultados"]
        reemplazar_resultado(resultados, resultado)
        contexto = metadata_licitaciones.get(licitacion.name, {})
        registrar_resultado(
            licitacion.name,
            contexto,
            resultado,
            todos,
            consumos,
            resumenes,
            archivos_ocr,
        )
        resumen_actual = resumenes[-1]
        escribir_log(
            ruta_log,
            "fin_proveedor",
            licitacion=licitacion.name,
            proveedor=resultado["proveedor"],
            rut=resultado["rut"],
            productos=resumen_actual["productos"],
            productos_validos=resumen_actual["productos_validos"],
            errores_modelo=resumen_actual["errores_modelo"],
            ocr=resumen_actual["ocr"],
            segundos=resultado.get("segundos"),
            segundos_lectura=resultado.get("segundos_lectura"),
            segundos_modelo=resultado.get("segundos_modelo"),
            llamadas=resultado.get("llamadas_modelo"),
            modo=resultado.get("modo"),
        )
        print(
            f"  [{completados}/{total_tareas}] {resultado.get('segundos', '?')}s | {licitacion.name} | "
            f"{resultado['proveedor'][:34]:34} "
            f"parciales={len(resultado['productos_parciales']):2} "
            f"finales={len(resultado['productos']):2} "
            f"ocr={len(resultado['necesita_ocr']):2} "
            f"consol={resultado['estado_consolidacion']}"
        )
        if resultado.get("segundos_modelo") is not None:
            print(f"      lectura={resultado['segundos_lectura']}s | modelo={resultado['segundos_modelo']}s | "
                  f"llamadas={resultado['llamadas_modelo']}")
        reproceso = resultado.get("reproceso") or {}
        if reproceso.get("modo") == "nada":
            print("      sin cambios: no habia archivos que rehacer")
        elif reproceso.get("modo") == "consolidar":
            print("      solo se repitio la consolidacion")
        elif reproceso.get("modo") == "parcial":
            for archivo_rehecho, motivo in sorted(reproceso["archivos"].items()):
                print(f"      rehecho [{motivo}] {archivo_rehecho[:60]}")
        elif reproceso.get("modo") == "completo":
            print("      rehecho completo (cambio de modelo o backend)")
        # Diagnostico visible por archivo: antes estos estados quedaban solo en JSON.
        for archivo in resultado["resultados_archivos"]:
            estado_ia = archivo.get("estado_ia", "")
            encontrados = archivo.get("productos_encontrados", 0)
            if archivo.get("paginas_vision") and not archivo.get("reutilizado"):
                print(
                    f"      [vision] {archivo.get('archivo','')[:47]} | "
                    f"paginas={', '.join(archivo['paginas_vision'])}"
                )
            archivo_sin_problema = resultado.get("modo") == "por_proveedor" and estado_ia in (
                "no_ejecutada", "ok", "ok_reparado", "ok_extraido")
            if not archivo_sin_problema and (
                    estado_ia not in ("no_ejecutada", "solo_texto", "ok", "ok_reparado") or encontrados == 0):
                print(
                    f"      [{archivo.get('tipo_documental','otro')}] "
                    f"{archivo.get('archivo','')[:47]} | lectura={archivo.get('estado_lectura')} "
                    f"| ia={estado_ia} | productos={encontrados}"
                )
                if archivo.get("observaciones"):
                    print(f"        observacion: {str(archivo['observaciones'])[:180]}")
        for producto in resultado["productos"]:
            print(f"    - {producto['producto'][:55]} | unitario={producto['precio_unitario']}")

        # checkpoint por proveedor para no perder el avance si el backend cae
        escribir_json_atomico(estado_licitacion["salida"], resultados)

    def esperar(futuros_pendientes):
        # wait() con timeout (en vez de as_completed sin timeout) para que Ctrl+C
        # se atienda tambien en Windows, donde una espera sin timeout lo bloquea.
        while futuros_pendientes:
            hechos, futuros_pendientes = wait(futuros_pendientes, timeout=1.0, return_when=FIRST_COMPLETED)
            for futuro in hechos:
                licitacion, carpeta = futuros[futuro]
                manejar(futuro, licitacion, carpeta)

    print(f"\nProveedores por procesar: {total_tareas} | en paralelo: {args.paralelo}")
    if tareas and not args.solo_texto and not args.sin_calentamiento:
        print("Calentando el modelo (activa LM Link y confirma que responde)...")
        segundos_calentamiento = calentar_modelo(args)
        if segundos_calentamiento is None:
            sys.exit(f"El modelo no respondio en {args.backend}. Revisa que este cargado (y LM Link conectado).")
        print(f"Modelo listo ({segundos_calentamiento} s).")
    pool = ThreadPoolExecutor(max_workers=args.paralelo)
    futuros = {pool.submit(trabajar, licitacion, carpeta, previo): (licitacion, carpeta)
               for licitacion, carpeta, previo in tareas}
    try:
        esperar(set(futuros))
    except KeyboardInterrupt:
        interrumpido = True
        detener = True
        DETENER.set()
        pool.shutdown(wait=False, cancel_futures=True)
        print("\nCtrl+C recibido: se cancelan los proveedores en cola y se abandona lo que")
        print("estaba en curso (se espera como maximo la respuesta de las llamadas activas).")
        print("Presiona Ctrl+C otra vez para salir de inmediato; los checkpoints guardados estan completos.")
        try:
            esperar({futuro for futuro in futuros if not futuro.done()})
        except KeyboardInterrupt:
            print("Salida forzada.")
            os._exit(130)
    finally:
        pool.shutdown(wait=True, cancel_futures=True)

    # Resultados "por revisar" cuyo reproceso no termino bien: siguen en el
    # checkpoint (nunca se quitaron) y se registran tal cual para el Excel.
    for licitacion, carpeta, previo in tareas:
        if previo is not None and carpeta not in terminados_ok:
            registrar_resultado(
                licitacion.name,
                metadata_licitaciones.get(licitacion.name, {}),
                previo,
                todos,
                consumos,
                resumenes,
                archivos_ocr,
            )
    if interrumpido:
        print("Corrida interrumpida por el usuario. Repite el comando para retomar.")

    if todos and not args.sin_cobertura:
        print("Calculando la cobertura documental de los productos...")
        try:
            cobertura_documental.anotar_productos(
                todos, raiz, workers=max(1, args.workers_cobertura), resultados=RESULTADOS_PROVEEDOR)
        except Exception as exc:  # la cobertura es informativa: nunca debe impedir el Excel
            print(f"  No se pudo calcular la cobertura documental: {exc}")
    if todos or resumenes:
        generar_excel(todos, consumos, resumenes, ruta_excel)

    print("\n" + "=" * 74)
    print(f"Productos consolidados: {len(todos)}")
    print(f"Productos validados OK: {sum(p.get('estado_validacion') == 'ok' for p in todos)}")
    print(f"Excel: {ruta_excel if todos or resumenes else 'no generado'}")
    print(f"Log: {ruta_log}")
    print(f"Archivos que requieren OCR: {len(archivos_ocr)}")
    for archivo_ocr in archivos_ocr:
        print(
            f"  - {archivo_ocr['licitacion']} | "
            f"{archivo_ocr['proveedor']} | {archivo_ocr['archivo']}"
        )
    print("=" * 74)
    escribir_log(
        ruta_log,
        "fin_corrida",
        productos=len(todos),
        productos_validos=sum(producto.get("estado_validacion") == "ok" for producto in todos),
        archivos_ocr=len(archivos_ocr),
        detenida=detener,
        interrumpida=interrumpido,
    )
    if interrumpido:
        sys.exit(130)


if __name__ == "__main__":
    main()
