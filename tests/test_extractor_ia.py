import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("extractor_ia", ROOT / "3_extraer_ia.py")
extractor = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(extractor)


class ExtractorIATest(unittest.TestCase):
    def test_alcance_comercial(self):
        casos = {
            "Notebook HP con mouse incluido": "equipo",
            "Asus B1503CVA-S75898X": "equipo",
            "Desktop Dell OptiPlex": "equipo",
            "All in One Lenovo": "equipo",
            "Monitor HP 24 pulgadas": "monitor",
            "Impresora Epson L5590": "impresora",
            "Notebook HP + Smart TV Samsung": None,
            "Smart TV con modo PC": None,
            "Toner para impresora HP": None,
            "Servicio de instalacion": None,
        }
        for descripcion, esperado in casos.items():
            with self.subTest(descripcion=descripcion):
                producto = {"producto": descripcion, "categoria": "equipo"}
                self.assertEqual(extractor.clasificar_producto(producto), esperado)

    def test_normaliza_separadores_chilenos(self):
        self.assertEqual(extractor.normalizar_numero("$ 18.527.280"), 18527280)
        self.assertEqual(extractor.normalizar_numero("$617,576"), 617576)
        self.assertEqual(extractor.normalizar_numero("123,50"), 123.5)

    def test_deriva_unitario_y_cantidad_solo_con_evidencia(self):
        desde_total = extractor.normalizar_producto(
            {
                "producto": "Notebook HP",
                "cantidad": 4,
                "precio_total": 400000,
                "precio_total_tipo": "linea",
            },
            "Proveedor", "1", "a.pdf", "economico",
        )
        desde_ambos = extractor.normalizar_producto(
            {
                "producto": "Monitor HP",
                "precio_unitario": 100000,
                "precio_total": 400000,
                "precio_total_tipo": "linea",
            },
            "Proveedor", "1", "b.pdf", "economico",
        )
        solo_total = extractor.normalizar_producto(
            {"producto": "Impresora Epson", "precio_total": 400000},
            "Proveedor", "1", "c.pdf", "economico",
        )
        self.assertEqual(desde_total["precio_unitario"], 100000)
        self.assertEqual(desde_ambos["cantidad"], 4)
        self.assertEqual(desde_ambos["cantidad_fuente"], "inferida_total_dividido_unitario")
        self.assertIsNone(solo_total["precio_unitario"])

        total_oferta = extractor.normalizar_producto(
            {
                "producto": "Notebook Dell",
                "cantidad": 10,
                "precio_total": 10000000,
                "precio_total_tipo": "oferta",
            },
            "Proveedor", "1", "d.pdf", "economico",
        )
        self.assertIsNone(total_oferta["precio_unitario"])

    def test_parsea_json_embebido_sin_regex_codiciosa(self):
        respuesta = 'diagnostico {"error": true}\nresultado {"productos": [{"producto": "Notebook"}]}'
        datos, estado = extractor.parsear_json(respuesta)
        self.assertEqual(estado, "ok_extraido")
        self.assertEqual(datos["productos"][0]["producto"], "Notebook")

    def test_fragmentacion_no_pierde_filas(self):
        texto = "\n".join(f"FILA {indice}: producto {indice}" for indice in range(100))
        fragmentos = extractor.dividir_texto(texto, 300, 2)
        combinado = "\n".join(fragmentos)
        self.assertGreater(len(fragmentos), 1)
        for indice in range(100):
            self.assertIn(f"FILA {indice}:", combinado)

    def test_detecta_suma_superior_al_total_ofertado(self):
        productos = [
            extractor.normalizar_producto(
                {
                    "producto": "Notebook ASUS VivoBook S16",
                    "categoria": "equipo",
                    "marca": "ASUS",
                    "cantidad": 1,
                    "precio_unitario": 4213990,
                    "precio_total": 4213990,
                    "precio_total_tipo": "linea",
                },
                "Austin", "1", "a.pdf", "economico",
            ),
            extractor.normalizar_producto(
                {
                    "producto": "NOTEBOOK 15,6",
                    "categoria": "equipo",
                    "cantidad": 30,
                    "precio_unitario": 807133,
                    "precio_total": 24213990,
                    "precio_total_tipo": "linea",
                },
                "Austin", "1", "b.pdf", "economico",
            ),
        ]
        validados = extractor.validar_productos(productos, productos, "25.125.550")
        self.assertTrue(all(
            "suma_productos_supera_total_oferta" in producto["alertas"]
            for producto in validados
        ))
        self.assertTrue(all(
            producto["estado_validacion"] == "inconsistente"
            for producto in validados
        ))

    def test_detecta_conflictos_entre_fuentes_del_mismo_modelo(self):
        productos = [
            {
                "producto": "Notebook HP EliteBook",
                "categoria": "equipo",
                "marca": "HP",
                "modelo": "840 G10",
                "cantidad": 10,
                "precio_unitario": 900000,
                "precio_total": 9000000,
                "precio_total_tipo": "linea",
                "confianza": "alta",
                "archivo_fuente": "a.pdf",
            },
            {
                "producto": "HP EliteBook 840 G10",
                "categoria": "equipo",
                "marca": "HP",
                "modelo": "840 G10",
                "cantidad": 12,
                "precio_unitario": 950000,
                "precio_total": 11400000,
                "precio_total_tipo": "linea",
                "confianza": "alta",
                "archivo_fuente": "b.pdf",
            },
        ]
        validados = extractor.validar_productos(productos, productos, None)
        self.assertTrue(all(
            producto["estado_validacion"] == "inconsistente"
            for producto in validados
        ))
        self.assertTrue(any(
            "conflicto_entre_fuentes_precio_unitario" in producto["alertas"]
            for producto in validados
        ))

    def test_consolidacion_por_lotes_no_corta_registros(self):
        registros = [
            {"producto": f"Notebook {indice}", "categoria": "equipo", "evidencia": "x" * 80}
            for indice in range(20)
        ]
        lotes = extractor.dividir_registros_json(registros, 500)
        self.assertEqual(sum(len(lote) for lote in lotes), len(registros))
        self.assertTrue(all(lote for lote in lotes))

    def test_deduplicacion_combina_datos_y_fuentes(self):
        tecnico = {
            "producto": "Notebook HP X1",
            "categoria": "equipo",
            "modelo": "X1",
            "cantidad": 2,
            "archivo_fuente": "tecnico.pdf",
        }
        economico = {
            "producto": "Notebook HP X1",
            "categoria": "equipo",
            "modelo": "X1",
            "precio_unitario": 100,
            "precio_total": 200,
            "archivo_fuente": "economico.pdf",
        }
        resultado = extractor.deduplicar_productos([tecnico, economico])
        self.assertEqual(len(resultado), 1)
        self.assertEqual(resultado[0]["cantidad"], 2)
        self.assertEqual(resultado[0]["precio_unitario"], 100)
        self.assertEqual(
            resultado[0]["fuentes_respaldo"],
            ["economico.pdf", "tecnico.pdf"],
        )

    def test_checkpoint_atomico(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "extraccion_ia.json"
            extractor.escribir_json_atomico(ruta, [{"rut": "1"}])
            self.assertEqual(json.loads(ruta.read_text(encoding="utf-8")), [{"rut": "1"}])
            self.assertFalse(ruta.with_suffix(".json.tmp").exists())

    def test_excel_simplificado_con_fuentes_por_campo(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "productos.xlsx"
            extractor.generar_excel(
                [{
                    "codigo": "1234-5-LE25",
                    "fecha_publicacion": "2025-11-01",
                    "proveedor": "Proveedor",
                    "producto": "Notebook HP",
                    "marca": "HP",
                    "cantidad": 2,
                    "precio_unitario": 500000,
                    "precio_total": 1000000,
                    "fuente_producto": "ficha.pdf",
                    "pagina_producto": 2,
                    "fuente_precio": "oferta.pdf",
                    "pagina_precio": 3,
                    "fuentes_respaldo": ["ficha.pdf", "oferta.pdf"],
                }],
                [], [], ruta,
            )
            libro = load_workbook(ruta, read_only=True, data_only=True)
            self.assertEqual(libro.sheetnames, ["Productos"])
            hoja = libro["Productos"]
            self.assertEqual(
                list(next(hoja.iter_rows(min_row=1, max_row=1, values_only=True))),
                [
                    "Licitación", "Fecha licitación", "Proveedor", "Producto", "Marca producto",
                    "Cantidad", "Precio unitario", "Precio total", "Fuente de información",
                ],
            )
            fila = list(next(hoja.iter_rows(min_row=2, max_row=2, values_only=True)))
            self.assertEqual(tuple(fila[0:8]), (
                "1234-5-LE25", "2025-11-01", "Proveedor", "Notebook HP", "HP", 2, 500000, 1000000,
            ))
            self.assertIn("Producto: ficha.pdf (p. 2)", fila[8])
            self.assertIn("Precio: oferta.pdf (p. 3)", fila[8])
            self.assertNotIn("Respaldo:", fila[8])
            libro.close()

    def test_prompt_proveedor_formatea_y_pide_reconstruir_tablas(self):
        prompt = extractor.PROMPT_PROVEEDOR.format(
            proveedor="Proveedor",
            total_oferta="1000",
            nota_imagenes="",
            documentos="tabla de prueba",
        )
        self.assertIn("Reconstruye cada fila", prompt)
        self.assertIn("archivo_precio", prompt)
        self.assertIn("tabla de prueba", prompt)

    def test_agrupacion_conserva_ultima_fila_y_encabezados_de_excel_largo(self):
        encabezado = "FILA 1: Producto || Cantidad || Unitario || Total"
        filas = [f"FILA {i}: Notebook HP {i} || 2 || 100 || 200" for i in range(2, 402)]
        bloque = {
            "archivo": "economico.xlsx", "pagina": None, "imagen": None,
            "texto": "\n".join(["[HOJA Oferta]", encabezado, *filas]),
        }
        grupos = extractor.agrupar_bloques([bloque], 1500, 4)
        self.assertGreater(len(grupos), 1)
        textos = [parte["texto"] for grupo in grupos for parte in grupo]
        combinado = "\n".join(textos)
        for fila in filas:
            self.assertEqual(combinado.count(fila + "\n") + int(combinado.endswith(fila)), 1)
        self.assertTrue(all(encabezado in texto for texto in textos))
        self.assertTrue(all("[HOJA Oferta]" in texto for texto in textos))

    def test_no_corta_una_fila_individual_mayor_que_el_limite(self):
        fila = "FILA 2: Notebook HP " + "x" * 2000
        bloque = {"archivo": "oferta.pdf", "pagina": 1, "imagen": None, "texto": fila}
        grupos = extractor.agrupar_bloques([bloque], 500, 4)
        self.assertEqual(grupos[0][0]["texto"], fila)

    def test_ocr_por_proveedor_abre_un_lote_y_conserva_paginas(self):
        args = SimpleNamespace(max_chars_proveedor=24000, vision=False, vision_solo_escaneadas=False,
                               ocr=True, max_paginas=20, tesseract_cmd=None, idioma_ocr="spa+eng",
                               dpi_ocr=180)
        def lectura(path, args, detalle):
            detalle.update(texto_paginas={0: "[PAGINA 1]", 1: "[PAGINA 2]"},
                           paginas_vision={0: "sin_texto", 1: "sin_texto"})
            return "", "texto_vision"

        def ocr(*posicionales, paginas, texto_paginas):
            self.assertEqual(paginas, {0, 1})
            for indice in paginas:
                texto_paginas[indice] = f"[PAGINA {indice + 1} - OCR]\nNotebook HP 2 $100.000 $200.000"
            return "texto OCR", "texto_ocr"

        carpeta = Path("proveedor")
        with mock.patch.object(extractor, "extraer_archivo", side_effect=lectura), \
                mock.patch.object(extractor, "extraer_pdf_ocr", side_effect=ocr) as llamada:
            bloques, registros, pendientes, _, _ = extractor.preparar_bloques_proveedor(
                carpeta, [carpeta / "economico.pdf"], args)
        llamada.assert_called_once()
        self.assertEqual([bloque["pagina"] for bloque in bloques], [1, 2])
        self.assertEqual(registros[0]["paginas_enviadas"], [1, 2])
        self.assertEqual(pendientes, [])

    def test_consumo_acumula_multiples_llamadas_del_mismo_archivo(self):
        args = SimpleNamespace(solo_texto=False, debug=False, max_chars_proveedor=300,
                               max_imagenes_proveedor=4, modelo="qwen-prueba", timeout=5,
                               num_ctx=8192, backend="lmstudio", max_tokens=1024,
                               reintentos_modelo=0, espera_reintento=0, pausa_archivo=0,
                               sin_consolidar=True, vision=False)
        bloque = {"archivo": "economico.txt", "pagina": None, "imagen": None,
                  "texto": "\n".join(["Notebook HP 2 $100.000 $200.000"] * 20)}
        registro = {"archivo": "economico.txt", "estado_ia": "no_ejecutada", "productos_encontrados": 0}
        with tempfile.TemporaryDirectory() as temporal:
            carpeta = Path(temporal)
            respuesta = ({"productos": []}, "ok", "{}", {"total_tokens": 110})
            with mock.patch.object(extractor, "descomprimir", return_value=[]), \
                    mock.patch.object(extractor, "preparar_bloques_proveedor",
                                      return_value=([bloque], [registro], [], [], set())), \
                    mock.patch.object(extractor, "consultar_modelo", return_value=respuesta) as llamada:
                resultado = extractor.procesar_oferta_por_proveedor(carpeta, args, None)
        consumo = resultado["resultados_archivos"][0]["consumo_local"]
        self.assertGreater(llamada.call_count, 1)
        self.assertEqual(consumo["llamadas"], llamada.call_count)
        self.assertEqual(consumo["total_tokens"], 110 * llamada.call_count)
        self.assertEqual(resultado["llamadas_modelo"], llamada.call_count)

    def test_proveedor_cruza_dos_anexos_en_una_llamada_y_reanuda_sin_llamar(self):
        import contextlib
        import io
        import sys

        producto = {
            "producto": "Notebook HP ProBook 440", "marca": "HP", "modelo": "ProBook 440",
            "cantidad": 2, "precio_unitario": 100000, "precio_total": 200000, "categoria": "equipo",
            "archivo_producto": "tecnico.txt", "archivo_precio": "economico.txt",
        }
        with tempfile.TemporaryDirectory() as temporal:
            raiz = Path(temporal)
            licitacion = raiz / "1234-5-LE25"
            carpeta = licitacion / "1__Proveedor"
            carpeta.mkdir(parents=True)
            extractor.escribir_json_atomico(carpeta / "oferta.json",
                                           {"proveedor": "Proveedor", "rut": "1", "total": "$200.000"})
            (carpeta / "economico.txt").write_text(
                "Producto Cantidad Unitario Total\nNotebook 2 $100.000 $200.000", encoding="utf-8")
            (carpeta / "tecnico.txt").write_text("Notebook HP ProBook 440", encoding="utf-8")
            (carpeta / "administrativo.txt").write_text(
                "Declaracion jurada: el oferente no tiene conflictos de interes.", encoding="utf-8")
            argumentos = ["3_extraer_ia.py", "--dir", str(raiz), "--por-proveedor",
                          "--backend", "lmstudio", "--modelo", "qwen-prueba", "--sin-calentamiento",
                          "--rampa", "0", "--metadata-csv", str(raiz / "sin_metadata.csv"),
                          "--excel", str(raiz / "productos.xlsx")]
            with mock.patch.object(extractor, "consultar_modelo",
                                   return_value=({"productos": [producto]}, "ok", "{}", {"total_tokens": 110})) as llamada, \
                    mock.patch.object(sys, "argv", argumentos), contextlib.redirect_stdout(io.StringIO()):
                extractor.main()
                llamada.assert_called_once()
                prompt = llamada.call_args.args[0]
                self.assertIn("economico.txt", prompt)
                self.assertIn("tecnico.txt", prompt)
                self.assertNotIn("administrativo.txt", prompt)
                extractor.main()
                llamada.assert_called_once()
            checkpoint = json.loads((licitacion / "extraccion_ia.json").read_text(encoding="utf-8"))
            self.assertEqual(checkpoint[0]["llamadas_modelo"], 1)
            self.assertEqual(checkpoint[0]["productos"][0]["precio_unitario"], 100000)
            self.assertEqual(checkpoint[0]["productos"][0]["fuente_producto"], "tecnico.txt")
            self.assertTrue((raiz / "productos.xlsx").is_file())

    def test_resultado_fallido_se_reprocesa(self):
        fallido = {
            "resultados_archivos": [{"estado_ia": "http_500"}],
            "estado_consolidacion": "omitida",
            "necesita_ocr": [],
        }
        completo = {
            "resultados_archivos": [{"estado_ia": "ok"}],
            "estado_consolidacion": "ok",
            "necesita_ocr": [],
        }
        self.assertTrue(extractor.resultado_requiere_reproceso(fallido))
        self.assertFalse(extractor.resultado_requiere_reproceso(completo))
        completo["modelo"] = "modelo-a"
        completo["backend"] = "lmstudio"
        self.assertTrue(
            extractor.resultado_requiere_reproceso(
                completo, modelo="modelo-b", backend="lmstudio"
            )
        )


if __name__ == "__main__":
    unittest.main()
