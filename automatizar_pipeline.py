#!/usr/bin/env python3
"""Ejecuta el pipeline completo en lotes mensuales y de forma reanudable."""

import argparse
import csv
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


RAIZ_PROYECTO = Path(__file__).resolve().parent
SCRIPT_FILTRAR = RAIZ_PROYECTO / "1_filtrar_licitaciones.py"
SCRIPT_SCRAPER = RAIZ_PROYECTO / "2_scraper_ofertas.py"
SCRIPT_OPENAI = RAIZ_PROYECTO / "3_extraer_openai_prueba.py"
ESTADOS_PRESERVAR = {
    "scraping_estado", "ia_estado", "n_ofertas", "n_archivos",
    "ia_proveedores", "ia_tokens", "ia_errores"
}


def leer_csv(ruta):
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with Path(ruta).open(encoding=encoding, newline="") as archivo:
                lector = csv.DictReader(archivo)
                return list(lector), list(lector.fieldnames or [])
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"No se pudo leer {ruta}")


def escribir_csv(ruta, filas, campos):
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with ruta.open("w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=campos, extrasaction="ignore")
        escritor.writeheader()
        escritor.writerows(filas)


def mes_publicacion(fecha):
    texto = (fecha or "").strip()
    formatos = (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"
    )
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato).strftime("%Y-%m")
        except ValueError:
            continue
    coincidencia = re.search(r"\b(20\d{2})[-/](\d{1,2})", texto)
    if coincidencia:
        return f"{coincidencia.group(1)}-{int(coincidencia.group(2)):02d}"
    coincidencia = re.search(r"\b\d{1,2}/(\d{1,2})/(20\d{2})", texto)
    if coincidencia:
        return f"{coincidencia.group(2)}-{int(coincidencia.group(1)):02d}"
    return "sin_fecha"


def dentro_del_rango(mes, desde, hasta):
    if mes == "sin_fecha":
        return not desde and not hasta
    return (not desde or mes >= desde) and (not hasta or mes <= hasta)


def agrupar_por_mes(filas, desde=None, hasta=None):
    grupos = {}
    for fila in filas:
        mes = mes_publicacion(fila.get("fecha_publicacion"))
        if dentro_del_rango(mes, desde, hasta):
            grupos.setdefault(mes, []).append(fila)
    return grupos


def combinar_estado(ruta, nuevas, campos_nuevos):
    existentes = []
    campos_existentes = []
    if Path(ruta).is_file():
        existentes, campos_existentes = leer_csv(ruta)

    campos = list(campos_nuevos)
    for campo in campos_existentes:
        if campo not in campos:
            campos.append(campo)
    for campo in ESTADOS_PRESERVAR:
        if campo not in campos:
            campos.append(campo)

    anteriores = {(fila.get("codigo") or "").strip(): fila for fila in existentes}
    combinadas = []
    codigos_nuevos = set()
    for nueva in nuevas:
        codigo = (nueva.get("codigo") or "").strip()
        codigos_nuevos.add(codigo)
        anterior = anteriores.get(codigo, {})
        fila = dict(nueva)
        for campo in ESTADOS_PRESERVAR:
            if anterior.get(campo) not in (None, ""):
                fila[campo] = anterior[campo]
        combinadas.append(fila)

    combinadas.extend(
        fila for codigo, fila in anteriores.items() if codigo not in codigos_nuevos
    )
    combinadas.sort(key=lambda fila: (fila.get("fecha_publicacion", ""), fila.get("codigo", "")))
    escribir_csv(ruta, combinadas, campos)


def ejecutar(comando, cwd, log):
    comando_visible = " ".join(f'"{parte}"' if " " in parte else parte for parte in comando)
    print(f"\n$ {comando_visible}")
    entorno = os.environ.copy()
    entorno["PYTHONUNBUFFERED"] = "1"
    entorno["PYTHONIOENCODING"] = "utf-8"
    with Path(log).open("a", encoding="utf-8") as archivo_log:
        archivo_log.write(f"\n[{datetime.now().isoformat(timespec='seconds')}] {comando_visible}\n")
        proceso = subprocess.Popen(
            comando,
            cwd=cwd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=entorno
        )
        assert proceso.stdout is not None
        for linea in proceso.stdout:
            salida_segura = linea.encode(
                sys.stdout.encoding or "utf-8", errors="replace"
            ).decode(sys.stdout.encoding or "utf-8", errors="replace")
            print(salida_segura, end="")
            archivo_log.write(linea)
        return proceso.wait()


def expandir_entradas(patrones):
    rutas = []
    for patron in patrones:
        coincidencias = glob.glob(patron)
        rutas.extend(coincidencias or [patron])
    unicas = []
    for ruta in rutas:
        absoluta = str(Path(ruta).resolve())
        if absoluta not in unicas:
            unicas.append(absoluta)
    faltantes = [ruta for ruta in unicas if not Path(ruta).is_file()]
    if faltantes:
        raise FileNotFoundError(f"No existen las entradas: {', '.join(faltantes)}")
    return unicas


def preparar_lotes(entradas, salida, desde, hasta, log):
    with tempfile.TemporaryDirectory(prefix="licitaciones_filtro_") as temporal:
        comando = [
            sys.executable, str(SCRIPT_FILTRAR), "--in", *entradas, "--out-dir", temporal
        ]
        if ejecutar(comando, RAIZ_PROYECTO, log) != 0:
            raise RuntimeError("Fallo la etapa de filtrado")

        ruta_todas = Path(temporal) / "licitaciones_computo.csv"
        ruta_scrapear = Path(temporal) / "para_scrapear.csv"
        todas, campos_todas = leer_csv(ruta_todas)
        scrapeables, campos_scrapear = leer_csv(ruta_scrapear)

        shutil.copy2(ruta_todas, salida / "catalogo_licitaciones.csv")
        shutil.copy2(ruta_scrapear, salida / "catalogo_para_scrapear.csv")

        grupos_todas = agrupar_por_mes(todas, desde, hasta)
        grupos_scrapear = agrupar_por_mes(scrapeables, desde, hasta)
        meses = sorted(set(grupos_todas) | set(grupos_scrapear))
        for mes in meses:
            carpeta = salida / mes
            carpeta.mkdir(parents=True, exist_ok=True)
            combinar_estado(
                carpeta / "licitaciones_computo.csv",
                grupos_todas.get(mes, []),
                campos_todas
            )
            combinar_estado(
                carpeta / "para_scrapear.csv",
                grupos_scrapear.get(mes, []),
                campos_scrapear
            )
        return meses


def estados_csv(ruta):
    filas, _ = leer_csv(ruta)
    estados = Counter((fila.get("scraping_estado") or "pendiente").strip().lower() for fila in filas)
    return filas, estados


def requiere_scraping(fila, reintentar_errores):
    estado = (fila.get("scraping_estado") or "pendiente").strip().lower()
    if estado in ("", "pendiente", "sin_cuadro"):
        return True
    return reintentar_errores and estado != "ok"


def es_error_ia(estado):
    texto = str(estado or "")
    return texto.startswith("error") or texto.startswith("http_")


def actualizar_estado_ia(ruta_csv, carpeta_ofertas):
    filas, campos = leer_csv(ruta_csv)
    for campo in ("ia_proveedores", "ia_tokens", "ia_errores"):
        if campo not in campos:
            campos.append(campo)

    for fila in filas:
        codigo = (fila.get("codigo") or "").strip()
        carpeta = carpeta_ofertas / codigo
        salida = carpeta / "extraccion_openai.json"
        if not salida.is_file():
            continue
        try:
            resultados = json.loads(salida.read_text(encoding="utf-8"))
        except Exception:
            fila["ia_estado"] = "error_json"
            continue

        esperados = len([
            path for path in carpeta.iterdir()
            if path.is_dir() and not path.name.startswith("_")
        ]) if carpeta.is_dir() else 0
        errores = 0
        tokens = 0
        for resultado in resultados:
            for archivo in resultado.get("resultados_archivos", []):
                if es_error_ia(archivo.get("estado_ia")):
                    errores += 1
                tokens += int((archivo.get("consumo_openai") or {}).get("total_tokens") or 0)
            if es_error_ia(resultado.get("estado_consolidacion")):
                errores += 1
            tokens += int((resultado.get("consumo_consolidacion") or {}).get("total_tokens") or 0)

        fila["ia_proveedores"] = len(resultados)
        fila["ia_tokens"] = tokens
        fila["ia_errores"] = errores
        if errores:
            fila["ia_estado"] = "error"
        elif esperados and len(resultados) >= esperados:
            fila["ia_estado"] = "ok"
        else:
            fila["ia_estado"] = "parcial"

    escribir_csv(ruta_csv, filas, campos)


def consolidar_reportes(salida, meses):
    reportes = [salida / mes / "resultados_openai.xlsx" for mes in meses]
    reportes = [ruta for ruta in reportes if ruta.is_file()]
    if not reportes:
        return None

    hojas = {"Productos": [], "Resumen": [], "Consumo": []}
    cabeceras = {}
    for ruta in reportes:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        for nombre in hojas:
            if nombre not in libro.sheetnames:
                continue
            hoja = libro[nombre]
            filas = hoja.iter_rows(values_only=True)
            encabezado = next(filas, None)
            if not encabezado:
                continue
            cabeceras.setdefault(nombre, list(encabezado))
            for valores in filas:
                fila = dict(zip(encabezado, valores))
                if nombre == "Resumen" and fila.get("codigo") == "TOTAL":
                    continue
                hojas[nombre].append(fila)
        libro.close()

    if hojas["Resumen"]:
        columnas_suma = {
            "proveedores_procesados", "productos", "llamadas_api", "errores_api",
            "archivos_ocr", "archivos_no_soportados", "prompt_tokens",
            "completion_tokens", "total_tokens", "cached_tokens", "reasoning_tokens"
        }
        total = {campo: "" for campo in cabeceras["Resumen"]}
        total["codigo"] = "TOTAL"
        total["nombre_licitacion"] = "Todos los meses"
        for campo in columnas_suma:
            total[campo] = sum(int(fila.get(campo) or 0) for fila in hojas["Resumen"])
        hojas["Resumen"].append(total)

    destino = salida / "resultados_consolidados.xlsx"
    libro_salida = Workbook()
    libro_salida.remove(libro_salida.active)
    for nombre in ("Productos", "Resumen", "Consumo"):
        hoja = libro_salida.create_sheet(nombre)
        columnas = cabeceras.get(nombre, [])
        for numero_columna, campo in enumerate(columnas, 1):
            celda = hoja.cell(1, numero_columna, campo)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
        for numero_fila, fila in enumerate(hojas[nombre], 2):
            for numero_columna, campo in enumerate(columnas, 1):
                hoja.cell(numero_fila, numero_columna, fila.get(campo))
        if columnas:
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
    libro_salida.save(destino)
    return destino


def crear_lock(salida):
    ruta = salida / ".pipeline.lock"
    try:
        descriptor = os.open(ruta, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise RuntimeError(
            f"Ya existe {ruta}. Si no hay otra corrida activa, elimina ese archivo y reintenta."
        ) from exc
    with os.fdopen(descriptor, "w", encoding="utf-8") as archivo:
        archivo.write(f"pid={os.getpid()}\ninicio={datetime.now().isoformat()}\n")
    return ruta


def main():
    parser = argparse.ArgumentParser(
        description="Filtra, divide por mes, scrapea, extrae con OpenAI y consolida resultados."
    )
    parser.add_argument("--in", dest="entradas", nargs="+", required=True,
                        help="CSV(s) descargados desde Mercado Publico")
    parser.add_argument("--salida", default="lotes", help="Carpeta de lotes mensuales")
    parser.add_argument("--modelo", default="gpt-4.1-mini")
    parser.add_argument("--desde", help="Primer mes incluido, formato YYYY-MM")
    parser.add_argument("--hasta", help="Ultimo mes incluido, formato YYYY-MM")
    parser.add_argument("--solo-preparar", action="store_true",
                        help="Solo filtra y crea los CSV mensuales; no usa red ni API")
    parser.add_argument("--omitir-scraping", action="store_true")
    parser.add_argument("--omitir-ia", action="store_true")
    parser.add_argument("--reintentar-errores", action="store_true")
    parser.add_argument("--todos-anexos", action="store_true",
                        help="Incluye anexos administrativos en la descarga")
    parser.add_argument("--sin-consolidar", action="store_true",
                        help="Omite la llamada OpenAI de consolidacion por proveedor")
    parser.add_argument("--rehacer-ia", action="store_true",
                        help="Descarta resultados OpenAI previos y vuelve a consumir tokens")
    args = parser.parse_args()

    for valor, nombre in ((args.desde, "--desde"), (args.hasta, "--hasta")):
        if valor and not re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", valor):
            parser.error(f"{nombre} debe usar el formato YYYY-MM")
    if args.desde and args.hasta and args.desde > args.hasta:
        parser.error("--desde no puede ser posterior a --hasta")

    entradas = expandir_entradas(args.entradas)
    salida = Path(args.salida).resolve()
    salida.mkdir(parents=True, exist_ok=True)
    lock = crear_lock(salida)
    log = salida / "pipeline.log"
    fallos = []

    try:
        meses = preparar_lotes(entradas, salida, args.desde, args.hasta, log)
        print(f"\nLotes preparados: {', '.join(meses) if meses else 'ninguno'}")
        if args.solo_preparar:
            return

        for indice, mes in enumerate(meses, 1):
            carpeta = salida / mes
            csv_mes = carpeta / "para_scrapear.csv"
            filas, estados_antes = estados_csv(csv_mes)
            if not filas:
                print(f"\n[{indice}/{len(meses)}] {mes}: sin licitaciones scrapeables")
                continue

            print("\n" + "=" * 78)
            print(f"[{indice}/{len(meses)}] LOTE {mes} | licitaciones={len(filas)} | estados={dict(estados_antes)}")
            print("=" * 78)

            pendientes = [
                fila for fila in filas if requiere_scraping(fila, args.reintentar_errores)
            ]
            if not args.omitir_scraping and pendientes:
                comando = [sys.executable, str(SCRIPT_SCRAPER), "--csv", "para_scrapear.csv"]
                if args.reintentar_errores:
                    comando.append("--reintentar-errores")
                comando.append("--incluye-tecnicos")
                if args.todos_anexos:
                    comando.append("--todos-anexos")
                codigo = ejecutar(comando, carpeta, log)
                if codigo != 0:
                    fallos.append(f"{mes}: scraper termino con codigo {codigo}")
            elif not args.omitir_scraping:
                print(f"{mes}: scraping ya completo; no se abre Chromium")

            carpeta_ofertas = carpeta / "ofertas"
            if args.omitir_ia or not carpeta_ofertas.is_dir():
                continue
            if not any(path.is_dir() for path in carpeta_ofertas.iterdir()):
                print(f"{mes}: no hay ofertas descargadas para extraer")
                continue

            comando = [
                sys.executable, str(SCRIPT_OPENAI),
                "--dir", str(carpeta_ofertas),
                "--modelo", args.modelo,
                "--limite-proveedores", "0",
                "--metadata-csv", str(csv_mes),
                "--excel", str(carpeta / "resultados_openai.xlsx")
            ]
            if args.sin_consolidar:
                comando.append("--sin-consolidar")
            if args.rehacer_ia:
                comando.append("--rehacer")
            codigo = ejecutar(comando, RAIZ_PROYECTO, log)
            if codigo != 0:
                fallos.append(f"{mes}: extractor OpenAI termino con codigo {codigo}")
            actualizar_estado_ia(csv_mes, carpeta_ofertas)

        meses_con_reporte = sorted(
            path.name for path in salida.iterdir()
            if path.is_dir() and re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", path.name)
        )
        reporte = consolidar_reportes(salida, meses_con_reporte)
        print("\n" + "=" * 78)
        print(f"Meses procesados: {len(meses)}")
        print(f"Reporte consolidado: {reporte or 'no generado'}")
        print(f"Log: {log}")
        if fallos:
            print("Incidencias:")
            for fallo in fallos:
                print(f"  - {fallo}")
            raise SystemExit(1)
        print("PIPELINE COMPLETADO")
    finally:
        lock.unlink(missing_ok=True)


if __name__ == "__main__":
    main()